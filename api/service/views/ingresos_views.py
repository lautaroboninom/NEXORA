from django.db import connection, transaction
import re
import time
import json
import logging
from contextlib import nullcontext
from types import SimpleNamespace
from datetime import date, datetime, timedelta
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from django.utils.decorators import method_decorator
from rest_framework import permissions
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView
from django.http import HttpResponse
from io import BytesIO
from openpyxl import Workbook
from django.conf import settings
from django.core.mail import send_mail, get_connection, EmailMessage

from .helpers import (
    _fetchall_dicts,
    _get_motivo_enum_values,
    _get_motivo_enum_values_raw,
    _map_motivo_to_db_label,
    _motivo_is_cotizacion_equipo,
    _norm_txt,
    _rol,
    _set_audit_user,
    _frontend_url,
    _email_append_footer_text,
    exec_returning,
    exec_void,
    last_insert_id,
    os_label,
    q,
    require_jefe,
    require_roles,
    require_permission,
    ensure_default_locations,
    money,
)
from .preventivos_sync import sync_plan_from_ingreso_fecha_servicio
from .quotes_views import marcar_quote_no_aplica
from .mg_state import is_mg_owner_device, normalize_mg, resolve_mg_flags
from ..ingreso_duplicados import ingreso_es_de_hoy
from ..mg_codes import normalize_mg_code, normalize_mg_input, sql_is_mg_code, sql_mg_key
from ..release_gate import sql_presupuesto_bloquea_salida
from ..rejected_budget import (
    ensure_rejected_budget_charge_schema,
    get_rejected_quote_summary,
    get_stored_rejected_budget_fields,
    has_rejected_budget_charge_schema,
)
from ..derivations import (
    active_derivation_block_payload,
    attach_active_derivations,
    get_active_derivation,
)
from ..bejerman_sync import (
    enqueue_client_ready_transfer_for_ingreso,
    enqueue_demo_ready_transfer_for_ingreso,
    enqueue_location_transfer_for_ingreso,
    enqueue_stock_alta_for_ingreso,
    enqueue_stock_baja_for_ingreso,
    enqueue_stock_transfer_for_ingreso,
    enqueue_stock_venta_for_ingreso,
    ingreso_is_demo_return,
    ingreso_is_internal_equipment,
    ingreso_targets_rental_stock,
)
from ..constants import bejerman_deposit_for_location, canonical_location_name
from ..rental_shelf import (
    RENTAL_SHELF_LOCATION_NAME,
    is_rental_shelf_location_name,
    move_ingreso_to_rental_shelf,
)
from ..bejerman_sdk import (
    BejermanSDKClient,
    BejermanSdkConfigError,
    BejermanSdkResponseError,
    BejermanSdkUnavailable,
    as_number,
    as_string,
    build_partida_rows,
    build_stock_rows,
    first_value,
    normalize_serial_for_lookup,
    records_from_response,
)

logger = logging.getLogger(__name__)
from ..serializers import (
    IngresoDetailSerializer,
    IngresoDetailWithAccesoriosSerializer,
    IngresoListItemSerializer,
)
from ..permissions import require_any_permission, user_has_any_permission, user_has_permission
from ..warranty import compute_warranty_from_sale_date
from ..bejerman_catalog_cache import (
    article_descriptions,
    deposit_response as cached_deposit_response,
    last_good_deposit_response,
    stock_rows as cached_stock_rows,
)
from ..bejerman_sales import find_latest_remito_by_serial, find_latest_sale_by_serial
from ..bejerman_companies import require_company
from ..bejerman_ris import (
    BejermanRisError,
    BejermanRisPreflightError,
    emit_or_get_ris_batch,
    force_stock_entry_for_ingreso,
    get_ris_status_for_ingreso,
    parse_ris_document_mode,
    preflight_ris_for_request_payload,
    register_ris_batch,
    RIS_DOCUMENT_MODE_REGISTER,
)
from ..bejerman_rental_remito import (
    RentalRemitoError,
    execute_rental_remito_action,
    fetch_rental_remito_pdf,
    get_rental_remito_state,
)
from ..bejerman_ingreso_conflicts import (
    attach_remito_number as attach_conflict_remito_number,
    record_conflict,
)
from ..notifications import (
    filter_email_recipients_by_preference,
    notify_estado_patrimonial,
    notify_ingreso_asignado,
    notify_ingreso_conflicto_abierto,
    notify_solicitud_asignacion,
    notify_solicitud_baja,
    resolve_solicitud_asignacion,
    solicitud_asignacion_email_recipients,
    solicitud_asignacion_user_ids,
)
from ..patrimonial_notifications import notify_patrimonial_email
from ..repair_ready_notifications import (
    notify_no_se_repara_for_remito,
    notify_repair_ready_for_remito,
    resolve_presupuesto_rechazado_reparado,
)


SALE_TICKET_STATES = ("vendido_pendiente_entrega", "vendido_entregado")
TECH_FINAL_REPAIR_STATES = ("reparado", "controlado_sin_defecto", "no_se_repara")
REPAIR_CLOSE_STATES = ("liberado", "entregado", "alquilado", "baja", *SALE_TICKET_STATES)
RESTOREABLE_REPAIR_STATES = ("diagnosticado", "presupuestado", "reparar")


def _compute_factory_warranty_from_bejerman_sale(
    numero_serie,
    *,
    brand_id=None,
    model_id=None,
    company_key=None,
    bejerman_sale_payload=None,
):
    serial = (numero_serie or "").strip()
    # La garantía de fábrica no depende de la empresa que factura la reparación: el mismo
    # equipo puede haber salido por SEPID o por MG BIO (o por SEPID a MG BIO y de ahí al
    # cliente final). Se busca la venta más reciente en todas las empresas; `company_key`
    # queda sólo como dato de contexto para el meta de la respuesta.
    sale_payload = find_latest_sale_by_serial(serial)
    estimated = False
    if not sale_payload.get("found") and isinstance(bejerman_sale_payload, dict):
        payload_date = (
            bejerman_sale_payload.get("issueDate")
            or bejerman_sale_payload.get("saleDate")
            or bejerman_sale_payload.get("fecha_venta")
        )
        payload_serial = (
            bejerman_sale_payload.get("serial")
            or bejerman_sale_payload.get("numero_serie")
            or bejerman_sale_payload.get("itemPartida")
            or serial
        )
        if payload_date and normalize_serial_for_lookup(payload_serial) == normalize_serial_for_lookup(serial):
            sale_payload = {
                "found": True,
                "source": "bejerman_sale_payload",
                "serial": payload_serial,
                "normalizedSerial": normalize_serial_for_lookup(payload_serial),
                "issueDate": payload_date,
                "documentLabel": bejerman_sale_payload.get("documentLabel") or "",
                "customerCode": bejerman_sale_payload.get("customerCode") or "",
                "customerName": bejerman_sale_payload.get("customerName") or "",
                "customerCuit": bejerman_sale_payload.get("customerCuit") or "",
                "articleCode": bejerman_sale_payload.get("articleCode") or "",
                "articleDescription": bejerman_sale_payload.get("articleDescription") or "",
                "cacheSaleItemId": bejerman_sale_payload.get("cacheSaleItemId") or "",
                "syncedAt": bejerman_sale_payload.get("syncedAt"),
            }

    if not sale_payload.get("found"):
        # Sin factura queda el remito de salida. No es la fecha de venta al cliente final
        # (el equipo pudo pasar meses en el distribuidor), pero es anterior a ella: la
        # garantía que sale de acá es la más corta posible, nunca una de más. Va marcada
        # como estimada para que la decida una persona.
        remito_payload = find_latest_remito_by_serial(serial)
        if remito_payload.get("found"):
            sale_payload = remito_payload
            estimated = True

    if not sale_payload.get("found"):
        return {
            "garantia": None,
            "vence_el": None,
            "fecha_venta": None,
            "days": None,
            "meta": {
                "source": "bejerman_sale",
                "found": False,
                "normalizedSerial": normalize_serial_for_lookup(serial),
            },
        }

    calc = compute_warranty_from_sale_date(
        sale_payload.get("issueDate"),
        numero_serie=sale_payload.get("serial") or serial,
        brand_id=brand_id,
        model_id=model_id,
        source="bejerman_sale",
    )
    meta = calc.get("meta") or {}
    meta.update(
        {
            "source": "bejerman_sale",
            "saleSource": sale_payload.get("source") or "bejerman_sale_cache",
            "documentLabel": sale_payload.get("documentLabel") or "",
            "customerCode": sale_payload.get("customerCode") or "",
            "customerName": sale_payload.get("customerName") or "",
            "customerCuit": sale_payload.get("customerCuit") or "",
            "articleCode": sale_payload.get("articleCode") or "",
            "articleDescription": sale_payload.get("articleDescription") or "",
            "cacheSaleItemId": sale_payload.get("cacheSaleItemId") or "",
            "cacheSyncedAt": sale_payload.get("syncedAt"),
            "companyKey": sale_payload.get("companyKey") or company_key or "",
            "found": True,
            "estimated": estimated,
            "dateBasis": "remito" if estimated else "factura",
            "documentType": sale_payload.get("documentType") or "",
        }
    )
    calc["meta"] = meta
    return calc


INGRESO_DETAIL_READ_PERMISSIONS = [
    "page.ingresos_history",
    "page.work_queues",
    "page.budget_queues",
    "page.logistics",
    "page.service_sheet_principal",
    "action.ingreso.edit_basics",
    "action.ingreso.edit_diagnosis",
    "action.ingreso.edit_location",
    "action.ingreso.edit_delivery",
    "action.ingreso.manage_derivations",
    "action.ingreso.repair_transitions",
    "action.ingreso.baja_alta",
    "action.presupuesto.manage",
]


BEJERMAN_STATUS_DEPOSIT_CODE_FIELDS = (
    "Deposito_Codigo",
    "DepositoCodigo",
    "Dep_Codigo",
    "dep_Cod",
    "Dep_Cod",
    "Codigo",
    "code",
)
BEJERMAN_STATUS_DEPOSIT_NAME_FIELDS = (
    "Deposito_Descripcion",
    "DepositoDescripcion",
    "Deposito_Nombre",
    "DepositoNombre",
    "Descripcion",
    "Nombre",
    "dep_Desc",
    "Dep_Desc",
    "name",
)
BEJERMAN_STATUS_ARTICLE_DESCRIPTION_FIELDS = (
    "Articulo_Descripcion",
    "ArticuloDescripcion",
    "Art_DescripcionGeneral",
    "Art_DescripcionReducida",
    "Art_DescricpionAdicional",
    "Art_DescripcionAdicional",
    "Descripcion",
    "Nombre",
    "article_description",
)
BEJERMAN_STATUS_EXPIRATION_FIELDS = (
    "Partida_FechaVtoIngreso",
    "PartidaFechaVtoIngreso",
    "FechaVtoIngreso",
    "FechaVencimientoIngreso",
    "FechaVencimiento",
    "FechaVto",
    "Partida_FechaVencimiento",
    "Partida_FechaVto",
    "expirationDate",
)


# Helpers locales (sin mover módulos por ahora)
def _equipolabel_row(r):
    try:
        tipo = (r.get("tipo_equipo") or "").strip()
        marca = (r.get("marca") or "").strip()
        modelo = (r.get("modelo") or "").strip()
        variante = (r.get("equipo_variante") or r.get("variante") or r.get("modelo_variante") or r.get("variante_nombre") or "").strip()
        modelo_comp = (f"{modelo} {variante}" if modelo else variante).strip()
        parts = [p for p in [tipo, marca, modelo_comp] if p]
        return " | ".join(parts) if parts else "-"
    except Exception:
        return "-"


def _mg_vendido(r) -> bool:
    try:
        mg_estado = (r.get("mg_estado") or "").strip().lower()
        estado = (r.get("estado") or "").strip().lower()
        return bool(r.get("mg_inactivo_venta")) or mg_estado == "inactivo_venta" or estado in SALE_TICKET_STATES
    except Exception:
        return False


def _ns_label(r):
    try:
        interno = (r.get("numero_interno") or "").strip()
        serie = (r.get("numero_serie") or "").strip()
        if _mg_vendido(r):
            return serie or interno or "-"
        return interno or serie or "-"
    except Exception:
        return "-"


def _mg_label(r):
    """MG para columnas de listado: el de un equipo vendido ya es histórico."""
    try:
        if _mg_vendido(r):
            return "-"
        return (r.get("numero_interno") or "").strip() or "-"
    except Exception:
        return "-"


# Tipear "26" en vez de "2026" en el año da una fecha perfectamente parseable:
# 0026-08-28. NEXORA la guardaba tal cual y después se la mandaba a Bejerman, que
# contestaba "Se han detectado inconsistencias en el archivo de texto" — un mensaje
# que no le dice a nadie que el problema es la fecha. Pasó con la OS 29664 (RDA de
# ALCLA, 31/08/2026) y con la 29288. Ninguna fecha de este negocio cae fuera de esta
# ventana, así que fuera de ella es un error de tipeo y hay que rechazarlo en la
# puerta, no doce pasos más adelante.
_MIN_PLAUSIBLE_YEAR = 1990
_MAX_PLAUSIBLE_YEAR_AHEAD = 5


def _year_is_plausible(year: int) -> bool:
    return _MIN_PLAUSIBLE_YEAR <= int(year) <= timezone.localdate().year + _MAX_PLAUSIBLE_YEAR_AHEAD


def _parse_datetime_or_date(value):
    parsed = _parse_datetime_or_date_raw(value)
    if parsed is not None and not _year_is_plausible(parsed.year):
        return None
    return parsed


def _parse_datetime_or_date_raw(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, 0, 0, 0)
    s = str(value).strip()
    if not s:
        return None
    s_norm = s[:-1] + "+00:00" if s.endswith("Z") else s
    dt = parse_datetime(s_norm) or parse_datetime(s)
    if dt:
        return dt
    d = parse_date(s)
    if d:
        return datetime(d.year, d.month, d.day, 0, 0, 0)
    for sep in ("/", "-", "."):
        parts = s.split(sep)
        if len(parts) == 3 and all(p.isdigit() for p in parts):
            if len(parts[0]) == 4:
                year, month, day = parts
            elif len(parts[2]) == 4:
                day, month, year = parts
            else:
                continue
            try:
                return datetime(int(year), int(month), int(day), 0, 0, 0)
            except ValueError:
                return None
    return None


def _has_mg_schema() -> bool:
    try:
        with (transaction.atomic() if connection.in_atomic_block else nullcontext()):
            if connection.vendor == "postgresql":
                row = q(
                    """
                    SELECT 1
                      FROM information_schema.columns
                     WHERE table_name='devices'
                       AND column_name='mg_estado'
                       AND table_schema = ANY(current_schemas(true))
                     LIMIT 1
                    """,
                    one=True,
                )
                return bool(row)
            row = q(
                """
                SELECT 1
                  FROM information_schema.columns
                 WHERE table_name='devices'
                   AND column_name='mg_estado'
                 LIMIT 1
                """,
                one=True,
            )
            return bool(row)
    except Exception:
        return False


def _has_table_column(table_name: str, column_name: str) -> bool:
    try:
        with (transaction.atomic() if connection.in_atomic_block else nullcontext()):
            if connection.vendor == "postgresql":
                row = q(
                    """
                    SELECT 1
                      FROM information_schema.columns
                     WHERE table_name=%s
                       AND column_name=%s
                       AND table_schema = ANY(current_schemas(true))
                     LIMIT 1
                    """,
                    [table_name, column_name],
                    one=True,
                )
                return bool(row)
            row = q(
                """
                SELECT 1
                  FROM information_schema.columns
                 WHERE table_name=%s
                   AND column_name=%s
                 LIMIT 1
                """,
                [table_name, column_name],
                one=True,
            )
            return bool(row)
    except Exception:
        return False


def _has_table(table_name: str) -> bool:
    try:
        with (transaction.atomic() if connection.in_atomic_block else nullcontext()):
            if connection.vendor == "postgresql":
                row = q(
                    """
                    SELECT 1
                      FROM information_schema.tables
                     WHERE table_name=%s
                       AND table_schema = ANY(current_schemas(true))
                     LIMIT 1
                    """,
                    [table_name],
                    one=True,
                )
                return bool(row)
            row = q(
                """
                SELECT 1
                  FROM information_schema.tables
                 WHERE table_name=%s
                 LIMIT 1
                """,
                [table_name],
                one=True,
            )
            return bool(row)
    except Exception:
        return False


def _has_mg_sale_extended_schema() -> bool:
    return _has_table_column("devices", "mg_venta_customer_id") and _has_table_column(
        "devices", "mg_venta_numero_alternativo"
    )


def _swap_list_select_sql(alias: str = "t") -> str:
    """Columnas del cambio de equipo para los listados de ingresos.

    Sin esto, el modal de salida abierto desde Reparados/Aprobados recibía
    `serial_cambio === undefined` y dejaba el campo del cambio muerto.
    """
    if _has_table_column("ingresos", "serial_cambio"):
        base = f"COALESCE({alias}.serial_cambio,'') AS serial_cambio,"
    else:
        base = "'' AS serial_cambio,"
    if _has_device_swap_schema():
        # `cambio_sin_entrega` no es opcional acá: una parte de pago no tiene equipo
        # entregado, así que sin esta columna el listado la ve como "sin registrar" y el
        # modal de salida bloquea la liberación pidiendo un cambio que ya está cargado.
        return f"""
          {base}
          {alias}.device_cambio_id AS device_cambio_id,
          COALESCE({alias}.cambio_destino_viejo,'') AS cambio_destino_viejo,
          COALESCE({alias}.cambio_sin_entrega, FALSE) AS cambio_sin_entrega,
        """
    return f"""
      {base}
      NULL::int AS device_cambio_id,
      '' AS cambio_destino_viejo,
      FALSE AS cambio_sin_entrega,
    """


def _owner_list_select_sql(ingreso_alias: str = "t", device_alias: str = "d") -> str:
    """Nombre del dueño real para los listados de ingresos.

    'Particular' y 'MG BIO' son clientes contenedores: la columna Cliente de las
    listas mostraba el contenedor y el operador no veía de quién es el equipo.
    El dato vive en el ingreso y se replica al equipo, así que se toma el del
    ingreso y se cae al del equipo para los ingresos viejos.
    """
    return f"{_owner_name_expr(ingreso_alias, device_alias)} AS propietario_nombre,"


def _owner_name_expr(ingreso_alias: str = "t", device_alias: str = "d") -> str:
    """La misma expresión, sin alias, para usar en WHERE y ORDER BY.

    El listado mostraba el dueño pero se buscaba y se ordenaba por la razón social del
    cliente contenedor: escribías en el buscador el nombre que estabas viendo en pantalla
    y no aparecía nada.
    """
    parts = []
    if _has_table_column("ingresos", "propietario_nombre"):
        parts.append(f"NULLIF({ingreso_alias}.propietario_nombre,'')")
    if _has_table_column("devices", "propietario_nombre"):
        parts.append(f"NULLIF({device_alias}.propietario_nombre,'')")
    if _has_table_column("ingresos", "cliente_original_id"):
        # Último recurso, para el equipo que tomamos en parte de pago: su `customer_id` ya
        # es MG BIO, y un cliente empresa no deja nombre en `propietario_nombre`. Va por
        # subconsulta y no por JOIN para no tocar el FROM de cada listado.
        parts.append(
            "NULLIF((SELECT co.razon_social FROM customers co"
            f" WHERE co.id = {ingreso_alias}.cliente_original_id), '')"
        )
    if not parts:
        return "''"
    return f"COALESCE({', '.join(parts)}, '')"


# Estados en los que el equipo sigue en el taller. Una orden liberada que hoy está en uno de
# éstos volvió atrás: la liberación se revirtió y el equipo nunca salió.
_ESTADOS_EN_TALLER = (
    "ingresado",
    "diagnosticado",
    "presupuestado",
    "reparar",
    "controlado_sin_defecto",
    "no_se_repara",
    "reparado",
    "derivado",
)


def _salida_taller_lateral_sql(ingreso_alias: str = "t") -> str:
    """JOIN que trae cuándo se liberó el ingreso, para acompañar a `_salida_taller_expr`.

    El `liberado` de una orden que después volvió al taller no cuenta. El evento queda en el
    historial aunque la liberación se revierta (limpieza del 14/07/2026, reversión del
    09/09/2026), y contarlo como salida le regalaba garantía de reparación al equipo que
    reingresaba: la OS 29875 la heredó del liberado revertido de la OS 27479, que quedaba
    justo a 90 días.
    """
    en_taller = ", ".join(f"'{estado}'" for estado in _ESTADOS_EN_TALLER)
    return f"""
      LEFT JOIN LATERAL (
        SELECT MAX(ev.ts) AS ts
          FROM ingreso_events ev
         WHERE ev.ingreso_id = {ingreso_alias}.id
           AND ev.a_estado = 'liberado'
           AND COALESCE({ingreso_alias}.estado::text, '') NOT IN ({en_taller})
      ) {ingreso_alias}_lib ON TRUE
    """


def _salida_taller_expr(ingreso_alias: str = "t") -> str:
    """Cuándo salió el equipo del taller.

    `fecha_entrega` sólo se escribe cuando se emite el RSS o cuando alguien marca
    entregado a mano; imprimir la orden de salida deja el ingreso en 'liberado' y
    nada más (ver `reportes_views.RemitoSalidaPdfView`). Como la mayoría de las
    entregas se cierran así, mirar sólo `fecha_entrega` daba por "nunca salido" a
    casi todo lo liberado, y con eso se perdían la garantía de reparación automática
    y los trabajos del último servicio.
    """
    return f"COALESCE({ingreso_alias}.fecha_entrega, {ingreso_alias}_lib.ts)"


def _customer_search_expr(customer_alias: str = "c", ingreso_alias: str = "t", device_alias: str = "d") -> str:
    """Lo que la columna Cliente muestra de verdad: el dueño si lo hay, si no el cliente.

    Espeja `_ingreso_customer_label`, que es la que decide en Python. Se usa para buscar,
    filtrar y ordenar, así que lo que se ve en pantalla es lo que se puede tipear.
    """
    owner = _owner_name_expr(ingreso_alias, device_alias)
    if owner == "''":
        return f"COALESCE({customer_alias}.razon_social,'')"
    # El dueño sólo desplaza al cliente cuando el cliente es un contenedor ('Particular' /
    # 'MG BIO'); para un cliente real la razón social sigue mandando.
    # Con STRPOS y no con LIKE '%mg%': el '%' hay que escaparlo como '%%' para psycopg2 y
    # esta expresión se interpola en consultas que a veces corren sin parámetros.
    razon = f"LOWER(COALESCE({customer_alias}.razon_social,''))"
    contenedor = (
        f"(TRIM({razon}) = 'particular'"
        f" OR (STRPOS({razon}, 'mg') > 0 AND STRPOS({razon}, 'bio') > 0))"
    )
    return (
        f"CASE WHEN {contenedor} AND NULLIF({owner},'') IS NOT NULL"
        f" THEN {owner} ELSE COALESCE({customer_alias}.razon_social,'') END"
    )


def _is_owner_display_customer(name: str) -> bool:
    """'Particular' y 'MG BIO' son clientes contenedores, no el dueño del equipo."""
    normalized = (name or "").strip().lower()
    if not normalized:
        return False
    if normalized == "particular":
        return True
    return "mg" in normalized and "bio" in normalized


def _ingreso_customer_label(row: dict) -> str:
    """Nombre a mostrar en la columna Cliente de listados y exports."""
    customer = (row.get("razon_social") or row.get("cliente") or "").strip()
    owner = (row.get("propietario_nombre") or "").strip()
    if owner and _is_owner_display_customer(customer):
        return owner
    return customer or owner


def _mg_list_select_sql(alias: str = "d") -> str:
    alt_sql = (
        f"COALESCE({alias}.mg_venta_numero_alternativo,'') AS mg_venta_numero_alternativo,"
        if _has_mg_sale_extended_schema()
        else "'' AS mg_venta_numero_alternativo,"
    )
    if _has_mg_schema():
        return f"""
          COALESCE({alias}.mg_estado,'activo') AS mg_estado,
          (COALESCE({alias}.mg_estado,'activo') = 'inactivo_venta') AS mg_inactivo_venta,
          {alt_sql}
        """
    return f"""
      'activo' AS mg_estado,
      FALSE AS mg_inactivo_venta,
      {alt_sql}
    """


def _historico_deleted_rows(cur, filters: dict | None = None) -> list[dict]:
    if connection.vendor != "postgresql":
        return []
    filters = filters or {}
    try:
        cur.execute("SELECT to_regclass('audit.change_log')")
        if not cur.fetchone()[0]:
            return []
    except Exception:
        return []

    cur.execute(
        """
        WITH deleted_notes AS (
          SELECT DISTINCT ON (ingreso_id)
                 id,
                 ingreso_id,
                 ts,
                 new_value AS eliminacion_motivo
            FROM audit.change_log
           WHERE table_name = 'ingresos'
             AND column_name = 'correccion_operativa'
             AND ingreso_id IS NOT NULL
           ORDER BY ingreso_id, id DESC
        ),
        ingreso_fields AS (
          SELECT ingreso_id,
                 MAX(old_value) FILTER (WHERE column_name = 'device_id') AS device_id,
                 MAX(old_value) FILTER (WHERE column_name = 'fecha_ingreso') AS fecha_ingreso,
                 MAX(old_value) FILTER (WHERE column_name = 'presupuesto_estado') AS presupuesto_estado,
                 MAX(old_value) FILTER (WHERE column_name = 'ubicacion_id') AS ubicacion_id
            FROM audit.change_log
           WHERE table_name = 'ingresos'
             AND new_value IS NULL
           GROUP BY ingreso_id
        ),
        device_fields AS (
          SELECT ingreso_id,
                 record_id::text AS device_id,
                 MAX(old_value) FILTER (WHERE column_name = 'customer_id') AS customer_id,
                 MAX(old_value) FILTER (WHERE column_name = 'marca_id') AS marca_id,
                 MAX(old_value) FILTER (WHERE column_name = 'model_id') AS model_id,
                 MAX(old_value) FILTER (WHERE column_name = 'numero_serie') AS numero_serie,
                 MAX(old_value) FILTER (WHERE column_name = 'numero_interno') AS numero_interno,
                 MAX(old_value) FILTER (WHERE column_name = 'tipo_equipo') AS tipo_equipo,
                 MAX(old_value) FILTER (WHERE column_name = 'variante') AS variante,
                 MAX(old_value) FILTER (WHERE column_name = 'mg_estado') AS mg_estado,
                 MAX(old_value) FILTER (WHERE column_name = 'mg_venta_numero_alternativo') AS mg_venta_numero_alternativo
            FROM audit.change_log
           WHERE table_name = 'devices'
             AND new_value IS NULL
           GROUP BY ingreso_id, record_id
        )
        SELECT
          dn.ingreso_id AS id,
          'eliminado' AS estado,
          COALESCE(NULLIF(i.presupuesto_estado, ''), 'no_aplica') AS presupuesto_estado,
          CASE WHEN COALESCE(i.fecha_ingreso, '') ~ '^\\d{4}-\\d{2}-\\d{2}'
               THEN i.fecha_ingreso::timestamptz
               ELSE NULL
          END AS fecha_ingreso,
          NULL::timestamptz AS fecha_liberacion,
          NULL::timestamptz AS fecha_entrega,
          CASE WHEN COALESCE(i.ubicacion_id, '') ~ '^\\d+$' THEN i.ubicacion_id::integer ELSE NULL END AS ubicacion_id,
          NULL::timestamptz AS presupuesto_fecha_emision,
          COALESCE(loc.nombre, '') AS ubicacion_nombre,
          CASE WHEN COALESCE(d.customer_id, '') ~ '^\\d+$' THEN d.customer_id::integer ELSE NULL END AS customer_id,
          COALESCE(c.razon_social, '') AS razon_social,
          COALESCE(d.numero_serie, '') AS numero_serie,
          COALESCE(d.numero_interno, '') AS numero_interno,
          COALESCE(NULLIF(d.mg_estado, ''), 'activo') AS mg_estado,
          (COALESCE(NULLIF(d.mg_estado, ''), 'activo') = 'inactivo_venta') AS mg_inactivo_venta,
          COALESCE(d.mg_venta_numero_alternativo, '') AS mg_venta_numero_alternativo,
          COALESCE(b.nombre, '') AS marca,
          COALESCE(m.nombre, '') AS modelo,
          COALESCE(NULLIF(d.tipo_equipo, ''), m.tipo_equipo, '') AS tipo_equipo,
          COALESCE(NULLIF(d.variante, ''), m.variante, '') AS equipo_variante,
          TRUE AS eliminado,
          dn.eliminacion_motivo,
          dn.ts AS eliminado_en,
          CASE WHEN COALESCE(i.device_id, '') ~ '^\\d+$' THEN i.device_id::integer ELSE NULL END AS eliminado_device_id
        FROM deleted_notes dn
        LEFT JOIN ingreso_fields i ON i.ingreso_id = dn.ingreso_id
        LEFT JOIN device_fields d ON d.ingreso_id = dn.ingreso_id AND d.device_id = i.device_id
        LEFT JOIN customers c ON c.id = CASE WHEN COALESCE(d.customer_id, '') ~ '^\\d+$' THEN d.customer_id::integer ELSE NULL END
        LEFT JOIN marcas b ON b.id = CASE WHEN COALESCE(d.marca_id, '') ~ '^\\d+$' THEN d.marca_id::integer ELSE NULL END
        LEFT JOIN models m ON m.id = CASE WHEN COALESCE(d.model_id, '') ~ '^\\d+$' THEN d.model_id::integer ELSE NULL END
        LEFT JOIN locations loc ON loc.id = CASE WHEN COALESCE(i.ubicacion_id, '') ~ '^\\d+$' THEN i.ubicacion_id::integer ELSE NULL END
        WHERE NOT EXISTS (SELECT 1 FROM ingresos t WHERE t.id = dn.ingreso_id)
        ORDER BY dn.ingreso_id DESC
        """
    )
    rows = _fetchall_dicts(cur)
    return [row for row in rows if _historico_deleted_row_matches(row, filters)]


def _historico_deleted_row_matches(row: dict, filters: dict) -> bool:
    import re as _re

    def txt(value) -> str:
        return str(value or "").strip()

    def low(value) -> str:
        return txt(value).lower()

    def contains(value, needle) -> bool:
        needle = txt(needle)
        return not needle or needle.lower() in low(value)

    def date_value(value):
        if value is None:
            return None
        if isinstance(value, datetime):
            try:
                return timezone.localtime(value).date()
            except Exception:
                return value.date()
        if isinstance(value, date):
            return value
        parsed_dt = parse_datetime(txt(value))
        if parsed_dt:
            try:
                return timezone.localtime(parsed_dt).date()
            except Exception:
                return parsed_dt.date()
        return parse_date(txt(value))

    def in_date_range(field_name: str, from_raw: str, to_raw: str) -> bool:
        if not (from_raw or to_raw):
            return True
        current = date_value(row.get(field_name))
        if current is None:
            return False
        if from_raw:
            from_date = parse_date(txt(from_raw))
            if from_date and current < from_date:
                return False
        if to_raw:
            to_date = parse_date(txt(to_raw))
            if to_date and current > to_date:
                return False
        return True

    estado = low(row.get("estado"))
    estados = filters.get("estados") or []
    excluir_estados = filters.get("excluir_estados") or []
    if estados and estado not in estados:
        return False
    if excluir_estados and estado in excluir_estados:
        return False

    ubic = canonical_location_name(filters.get("ubic")) if filters.get("ubic") else ""
    ubic_id_raw = txt(filters.get("ubic_id_raw"))
    row_ubic_canonical = canonical_location_name(row.get("ubicacion_nombre")) if row.get("ubicacion_nombre") else ""
    if ubic and row_ubic_canonical.lower() != ubic.lower():
        return False
    if not ubic and ubic_id_raw.isdigit() and int(row.get("ubicacion_id") or 0) != int(ubic_id_raw):
        return False
    if (
        not ubic
        and not ubic_id_raw.isdigit()
        and filters.get("solo_taller_param_present")
        and filters.get("solo_taller")
        and low(row.get("ubicacion_nombre")) != "taller"
    ):
        return False

    if filters.get("delivered"):
        return False
    if not in_date_range("fecha_ingreso", filters.get("fecha_ingreso_from_raw"), filters.get("fecha_ingreso_to_raw")):
        return False
    if not in_date_range("fecha_liberacion", filters.get("fecha_liberacion_from_raw"), filters.get("fecha_liberacion_to_raw")):
        return False
    if not in_date_range("fecha_entrega", filters.get("fecha_entrega_from_raw"), filters.get("fecha_entrega_to_raw")):
        return False

    searchable = " ".join(
        txt(row.get(key))
        for key in (
            "id",
            "razon_social",
            # El dueño real, por el mismo motivo que en el WHERE del SQL: es el nombre que
            # muestra la columna Cliente.
            "propietario_nombre",
            "marca",
            "modelo",
            "tipo_equipo",
            "equipo_variante",
            "estado",
            "numero_serie",
            "numero_interno",
            "ubicacion_nombre",
            "eliminacion_motivo",
        )
    )
    q_raw = txt(filters.get("q_raw"))
    if q_raw:
        # Criterio único de número interno (`mg_codes`): "MG1678", "MG 1678" y "MG 001678" son
        # el mismo equipo.
        mg_buscado = normalize_mg_code(q_raw)
        m_os = _re.match(r"^(?:OS\s*)?(\d+)$", q_raw, _re.IGNORECASE)
        if mg_buscado:
            if mg_buscado not in (normalize_mg_code(row.get("numero_interno")), normalize_mg_code(row.get("numero_serie"))):
                return False
        elif m_os:
            if int(row.get("id") or 0) != int(m_os.group(1)):
                return False
        elif not contains(searchable, q_raw):
            return False

    os_raw = txt(filters.get("os_raw"))
    if os_raw:
        m_os = _re.match(r"^(?:OS\s*)?(\d+)$", os_raw, _re.IGNORECASE)
        if m_os:
            if int(row.get("id") or 0) != int(m_os.group(1)):
                return False
        elif not contains(str(row.get("id") or ""), os_raw):
            return False

    # El filtro de la columna Cliente mira el dueño además de la razón social, igual que el
    # WHERE del SQL: si no, una OS eliminada de un particular no se encuentra por el nombre
    # que la propia columna muestra.
    cliente_filtro = filters.get("cliente_raw")
    if cliente_filtro and not (
        contains(row.get("razon_social"), cliente_filtro)
        or contains(row.get("propietario_nombre"), cliente_filtro)
    ):
        return False

    checks = [
        ("tipo_equipo", "tipo_equipo_raw"),
        ("marca", "marca_raw"),
        ("modelo", "modelo_raw"),
        ("equipo_variante", "variante_raw"),
        ("estado", "estado_q_raw"),
        ("numero_serie", "numero_serie_raw"),
        ("numero_interno", "numero_interno_raw"),
    ]
    for field_name, filter_name in checks:
        if not contains(row.get(field_name), filters.get(filter_name)):
            return False
    return True


def _historico_python_sort_value(row: dict, field: str):
    accessors = {
        "os": lambda item: int(item.get("id") or 0),
        # Mismo criterio que el ORDER BY del SQL: el nombre que muestra la columna.
        "cliente": _ingreso_customer_label,
        "tipo": lambda item: item.get("tipo_equipo"),
        "marca": lambda item: item.get("marca"),
        "modelo": lambda item: item.get("modelo"),
        "variante": lambda item: item.get("equipo_variante"),
        "estado": lambda item: item.get("estado"),
        "numero_serie": lambda item: item.get("numero_serie"),
        "numero_interno": lambda item: item.get("numero_interno"),
        "fecha_ingreso": lambda item: item.get("fecha_ingreso"),
        "fecha_liberacion": lambda item: item.get("fecha_liberacion"),
        "fecha_entrega": lambda item: item.get("fecha_entrega"),
    }
    value = accessors.get(field, accessors["os"])(row)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip().casefold()
    return value


def _merge_historico_rows(rows: list[dict], deleted_rows: list[dict], sort_raw: str = "") -> list[dict]:
    if not deleted_rows:
        return rows
    seen = {int(row.get("id") or 0) for row in rows if str(row.get("id") or "").isdigit()}
    merged = list(rows)
    for row in deleted_rows:
        try:
            row_id = int(row.get("id") or 0)
        except Exception:
            row_id = 0
        if row_id and row_id not in seen:
            merged.append(row)
            seen.add(row_id)
    raw = str(sort_raw or "").strip()
    direction_desc = raw.startswith("-")
    field = raw[1:] if direction_desc else raw
    valid_fields = {
        "os", "cliente", "tipo", "marca", "modelo", "variante", "estado",
        "numero_serie", "numero_interno", "fecha_ingreso", "fecha_liberacion", "fecha_entrega",
    }
    if field not in valid_fields:
        return sorted(merged, key=lambda item: int(item.get("id") or 0), reverse=True)
    non_empty = []
    empty = []
    for item in merged:
        value = _historico_python_sort_value(item, field)
        (empty if value in (None, "") else non_empty).append(item)
    non_empty.sort(key=lambda item: int(item.get("id") or 0), reverse=direction_desc)
    non_empty.sort(
        key=lambda item: _historico_python_sort_value(item, field),
        reverse=direction_desc,
    )
    empty.sort(key=lambda item: int(item.get("id") or 0), reverse=direction_desc)
    return non_empty + empty


def _has_historical_corrections_schema() -> bool:
    return (
        _table_exists("ingreso_historical_corrections")
        and _has_table_column("ingreso_historical_corrections", "accion")
        and _has_table_column("ingreso_historical_corrections", "fecha_efectiva")
        and _has_table_column("ingreso_historical_corrections", "motivo")
        and _has_table_column("ingreso_historical_corrections", "payload")
    )


def _ingresos_has_permite_reparacion_col() -> bool:
    try:
        if connection.vendor == "postgresql":
            row = q(
                """
                SELECT 1
                  FROM information_schema.columns
                 WHERE table_name='ingresos'
                   AND column_name='permite_reparacion'
                   AND table_schema = ANY(current_schemas(true))
                 LIMIT 1
                """,
                one=True,
            )
            return bool(row)
        row = q(
            """
            SELECT 1
              FROM information_schema.columns
             WHERE table_name='ingresos'
               AND column_name='permite_reparacion'
             LIMIT 1
            """,
            one=True,
        )
        return bool(row)
    except Exception:
        return False


def _is_cotizacion_bloqueada_para_reparar(motivo, permite_reparacion) -> bool:
    return _motivo_is_cotizacion_equipo(motivo) and not bool(permite_reparacion)


def _motivo_is_revision_tecnica(value: str) -> bool:
    key = _norm_txt(value or "")
    return key == _norm_txt("Revisión Técnica")


def _mg_owner_customer_id():
    row = q(
        "SELECT id FROM customers WHERE LOWER(razon_social) LIKE %s ORDER BY id ASC LIMIT 1",
        ["%mg%bio%"],
        one=True,
    )
    return row.get("id") if row else None


def _is_mg_owner_customer(customer_id, mg_owner_id=None) -> bool:
    try:
        if not customer_id:
            return False
        owner_id = mg_owner_id if mg_owner_id is not None else _mg_owner_customer_id()
        if not owner_id:
            return False
        return int(customer_id) == int(owner_id)
    except Exception:
        return False


def _fetch_ingreso_device_mg_context(ingreso_id):
    n_de_control_sql = (
        "COALESCE(d.n_de_control,'') AS n_de_control,"
        if _has_table_column("devices", "n_de_control")
        else "'' AS n_de_control,"
    )
    mg_estado_sql = (
        "COALESCE(d.mg_estado,'activo') AS mg_estado,"
        if _has_mg_schema()
        else "'activo' AS mg_estado,"
    )
    row = q(
        f"""
        SELECT
          i.estado,
          d.customer_id,
          COALESCE(d.numero_interno,'') AS numero_interno,
          COALESCE(d.numero_serie,'') AS numero_serie,
          {n_de_control_sql}
          {mg_estado_sql}
          COALESCE(d.alquilado,false) AS alquilado
          FROM ingresos i
          LEFT JOIN devices d ON d.id = i.device_id
         WHERE i.id=%s
        """,
        [ingreso_id],
        one=True,
    )
    if row:
        row.update(resolve_mg_flags(row, _mg_owner_customer_id()))
    return row


def _parse_rejected_budget_charge_or_raise(raw_value, *, required: bool) -> object:
    if raw_value is None or (isinstance(raw_value, str) and raw_value.strip() == ""):
        if required:
            raise ValidationError("presupuesto_rechazado_cobro_neto requerido")
        return None
    try:
        cobro_neto = money(raw_value)
    except (TypeError, ValueError):
        raise ValidationError("presupuesto_rechazado_cobro_neto debe ser numérico")
    if cobro_neto < 0:
        raise ValidationError("presupuesto_rechazado_cobro_neto no puede ser negativo")
    return cobro_neto


def _send_mail_with_fallback(subject, body, recipients):
    debug = {}
    try:
        debug.update({
            "backend": getattr(settings, "EMAIL_BACKEND", None),
            "host": getattr(settings, "EMAIL_HOST", None),
            "port": getattr(settings, "EMAIL_PORT", None),
            "use_tls": getattr(settings, "EMAIL_USE_TLS", None),
            "use_ssl": getattr(settings, "EMAIL_USE_SSL", None),
            "from": getattr(settings, "DEFAULT_FROM_EMAIL", None),
            "recipients": list(recipients or []),
        })
    except Exception:
        pass
    if not recipients:
        logger.warning("email no recipients configured")
        return False, debug
    try:
        sent = send_mail(subject, body, getattr(settings, "DEFAULT_FROM_EMAIL", None), recipients, fail_silently=False)
        ok = bool(sent and sent > 0)
        return ok, debug
    except Exception as e:
        try:
            debug["error"] = str(e)
            debug["exception"] = e.__class__.__name__
        except Exception:
            pass
        try:
            port_cfg = int(getattr(settings, "EMAIL_PORT", 0) or 0)
        except Exception:
            port_cfg = 0
        if port_cfg == 587:
            try:
                conn = get_connection(
                    backend=getattr(settings, "EMAIL_BACKEND", "django.core.mail.backends.smtp.EmailBackend"),
                    host=getattr(settings, "EMAIL_HOST", None),
                    port=465,
                    username=getattr(settings, "EMAIL_HOST_USER", None),
                    password=getattr(settings, "EMAIL_HOST_PASSWORD", None),
                    use_tls=False,
                    use_ssl=True,
                    fail_silently=False,
                )
                msg = EmailMessage(subject, body, getattr(settings, "DEFAULT_FROM_EMAIL", None), recipients, connection=conn)
                sent2 = msg.send()
                ok2 = bool(sent2 and sent2 > 0)
                debug["fallback"] = {"mode": "ssl_465", "sent": ok2}
                return ok2, debug
            except Exception as e2:
                try:
                    debug.setdefault("fallback", {})["error"] = str(e2)
                    debug.setdefault("fallback", {})["exception"] = e2.__class__.__name__
                except Exception:
                    pass
        return False, debug


def _dash_location_id():
    try:
        ensure_default_locations()
    except Exception:
        pass
    try:
        row = q("SELECT id FROM locations WHERE nombre='-' LIMIT 1", [], one=True)
        return row and row.get("id")
    except Exception:
        return None


def _taller_location_id():
    try:
        ensure_default_locations()
    except Exception:
        pass
    try:
        row = q(
            "SELECT id FROM locations WHERE LOWER(nombre)=LOWER(%s) LIMIT 1",
            ["Taller"],
            one=True,
        )
        return row and row.get("id")
    except Exception:
        return None


def _pending_baja_request(ingreso_id: int):
    if not _table_exists("ingreso_baja_requests"):
        return None
    try:
        return q(
            """
            SELECT r.id,
                   r.usuario_id AS uid,
                   COALESCE(u.nombre, '') AS nombre,
                   COALESCE(r.motivo, '') AS motivo,
                   r.created_at
              FROM ingreso_baja_requests r
              LEFT JOIN users u ON u.id = r.usuario_id
             WHERE r.ingreso_id = %s
               AND r.accepted_at IS NULL
               AND r.canceled_at IS NULL
             ORDER BY r.created_at DESC, r.id DESC
             LIMIT 1
            """,
            [ingreso_id],
            one=True,
        )
    except Exception:
        try:
            transaction.set_rollback(False)
        except Exception:
            pass
        return None


def _table_exists(table_name: str) -> bool:
    try:
        row = q(
            """
            SELECT 1
              FROM information_schema.tables
             WHERE table_schema = ANY(current_schemas(true))
               AND table_name = %s
             LIMIT 1
            """,
            [table_name],
            one=True,
        )
        return bool(row)
    except Exception:
        return False


def _service_release_remito_salida_for_ingreso(ingreso_id: int) -> str:
    if not (_table_exists("delivery_orders") and _table_exists("bejerman_remito_groups")):
        return ""
    delivery_columns = {"id", "ingreso_id", "delivery_type"}
    group_columns = {
        "id",
        "status",
        "comprobante_tipo",
        "comprobante_letra",
        "comprobante_pto_venta",
        "comprobante_numero",
        "remito_number",
        "response_summary",
        "created_at",
        "generated_at",
    }
    if not all(_has_table_column("delivery_orders", column) for column in delivery_columns):
        return ""
    if not all(_has_table_column("bejerman_remito_groups", column) for column in group_columns):
        return ""

    ingreso_match_conditions = ["o.ingreso_id = target.ingreso_id"]
    has_item_ingreso = _table_exists("delivery_order_items") and all(
        _has_table_column("delivery_order_items", column) for column in ("id", "order_id", "ingreso_id")
    )
    if has_item_ingreso:
        ingreso_match_conditions.append(
            """
            EXISTS (
              SELECT 1
                FROM delivery_order_items item
               WHERE item.order_id = o.id
                 AND item.ingreso_id = target.ingreso_id
            )
            """
        )
    has_partida_ingreso = has_item_ingreso and _table_exists("delivery_order_item_partidas") and all(
        _has_table_column("delivery_order_item_partidas", column)
        for column in ("order_item_id", "ingreso_id")
    )
    if has_partida_ingreso:
        ingreso_match_conditions.append(
            """
            EXISTS (
              SELECT 1
                FROM delivery_order_item_partidas partida
                JOIN delivery_order_items item ON item.id = partida.order_item_id
               WHERE item.order_id = o.id
                 AND partida.ingreso_id = target.ingreso_id
            )
            """
        )
    ingreso_match_sql = "\nOR ".join(ingreso_match_conditions)

    link_columns = {"id", "order_id", "remito_group_id", "batch_id", "sequence", "created_at"}
    has_canonical_links = _table_exists("delivery_order_remito_links") and all(
        _has_table_column("delivery_order_remito_links", column) for column in link_columns
    )
    if has_canonical_links:
        try:
            row = q(
                f"""
                WITH target AS (
                  SELECT %s::BIGINT AS ingreso_id
                ),
                remitos AS (
                  SELECT DISTINCT ON (g.id)
                         g.id,
                         NULLIF(BTRIM(COALESCE(
                           NULLIF(g.remito_number, ''),
                           NULLIF(CONCAT_WS(' ', NULLIF(g.comprobante_tipo, ''), NULLIF(g.comprobante_letra, ''), NULLIF(g.comprobante_pto_venta, ''), NULLIF(g.comprobante_numero, '')), '')
                         )), '') AS remito_salida,
                         l.batch_id,
                         l.sequence,
                         l.created_at AS linked_at,
                         MIN(l.created_at) OVER (PARTITION BY o.id, l.batch_id) AS batch_created_at,
                         COALESCE(g.generated_at, g.created_at, l.created_at) AS document_at
                    FROM delivery_orders o
                    JOIN delivery_order_remito_links l ON l.order_id = o.id
                    JOIN bejerman_remito_groups g ON g.id = l.remito_group_id
                    CROSS JOIN target
                   WHERE ({ingreso_match_sql})
                     AND o.delivery_type = 'service_release'
                     AND g.status IN ('generated', 'reconciled')
                     AND (
                       NOT EXISTS (
                         SELECT 1 FROM delivery_order_remito_item_allocations AS allocation
                          WHERE allocation.link_id = l.id
                       )
                       OR EXISTS (
                         SELECT 1
                           FROM delivery_order_remito_item_allocations AS allocation
                           JOIN delivery_order_items AS assigned_item ON assigned_item.id = allocation.order_item_id
                           LEFT JOIN delivery_order_item_partidas AS assigned_partida ON assigned_partida.id = allocation.order_item_partida_id
                          WHERE allocation.link_id = l.id
                            AND COALESCE(assigned_partida.ingreso_id, assigned_item.ingreso_id, o.ingreso_id) = target.ingreso_id
                       )
                     )
                     AND UPPER(COALESCE(
                           NULLIF(g.comprobante_tipo, ''),
                           NULLIF(g.response_summary -> 'profile' ->> 'type', ''),
                           NULLIF(g.response_summary ->> 'comprobanteTipo', '')
                         )) = 'RSS'
                   ORDER BY g.id, l.created_at, l.sequence, l.id
                )
                SELECT STRING_AGG(
                         remito_salida,
                         ' / ' ORDER BY batch_created_at, sequence, linked_at, document_at, id
                       ) AS remito_salida
                  FROM remitos
                 WHERE remito_salida IS NOT NULL
                """,
                [ingreso_id],
                one=True,
            )
            return (row or {}).get("remito_salida") or ""
        except Exception:
            return ""

    if not _has_table_column("delivery_orders", "bejerman_remito_group_id"):
        return ""
    try:
        row = q(
            f"""
            WITH target AS (
              SELECT %s::BIGINT AS ingreso_id
            ),
            remitos AS (
              SELECT DISTINCT ON (g.id)
                     g.id,
                     NULLIF(BTRIM(COALESCE(
                       NULLIF(g.remito_number, ''),
                       NULLIF(CONCAT_WS(' ', NULLIF(g.comprobante_tipo, ''), NULLIF(g.comprobante_letra, ''), NULLIF(g.comprobante_pto_venta, ''), NULLIF(g.comprobante_numero, '')), '')
                     )), '') AS remito_salida,
                     COALESCE(g.generated_at, g.created_at) AS document_at
                FROM delivery_orders o
                JOIN bejerman_remito_groups g ON g.id = o.bejerman_remito_group_id
                CROSS JOIN target
               WHERE ({ingreso_match_sql})
                 AND o.delivery_type = 'service_release'
                 AND g.status IN ('generated', 'reconciled')
                 AND UPPER(COALESCE(
                       NULLIF(g.comprobante_tipo, ''),
                       NULLIF(g.response_summary -> 'profile' ->> 'type', ''),
                       NULLIF(g.response_summary ->> 'comprobanteTipo', '')
                     )) = 'RSS'
               ORDER BY g.id, COALESCE(g.generated_at, g.created_at), g.created_at
            )
            SELECT STRING_AGG(remito_salida, ' / ' ORDER BY document_at, id) AS remito_salida
              FROM remitos
             WHERE remito_salida IS NOT NULL
            """,
            [ingreso_id],
            one=True,
        )
        return (row or {}).get("remito_salida") or ""
    except Exception:
        return ""


def _estado_actual_txt(ingreso_id: int) -> str:
    """Estado actual de la OS en minúsculas, o '' si no se pudo leer."""
    try:
        row = q("SELECT estado::text AS estado FROM ingresos WHERE id=%s", [ingreso_id], one=True) or {}
        return (row.get("estado") or "").strip().lower()
    except Exception:
        return ""


def _presupuesto_estado_actual_txt(ingreso_id: int) -> str:
    """Estado de presupuesto que quedó grabado, para no responder uno de fantasía.

    Los cierres sin presupuesto respetan un presupuesto ya emitido o rechazado, así
    que lo que quedó no siempre es 'no_aplica'.
    """
    try:
        row = q(
            "SELECT presupuesto_estado FROM ingresos WHERE id=%s",
            [ingreso_id],
            one=True,
        ) or {}
        return str(row.get("presupuesto_estado") or "").strip().lower()
    except Exception:
        return ""


# Valores válidos del enum `ticket_state`. `de_estado` es informativo, así que cualquier
# cosa fuera de esta lista (p. ej. el 'asignado' legacy, que nunca entró al enum) se
# guarda como NULL en vez de hacer fallar el INSERT.
TICKET_STATES = frozenset(
    (
        "ingresado", "diagnosticado", "presupuestado", "reparar", "controlado_sin_defecto",
        "no_se_repara", "reparado", "entregado", "baja", "derivado", "liberado", "alquilado",
        "vendido_pendiente_entrega", "vendido_entregado",
    )
)


def registrar_ingreso_event(
    ingreso_id: int,
    a_estado: str,
    *,
    de_estado: str | None = None,
    usuario_id=None,
    comentario: str = "",
) -> None:
    """Deja la traza del cambio de estado en `ingreso_events`.

    Los hitos del panel de la hoja de servicio y las métricas de tiempos se calculan
    desde esta tabla, así que todo endpoint que mueva `ingresos.estado` tiene que
    pasar por acá. La escritura va en su propio savepoint y nunca propaga el error:
    perder la traza es feo, pero no puede tumbar la operación que sí se hizo.
    """
    if not _table_exists("ingreso_events"):
        return
    de_estado_txt = (de_estado or "").strip().lower()
    if de_estado_txt not in TICKET_STATES:
        de_estado_txt = ""
    try:
        with transaction.atomic():
            exec_void(
                """
                INSERT INTO ingreso_events (ticket_id, de_estado, a_estado, usuario_id, comentario)
                VALUES (%s, NULLIF(%s,'')::ticket_state, %s::ticket_state, %s, %s)
                """,
                [ingreso_id, de_estado_txt, a_estado, usuario_id, comentario or None],
            )
        return
    except Exception:
        pass
    # Fallback para esquemas donde las columnas son TEXT (sin el enum `ticket_state`).
    try:
        with transaction.atomic():
            exec_void(
                """
                INSERT INTO ingreso_events (ticket_id, a_estado, usuario_id, comentario)
                VALUES (%s, %s, %s, %s)
                """,
                [ingreso_id, a_estado, usuario_id, comentario or None],
            )
    except Exception:
        logger.exception(
            "No se pudo registrar el evento de estado",
            extra={"ingreso_id": ingreso_id, "a_estado": a_estado},
        )


def _ingreso_event_date_hitos(ingreso_id: int) -> dict:
    if not _table_exists("ingreso_events"):
        return {}
    where_parts = []
    where_params = []
    if _has_table_column("ingreso_events", "ingreso_id"):
        where_parts.append("ingreso_id = %s")
        where_params.append(ingreso_id)
    if _has_table_column("ingreso_events", "ticket_id"):
        where_parts.append("ticket_id = %s")
        where_params.append(ingreso_id)
    if not where_parts:
        return {}
    try:
        # La etapa de reparación cierra por cualquiera de los tres estados finales del
        # técnico, igual que en la cola de Reparados: quedarse sólo con 'reparado' dejaba
        # sin fecha a todo equipo controlado sin defecto o que no se repara.
        # El placeholder de la lista de estados va primero: está en el SELECT, antes del WHERE.
        return q(
            f"""
            SELECT
              MIN(CASE WHEN CAST(a_estado AS TEXT) = 'diagnosticado' THEN ts END) AS fecha_diagnosticado,
              MIN(CASE WHEN CAST(a_estado AS TEXT) = ANY(%s) THEN ts END) AS fecha_reparado,
              MAX(CASE WHEN CAST(a_estado AS TEXT) = 'liberado' THEN ts END) AS fecha_liberacion
              FROM ingreso_events
             WHERE {" OR ".join(where_parts)}
            """,
            [list(TECH_FINAL_REPAIR_STATES), *where_params],
            one=True,
        ) or {}
    except Exception:
        try:
            transaction.set_rollback(False)
        except Exception:
            pass
        return {}


def _ingreso_quote_date_hitos(ingreso_id: int) -> dict:
    """Fechas del presupuesto de la OS: cuándo se emitió, se aprobó y se rechazó.

    Se toma el valor más reciente **de cada fecha** entre las versiones no anuladas, en
    vez de las fechas de la última versión. El motivo: versionar un presupuesto clona la
    fila con `fecha_emitido`, `fecha_aprobado` y `fecha_rechazado` en NULL
    (`_clone_quote_version`), así que mirar sólo la última dejaba el hito "Presupuestado"
    en blanco en toda OS re-versionada, como si nunca se hubiera presupuestado.

    Si todas están anuladas se cae al mismo cálculo sobre todas las versiones, para no
    perder el dato.
    """
    if not (
        _has_table("quotes")
        and _has_table_column("quotes", "ingreso_id")
        and _has_table_column("quotes", "fecha_emitido")
    ):
        return {}
    estado_ok = _has_table_column("quotes", "estado")
    hitos_sql = """
        SELECT MAX(fecha_emitido)   AS fecha_emitido,
               MAX(fecha_aprobado)  AS fecha_aprobado,
               MAX(fecha_rechazado) AS fecha_rechazado
          FROM quotes
         WHERE ingreso_id=%s
    """
    try:
        row = None
        if estado_ok:
            row = q(
                f"{hitos_sql} AND LOWER(COALESCE(estado::text,'')) <> 'anulado'",
                [ingreso_id],
                one=True,
            )
        # El agregado siempre devuelve una fila: si las versiones vivas no tienen ninguna
        # fecha (todas anuladas, o ninguna emitida todavía) se reintenta sobre todas.
        if not row or not any(
            row.get(k) for k in ("fecha_emitido", "fecha_aprobado", "fecha_rechazado")
        ):
            row = q(hitos_sql, [ingreso_id], one=True)
        row = row or {}
        return {
            "fecha_presupuestado": row.get("fecha_emitido"),
            "fecha_presupuesto_aprobado": row.get("fecha_aprobado"),
            "fecha_presupuesto_rechazado": row.get("fecha_rechazado"),
        }
    except Exception:
        try:
            transaction.set_rollback(False)
        except Exception:
            pass
        return {}


def _parse_bool_or_default(raw, default=True) -> bool:
    if raw is None:
        return bool(default)
    if isinstance(raw, bool):
        return raw
    val = str(raw).strip().lower()
    if val in ("1", "true", "yes", "y", "si", "sí", "on"):
        return True
    if val in ("0", "false", "no", "n", "off"):
        return False
    return bool(default)


def _parse_historical_effective_datetime(raw):
    dt = _parse_datetime_or_date(raw)
    if not dt:
        return None
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def _historical_rental_customer(customer_id):
    try:
        customer_id_int = int(customer_id)
    except (TypeError, ValueError):
        return None
    if customer_id_int <= 0:
        return None
    alias_sql = (
        "COALESCE(c.alias_interno, '') AS alias_interno"
        if _has_table_column("customers", "alias_interno")
        else "'' AS alias_interno"
    )
    return q(
        f"""
        SELECT c.id,
               COALESCE(c.razon_social, '') AS razon_social,
               COALESCE(c.cod_empresa, '') AS cod_empresa,
               {alias_sql}
          FROM customers c
         WHERE c.id=%s
        """,
        [customer_id_int],
        one=True,
    )


def _sync_historical_device_snapshot(
    *,
    ingreso_id: int,
    device_id: int,
    accion: str,
    alquiler_a: str = "",
    dash_id=None,
    taller_id=None,
) -> bool:
    """Mantiene el snapshot del equipo solo si la OS corregida es la más reciente."""
    latest = q(
        """
        SELECT id
          FROM ingresos
         WHERE device_id=%s
         ORDER BY COALESCE(fecha_ingreso, fecha_creacion) DESC, id DESC
         LIMIT 1
        """,
        [device_id],
        one=True,
    ) or {}
    if int(latest.get("id") or 0) != int(ingreso_id):
        return False

    sets = []
    params = []
    has_alquilado = _has_table_column("devices", "alquilado")
    has_alquiler_a = _has_table_column("devices", "alquiler_a")
    has_ubicacion = _has_table_column("devices", "ubicacion_id")

    if accion == "alta_alquiler":
        if has_alquilado:
            sets.append("alquilado=TRUE")
        if has_alquiler_a:
            sets.append("alquiler_a=NULLIF(%s,'')")
            params.append(alquiler_a)
    elif accion == "baja_alquiler":
        if has_alquilado:
            sets.append("alquilado=FALSE")
        if has_alquiler_a:
            sets.append("alquiler_a=NULL")

    target_location = taller_id if accion in ("baja_alquiler", "alta_ingreso") else dash_id
    if has_ubicacion and target_location:
        sets.append("ubicacion_id=COALESCE(%s, ubicacion_id)")
        params.append(target_location)

    if not sets:
        return False
    params.append(device_id)
    exec_void(f"UPDATE devices SET {', '.join(sets)} WHERE id=%s", params)
    return True


def _historical_bejerman_reference_effect(accion: str, document_number: str = "") -> dict:
    labels = {
        "entrega": "remito de salida",
        "alta_alquiler": "RTA de alquiler",
        "baja_alquiler": "RDA de devolución",
    }
    document_label = labels.get(accion, "comprobante")
    document_number = str(document_number or "").strip()
    if document_number:
        detail = (
            f"Se registró la referencia {document_number}. "
            f"No se emitió un nuevo {document_label} en Bejerman por tratarse de una corrección histórica."
        )
        status = "reference_registered"
    else:
        detail = (
            f"No se emitió un nuevo {document_label} en Bejerman. "
            "La corrección quedó asentada solo en NEXORA y requiere una referencia histórica para poder rastrearse."
        )
        status = "reference_missing"
    return {
        "status": status,
        "document_number": document_number or None,
        "detail": detail,
    }


def _notify_estado_patrimonial(request, ingreso_id: int, evento: str, fecha_evento=None):
    """Aviso patrimonial (baja / alta / alta MG) con el resultado real del envío."""
    return notify_patrimonial_email(
        request,
        evento=evento,
        ingreso_id=ingreso_id,
        fecha_evento=fecha_evento,
    )


class MisPendientesView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        require_roles(request, ["jefe", "tecnico", "jefe_veedor"])
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT t.id,
                       t.estado,
                       t.presupuesto_estado,
                       t.motivo,
                       c.razon_social,
                       {owner_list_sql}
                       d.numero_serie,
                       COALESCE(d.numero_interno,'') AS numero_interno,
                       {mg_list_sql}
                       COALESCE(b.nombre,'') AS marca,
                       COALESCE(m.nombre,'') AS modelo,
                       COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                       COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                       t.asignado_a,
                       COALESCE(u.nombre,'') AS tecnico,
                       COALESCE(u.nombre,'') AS tecnico_nombre,
                       COALESCE(u.nombre,'') AS asignado_a_nombre,
                       t.fecha_ingreso,
                       CASE WHEN ed.estado = 'devuelto' THEN true ELSE false END AS derivado_devuelto
                FROM ingresos t
                JOIN devices d ON d.id=t.device_id
                JOIN customers c ON c.id=d.customer_id
                LEFT JOIN marcas b ON b.id=d.marca_id
                LEFT JOIN models m ON m.id=d.model_id
                LEFT JOIN users u ON u.id=t.asignado_a
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                LEFT JOIN (
                  SELECT e.*, ROW_NUMBER() OVER (
                    PARTITION BY e.ingreso_id ORDER BY e.fecha_deriv DESC, e.id DESC
                  ) AS rn
                  FROM equipos_derivados e
                ) ed ON ed.ingreso_id = t.id AND ed.rn = 1
                WHERE t.asignado_a = %s
                  AND LOWER(loc.nombre) = LOWER(%s)
                  AND t.estado NOT IN ('entregado','liberado','alquilado','baja','no_se_repara','vendido_pendiente_entrega','vendido_entregado')
                ORDER BY
                   (CASE WHEN ed.estado = 'devuelto' THEN 1 ELSE 0 END) DESC,
                   (t.motivo = 'urgente control') DESC,
                   t.fecha_ingreso ASC;
                """,
                [
                    getattr(getattr(request, "user", None), "id", None)
                    or getattr(request, "user_id", None),
                    "taller",
                ],
            )
            data = _fetchall_dicts(cur)
        attach_active_derivations(data)
        return Response(IngresoListItemSerializer(data, many=True).data)


class PendientesPresupuestoView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        from .. import technical_interventions as technical
        has_technical = technical.available()
        technical_extra = """ OR EXISTS (
          SELECT 1 FROM technical_interventions ti WHERE ti.ingreso_id=t.id
          AND ti.scope_revision IS DISTINCT FROM q.technical_revision
        )""" if has_technical else ""
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT t.id, t.estado, t.presupuesto_estado,
                       c.razon_social,
                       {owner_list_sql}
                       d.numero_serie,
                       COALESCE(d.numero_interno,'') AS numero_interno,
                       {mg_list_sql}
                       COALESCE(b.nombre,'') AS marca,
                       COALESCE(m.nombre,'') AS modelo,
                       COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                       COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                       t.motivo,
                       t.fecha_ingreso,
                       t.fecha_servicio,
                       q.fecha_emitido AS presupuesto_fecha_emision,
                       t.ubicacion_id,
                       COALESCE(loc.nombre,'') AS ubicacion_nombre
                FROM ingresos t
                JOIN devices d ON d.id=t.device_id
                JOIN customers c ON c.id=d.customer_id
                LEFT JOIN marcas b ON b.id=d.marca_id
                LEFT JOIN models m ON m.id=d.model_id
                LEFT JOIN quotes q ON q.id = (
                  SELECT q2.id FROM quotes q2
                  WHERE q2.ingreso_id = t.id
                  ORDER BY COALESCE(q2.version_num, 1) DESC, q2.id DESC
                  LIMIT 1
                )
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                WHERE (COALESCE(t.presupuesto_estado, 'pendiente') = 'pendiente' {technical_extra})
                  -- Al volver a alquilar un equipo propio, la OS que documenta la
                  -- reparación puede pasar a estado "alquilado". Eso no resuelve el
                  -- presupuesto por los daños del alquiler anterior: debe seguir en
                  -- esta cola hasta que Comercial lo resuelva. Lo mismo si el equipo
                  -- se da de baja. Las bajas que ya existían al cambiar la regla
                  -- (15/09/2026) se pasaron a "no aplica" por script en sql/ops/.
                  AND t.estado NOT IN ('ingresado','entregado','liberado','no_se_repara','vendido_pendiente_entrega','vendido_entregado')
                  AND (
                    t.estado <> 'alquilado'
                    OR (
                      {sql_is_mg_code('d.numero_interno')}
                      AND LOWER(BTRIM(COALESCE(d.mg_estado,'activo'))) <> 'inactivo_venta'
                    )
                    OR LOWER(BTRIM(COALESCE(t.motivo::text,''))) IN ('baja alquiler','reparación alquiler','reparacion alquiler')
                  )
                ORDER BY COALESCE(t.fecha_servicio, t.fecha_ingreso) ASC;
                """
            )
            data = _fetchall_dicts(cur)
        statuses = technical.queue_status([row['id'] for row in data]) if has_technical else {}
        for row in data:
            row['estado_revision_tecnica'] = statuses.get(row['id'], 'historico')
        attach_active_derivations(data)
        return Response(IngresoListItemSerializer(data, many=True).data)


class PresupuestadosView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT
                  t.id,
                  t.estado,
                  CASE
                    WHEN t.presupuesto_estado IS NOT NULL THEN t.presupuesto_estado::text
                    WHEN q.estado::text IN ('emitido','enviado','presupuestado') THEN 'presupuestado'
                    ELSE q.estado::text
                  END AS presupuesto_estado,
                  c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                  t.motivo,
                  t.fecha_ingreso,
                  q.id AS presupuesto_id,
                  q.id AS presupuesto_numero,
                  -- Monto mostrado: Subtotal del presupuesto (mano de obra + repuestos), sin IVA
                  COALESCE((
                    SELECT ROUND(SUM(qi.qty * qi.precio_u), 2)
                    FROM quote_items qi
                    WHERE qi.quote_id = q.id
                  ), 0) AS presupuesto_monto,
                  COALESCE(q.moneda, 'ARS') AS presupuesto_moneda,
                  q.fecha_emitido AS presupuesto_fecha_emision,
                  NULL AS presupuesto_fecha_envio,
                  t.ubicacion_id,
                  COALESCE(loc.nombre,'') AS ubicacion_nombre
                FROM ingresos t
                JOIN devices d ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN quotes q ON q.id = (
                  SELECT q2.id FROM quotes q2
                  WHERE q2.ingreso_id = t.id
                  ORDER BY COALESCE(q2.version_num, 1) DESC, q2.id DESC
                  LIMIT 1
                )
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                WHERE (
                        q.estado::text IN ('emitido','enviado','presupuestado')
                        OR t.presupuesto_estado = 'presupuestado'
                      )
                  -- `no_se_repara` se muestra: un presupuesto emitido frena la salida aunque el
                  -- equipo no se repare (`release_gate`), y si no figura acá nadie lo resuelve.
                  AND t.estado NOT IN ('entregado','liberado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                ORDER BY COALESCE(q.fecha_emitido, t.fecha_ingreso) ASC;
                """
            )
            data = _fetchall_dicts(cur)
        attach_active_derivations(data)
        return Response(data)


class PresupuestadosExportView(APIView):
    """
    Exporta a Excel (.xlsx) filas de 'presupuestados' dadas por sus IDs.

    Parametros query:
      - ids: lista separada por comas de IDs de ingresos. Ej: ?ids=10,11,15

    Columnas:
      OS, Cliente, Equipo, N/S, Monto sin IVA
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        ids_raw = (request.GET.get("ids") or "").strip()
        if not ids_raw:
            return Response({"detail": "Parámetro 'ids' requerido"}, status=400)

        try:
            ids = [int(x) for x in ids_raw.split(",") if x.strip()]
        except Exception:
            return Response({"detail": "Parámetro 'ids' inválido"}, status=400)

        # Evitar excesos accidentales
        if len(ids) > 1000:
            return Response({"detail": "Demasiados IDs (maximo 1000)"}, status=400)

        # Traer datos necesarios. Reutilizamos estructura del listado 'presupuestados'.
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT
                  t.id,
                  c.razon_social AS cliente,
                  {owner_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE((
                    SELECT ROUND(SUM(qi.qty * qi.precio_u), 2)
                      FROM quote_items qi
                     WHERE qi.quote_id = q.id
                  ), 0) AS subtotal_sin_iva
                FROM ingresos t
                JOIN devices d   ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN quotes q ON q.id = (
                  SELECT q2.id FROM quotes q2
                  WHERE q2.ingreso_id = t.id
                  ORDER BY COALESCE(q2.version_num, 1) DESC, q2.id DESC
                  LIMIT 1
                )
               WHERE t.id = ANY(%s)
               ORDER BY COALESCE(q.fecha_emitido, t.fecha_ingreso, NOW()) ASC
                """,
                [ids],
            )
            rows = _fetchall_dicts(cur)

        # Construir Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Presupuestados"

        headers = [
            "OS",
            "Cliente",
            "Equipo",
            "N/S",
            "Monto sin IVA",
        ]
        ws.append(headers)

        # formato de equipo y N/S via helpers _equipolabel_row y _ns_label
        currency_format = '"$" #,##0.00'

        for r in rows:
            os_txt = os_label(r.get("id"))
            equipo = _equipolabel_row(r)
            ns_val = _ns_label(r)
            monto = r.get("subtotal_sin_iva")

            ws.append([
                os_txt,
                _ingreso_customer_label(r) or "-",
                equipo,
                ns_val,
                float(monto) if (monto is not None) else None,
            ])
            ws.cell(row=ws.max_row, column=5).number_format = currency_format

        # Ajuste simple de ancho de columnas (opcional)
        try:
            widths = [10, 40, 40, 20, 18]
            for idx, w in enumerate(widths, start=1):
                col = ws.column_dimensions[chr(64 + idx)]
                col.width = w
        except Exception:
            pass

        # Serializar a binario y responder
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)

        fname = f"presupuestados_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
        resp = HttpResponse(
            bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f"attachment; filename=\"{fname}\""
        return resp


class MarcarReparadoView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def post(self, request, ingreso_id: int):
        require_roles(request, ["tecnico", "jefe", "jefe_veedor"])
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        permite_reparacion_sql = (
            "COALESCE(permite_reparacion, TRUE) AS permite_reparacion"
            if has_permite_reparacion
            else "TRUE AS permite_reparacion"
        )
        # Solo el técnico asignado puede marcar como reparado (salvo jefes)
        try:
            if _rol(request) in ("tecnico", "jefe_veedor"):
                row = q(
                    f"""
                    SELECT asignado_a, motivo, {permite_reparacion_sql}
                      FROM ingresos
                     WHERE id=%s
                    """,
                    [ingreso_id],
                    one=True,
                )
                uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
                if not row or int(row.get("asignado_a") or 0) != int(uid or 0):
                    raise PermissionDenied("Solo el técnico asignado puede marcar como reparado")
                if _is_cotizacion_bloqueada_para_reparar(row.get("motivo"), row.get("permite_reparacion")):
                    raise ValidationError("Ingreso en cotización: primero habilite la reparación.")
        except PermissionDenied:
            raise
        except ValidationError:
            raise
        except Exception:
            # En duda, permitir solo a roles superiores
            if _rol(request) in ("tecnico", "jefe_veedor"):
                raise PermissionDenied("Solo el técnico asignado puede marcar como reparado")
        if _rol(request) not in ("tecnico", "jefe_veedor"):
            row = q(
                f"""
                SELECT motivo, {permite_reparacion_sql}
                  FROM ingresos
                 WHERE id=%s
                """,
                [ingreso_id],
                one=True,
            )
            if not row:
                raise ValidationError("Ingreso no encontrado")
            if _is_cotizacion_bloqueada_para_reparar(row.get("motivo"), row.get("permite_reparacion")):
                raise ValidationError("Ingreso en cotización: primero habilite la reparación.")
        # Leer estado/presupuesto previos para decidir envio (idempotencia)
        try:
            _prev_row = q(
                "SELECT estado::text AS estado, COALESCE(presupuesto_estado::text,'') AS presupuesto_estado FROM ingresos WHERE id=%s",
                [ingreso_id],
                one=True,
            ) or {}
            _prev_estado = (_prev_row.get("estado") or "").lower()
            _prev_presu = (_prev_row.get("presupuesto_estado") or "").lower()
        except Exception:
            _prev_estado = ""
            _prev_presu = ""
        _set_audit_user(request)
        # El presupuesto rechazado manda sobre el resultado del trabajo. El equipo puede
        # terminar armado y andando igual (se rechaza con la reparación ya hecha, o el
        # técnico cierra después del rechazo), pero la OS no se resolvió reparando: pisar
        # la resolución con 'reparado' hacía salir la orden de salida con "Resolución:
        # Reparado" sobre un presupuesto que el cliente no aceptó, y de paso hacía que
        # facturación y métricas tomaran el neto del presupuesto en vez del cobro pactado.
        resolucion_set = (
            "resolucion = CASE "
            "WHEN LOWER(COALESCE(presupuesto_estado::text,'')) = 'rechazado' "
            "THEN 'presupuesto_rechazado' ELSE 'reparado' END"
        )
        sets = sets_con_limpieza_de_cambio("estado='reparado'", resolucion_set)
        exec_void(f"UPDATE ingresos SET {sets} WHERE id=%s", [ingreso_id])
        if _prev_estado != "reparado":
            registrar_ingreso_event(
                ingreso_id,
                "reparado",
                de_estado=_prev_estado,
                usuario_id=getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None),
                comentario="Marcado como reparado",
            )

        # El equipo propio que sale del taller vuelve al parque de alquiler: la estantería
        # en NEXORA y el stock en STL son las dos mitades de "disponible para alquilar".
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        shelf = move_ingreso_to_rental_shelf(ingreso_id, actor_user_id=actor_user_id)
        if not shelf.get("applies"):
            try:
                if ingreso_is_demo_return(ingreso_id):
                    enqueue_demo_ready_transfer_for_ingreso(ingreso_id, actor_user_id=actor_user_id)
            except Exception:
                logger.exception("No se pudo encolar transferencia Bejerman STR a VAL", extra={"ingreso_id": ingreso_id})

        # Disparar aviso si pasó a 'reparado' y el presupuesto ya estaba aprobado.
        _should_send_mail = (_prev_estado != "reparado" and _prev_presu == "aprobado")
        if _should_send_mail:
            try:
                notify_repair_ready_for_remito(
                    ingreso_id,
                    request=request,
                    actor_name=getattr(request.user, "nombre", ""),
                )
            except Exception:
                logger.exception(
                    "No se pudo avisar que la reparación quedó lista para remito",
                    extra={"ingreso_id": ingreso_id},
                )

        resp = {"ok": True}
        if shelf.get("moved"):
            resp.update(
                {
                    "auto_moved": True,
                    "auto_moved_to": shelf.get("location_name") or RENTAL_SHELF_LOCATION_NAME,
                    "ubicacion_id": shelf.get("location_id"),
                    "ubicacion_nombre": shelf.get("location_name") or RENTAL_SHELF_LOCATION_NAME,
                }
            )
        return Response(resp)


class MarcarParaRepararView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe"])
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        permite_reparacion_sql = (
            "COALESCE(permite_reparacion, TRUE) AS permite_reparacion"
            if has_permite_reparacion
            else "TRUE AS permite_reparacion"
        )
        row = q(
            f"""
            SELECT
              estado::text AS estado,
              COALESCE(presupuesto_estado::text,'') AS presupuesto_estado,
              asignado_a,
              motivo,
              {permite_reparacion_sql}
            FROM ingresos
            WHERE id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row:
            raise ValidationError("Ingreso no encontrado")
        estado_cur = (row.get("estado") or "").lower()
        presu_cur = (row.get("presupuesto_estado") or "").lower()
        asignado_a = row.get("asignado_a")
        if _rol(request) == "jefe_veedor":
            uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
            if int(asignado_a or 0) != int(uid or 0):
                raise PermissionDenied("Solo el técnico asignado puede editar diagnóstico y reparación")
        if not asignado_a:
            raise ValidationError("Antes de reparar, asigna un técnico al ingreso")
        if estado_cur in ("reparado", "liberado", "entregado", "baja", "alquilado", "controlado_sin_defecto", "no_se_repara", *SALE_TICKET_STATES):
            raise ValidationError("No se puede marcar para reparar desde el estado actual")
        if _is_cotizacion_bloqueada_para_reparar(row.get("motivo"), row.get("permite_reparacion")):
            raise ValidationError("Ingreso en cotización: primero habilite la reparación.")

        _set_audit_user(request)
        try:
            exec_void(
                """
                UPDATE ingresos
                   SET estado='reparar',
                       presupuesto_estado = CASE
                         WHEN presupuesto_estado='presupuestado' THEN 'aprobado'
                         ELSE presupuesto_estado
                       END
                 WHERE id=%s
                """,
                [ingreso_id],
            )
        except Exception:
            exec_void("UPDATE ingresos SET estado='reparar' WHERE id=%s", [ingreso_id])

        # El MTTR arranca acá. Sin este evento el JOIN de `metricas_views` contra
        # `a_estado='reparar'` no matchea nada y la métrica queda en cero para siempre.
        if estado_cur != "reparar":
            registrar_ingreso_event(
                ingreso_id,
                "reparar",
                de_estado=estado_cur,
                usuario_id=getattr(getattr(request, "user", None), "id", None)
                or getattr(request, "user_id", None),
            )

        if presu_cur == "presupuestado":
            try:
                qid_row = q(
                    "SELECT id FROM quotes WHERE ingreso_id=%s ORDER BY COALESCE(version_num, 1) DESC, id DESC LIMIT 1",
                    [ingreso_id],
                    one=True,
                )
                qid = qid_row and qid_row.get("id")
                if qid:
                    exec_void(
                        """
                        UPDATE quotes
                           SET estado='aprobado',
                               fecha_aprobado=COALESCE(fecha_aprobado, now())
                         WHERE id=%s
                        """,
                        [qid],
                    )
            except Exception:
                pass

        email_sent = False
        try:
            info = q(
                """
                SELECT
                  u.email,
                  COALESCE(u.nombre,'') AS tecnico_nombre,
                  c.razon_social AS cliente,
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(d.numero_serie,'') AS numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno
                FROM ingresos t
                LEFT JOIN users   u ON u.id = t.asignado_a
                JOIN devices      d ON d.id = t.device_id
                JOIN customers    c ON c.id = d.customer_id
                LEFT JOIN marcas  b ON b.id = d.marca_id
                LEFT JOIN models  m ON m.id = d.model_id
                WHERE t.id=%s
                """,
                [ingreso_id],
                one=True,
            ) or {}
            to_email = (info.get("email") or "").strip()
            if to_email:
                os_txt = os_label(ingreso_id)
                try:
                    link = _frontend_url(request, f"/ingresos/{ingreso_id}") + "?tab=diagnostico"
                except Exception:
                    link = ""
                subject = f"{os_txt} - Autorizado para reparar"
                lines = [
                    f"Hola {info.get('tecnico_nombre') or ''},",
                    "",
                    f"Puede reparar la {os_txt}.",
                    "",
                    "Detalle del equipo:",
                    f"- Cliente: {info.get('cliente') or '-'}",
                    f"- Marca/Modelo: {info.get('marca') or '-'} / {info.get('modelo') or '-'}",
                    f"- Tipo: {info.get('tipo_equipo') or '-'}",
                    f"- Número de serie: {info.get('numero_interno') or info.get('numero_serie') or '-'}",
                ]
                if link:
                    lines.append("")
                    lines.append(f"Abrir hoja de servicio: {link}")
                try:
                    lines.append("")
                    lines.append("Aviso automático - no responder a este correo.")
                except Exception:
                    pass
                body = "\n".join(lines)
                try:
                    body = _email_append_footer_text(body)
                except Exception:
                    pass
                sent = send_mail(subject, body, getattr(settings, "DEFAULT_FROM_EMAIL", None), [to_email], fail_silently=True)
                email_sent = bool(sent and sent > 0)
        except Exception:
            pass

        resp = {
            "ok": True,
            "estado": "reparar",
            "presupuesto_estado": "aprobado" if presu_cur == "presupuestado" else presu_cur,
            "email_sent": bool(email_sent),
        }
        return Response(resp)


class QuitarRepararView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        # Pasa por la matriz (`action.ingreso.authorize_repair`); `require_jefe` la salteaba
        # y el permiso no se podía dar ni quitar por usuario.
        require_roles(request, ["jefe"])
        row = q(
            """
            SELECT
              estado::text AS estado,
              presupuesto_estado::text AS presupuesto_estado
            FROM ingresos
            WHERE id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row:
            raise ValidationError("Ingreso no encontrado")

        estado_cur = (row.get("estado") or "").strip().lower()
        if estado_cur != "reparar":
            raise ValidationError("Solo se puede quitar reparar cuando la OS está en estado reparar")

        presupuesto_actual = row.get("presupuesto_estado")
        estado_next, estado_transition = _estado_previo_a_estado(ingreso_id, "reparar")
        presupuesto_ya_aprobado_en_quote = (
            (presupuesto_actual or "").strip().lower() == "aprobado"
            and _latest_quote_is_approved(ingreso_id)
        )
        if presupuesto_ya_aprobado_en_quote:
            # El presupuesto está aprobado de verdad en `quotes` (aprobación formal,
            # no un efecto colateral de marcar "reparar" a mano): sólo "Anular
            # presupuesto" puede tocar eso, "Quitar reparar" no debe des-aprobarlo.
            presupuesto_next = presupuesto_actual
        else:
            presupuesto_next = _presupuesto_previo_en_misma_transicion(
                ingreso_id,
                estado_transition,
                presupuesto_actual,
            )

        _set_audit_user(request)
        # Al anular la resolución hay que soltar también el cambio de equipo: si no, la OS
        # queda arrastrando la serie de un cambio que ya no existe.
        resolucion_next = resolucion_al_anular_estado_final(presupuesto_next)
        exec_void(
            f"""
            UPDATE ingresos
               SET {sets_con_limpieza_de_cambio('estado=%s', 'resolucion=%s', 'presupuesto_estado=%s')}
             WHERE id=%s
            """,
            [estado_next, resolucion_next, presupuesto_next, ingreso_id],
        )
        return Response(
            {
                "ok": True,
                "estado_anterior": "reparar",
                "estado": estado_next,
                "resolucion": resolucion_next,
                "presupuesto_estado": presupuesto_next,
            }
        )


class HabilitarReparacionCotizacionView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe"])
        if not _ingresos_has_permite_reparacion_col():
            return Response(
                {"detail": "Schema no aplicado: falta columna ingresos.permite_reparacion."},
                status=400,
            )

        row = q(
            """
            SELECT motivo, COALESCE(permite_reparacion, TRUE) AS permite_reparacion
              FROM ingresos
             WHERE id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row:
            raise ValidationError("Ingreso no encontrado")

        motivo = row.get("motivo")
        if not _motivo_is_cotizacion_equipo(motivo):
            return Response(
                {
                    "ok": True,
                    "motivo": motivo,
                    "permite_reparacion": True,
                    "already_enabled": True,
                    "not_required": True,
                }
            )

        if bool(row.get("permite_reparacion")):
            return Response(
                {
                    "ok": True,
                    "motivo": motivo,
                    "permite_reparacion": True,
                    "already_enabled": True,
                }
            )

        _set_audit_user(request)
        exec_void(
            "UPDATE ingresos SET permite_reparacion=TRUE WHERE id=%s",
            [ingreso_id],
        )
        return Response(
            {
                "ok": True,
                "motivo": motivo,
                "permite_reparacion": True,
                "already_enabled": False,
            }
        )


# Un remito de venta cierra la OS como vendida: el equipo salió por el circuito comercial
# (RT), no por el de servicio técnico (RSS) ni por el de alquiler (RTA). El tipo va al
# principio del número guardado ("RT R 00001-00024051"), y `\b` evita que RTA/RTN entren.
_SALE_EXIT_REMITO_RE = re.compile(r"^\s*RT\b", re.IGNORECASE)
# Lo que puede documentar una salida que el operador marcó como venta: el RT o el RTN. Los
# aportes a NOVAMED salen con RTN, que no se factura (OS 29131, 14/09/2026). El RTN solo no
# alcanza para reconocer una venta -también es la entrega en alquiler-, así que vale únicamente
# junto con la marca explícita.
_SALE_MARKED_REMITO_RE = re.compile(r"^\s*RTN?\b", re.IGNORECASE)


def is_sale_exit_remito(remito) -> bool:
    """¿El remito de salida es una venta (RT)?

    Es el único dato que distingue una entrega de servicio de una venta: el equipo se
    factura y sale por comercial, así que la OS no cierra como 'entregado' sino como
    'vendido_entregado', y NEXORA no mueve stock (el RT ya lo movió en Bejerman).
    """
    return bool(_SALE_EXIT_REMITO_RE.match(str(remito or "")))


def resolve_ingreso_entrega_estado(ingreso_id: int, remito_salida=None, *, venta: bool = False) -> dict:
    """Determina a que estado debe pasar un ingreso al entregarse.

    Replica la deteccion de alquiler/MG/venta pendiente que usa
    EntregarIngresoView, para que cualquier otro flujo que registre una
    entrega (p.ej. la sincronizacion desde ordenes de entrega) aplique
    exactamente las mismas reglas.

    `remito_salida` es opcional para no romper a los llamadores que no lo tienen a mano;
    cuando viene y es un RT manda por encima de la heurística de alquiler/MG. `venta` es la
    marca explícita del operador ("salió vendido"), que también admite el RTN de un aporte.
    """
    # El tipo de comprobante no depende de la base: si la lectura de abajo falla, la venta
    # tiene que seguir cerrando como venta y no caerse a 'entregado'.
    es_venta = bool(venta) or is_sale_exit_remito(remito_salida)
    estado_to = "vendido_entregado" if es_venta else "entregado"
    set_alquilado = False
    clear_alquilado = es_venta
    try:
        sin_entrega_sql = (
            "COALESCE(t.cambio_sin_entrega, FALSE)" if _has_device_swap_schema() else "FALSE"
        )
        info = q(
            f"""
            SELECT t.alquilado,
                   t.estado,
                   LOWER(COALESCE(t.resolucion,'')) AS resolucion,
                   {sin_entrega_sql} AS cambio_sin_entrega,
                   COALESCE(d.numero_serie,'') AS numero_serie,
                   COALESCE(d.numero_interno,'') AS numero_interno,
                   COALESCE(d.mg_estado,'activo') AS mg_estado,
                   COALESCE(loc.nombre,'') AS ubicacion_nombre
              FROM ingresos t
              LEFT JOIN devices d ON d.id = t.device_id
              LEFT JOIN locations loc ON loc.id = t.ubicacion_id
             WHERE t.id=%s
            """,
            [ingreso_id],
            one=True,
        ) or {}
        # En un cambio de equipo el que queda acá es el equipo del cliente, no uno que se
        # alquila. Si además se lo convirtió en MG, la heurística de abajo lo marcaría
        # 'alquilado' y quedaría figurando en poder de un cliente que ya no lo tiene.
        #
        # La parte de pago se chequea por su marca y no por la resolución: al equipo tomado
        # se lo repara, y ahí la resolución pasa a 'reparado'.
        if (info.get("resolucion") or "") == "cambio" or info.get("cambio_sin_entrega"):
            return {"estado_to": "entregado", "set_alquilado": False, "clear_alquilado": False}
        # Una venta cierra la OS aunque el equipo sea MG o venga de la estantería: el RT ya
        # lo sacó del patrimonio, así que no puede quedar marcado como alquilado.
        if es_venta:
            return {"estado_to": "vendido_entregado", "set_alquilado": False, "clear_alquilado": True}
        ns = (info.get("numero_serie") or "").strip()
        ni = (info.get("numero_interno") or "").strip()
        pat = re.compile(r"\bMG \d{4}\b", re.IGNORECASE)
        is_mg = bool(pat.search(ns) or pat.search(ni)) or (
            ns.strip().upper().startswith("MG ") or ni.strip().upper().startswith("MG ")
        )
        loc_norm = _norm_txt(info.get("ubicacion_nombre") or "")
        target_loc_norm = _norm_txt("Estanteria de Alquiler")
        estado_actual = (info.get("estado") or "").strip().lower()
        mg_vendido = (info.get("mg_estado") or "").strip().lower() == "inactivo_venta"
        if estado_actual == "vendido_pendiente_entrega":
            estado_to = "vendido_entregado"
            clear_alquilado = True
        elif bool(info.get("alquilado")) or (is_mg and not mg_vendido) or loc_norm == target_loc_norm:
            estado_to = "alquilado"
            set_alquilado = True
    except Exception:
        # No bloquear la entrega por errores de deteccion
        pass
    return {"estado_to": estado_to, "set_alquilado": set_alquilado, "clear_alquilado": clear_alquilado}


class EntregarIngresoView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    @transaction.atomic
    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe", "jefe_veedor", "admin", "recepcion"])
        ingreso = q("SELECT id FROM ingresos WHERE id=%s FOR UPDATE", [ingreso_id], one=True)
        if not ingreso:
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        derivacion_activa = get_active_derivation(ingreso_id, for_update=True)
        if derivacion_activa:
            return Response(active_derivation_block_payload(derivacion_activa), status=409)
        data = request.data or {}
        remito = (data.get("remito_salida") or "").strip()
        if not remito:
            return Response({"detail": "remito_salida requerido"}, status=400)
        # "El equipo salió vendido": la marca la pone el operador en la hoja de servicio, y
        # entonces el comprobante puede ser un RT o el RTN de un aporte. Cualquier otro no
        # documenta una venta.
        marca_venta = _parse_bool_or_default(data.get("venta"), default=False)
        if marca_venta and not _SALE_MARKED_REMITO_RE.match(remito):
            return Response(
                {"detail": "Para cerrar por venta el remito tiene que ser un RT o un RTN."},
                status=400,
            )
        factura = (data.get("factura_numero") or "").strip() or None
        fecha_entrega = data.get("fecha_entrega") or None
        retira_persona = (data.get("retira_persona") or "").strip()
        estado_to = "entregado"
        set_alquilado = False
        clear_alquilado = False
        # Si es CAMBIO, verificar la serie contra el equipo vinculado al cerrar.
        #
        # Antes esto vivía dentro de un `try/except Exception: pass`, así que cualquier
        # error de lectura dejaba pasar la entrega sin verificar nada: justo lo contrario
        # de lo que la verificación existe para evitar (entregar el equipo equivocado).
        serial_confirm = (data.get("serial_confirm") or "").strip()
        es_cambio_de_equipo = False
        if _has_table_column("ingresos", "resolucion"):
            # El helper también reconoce la parte de pago, que sobrevive a que la
            # resolución pase a 'reparado'. Si acá se mirara sólo `resolucion`, la entrega
            # volvería a mover stock al cliente por un equipo que ya es nuestro.
            es_cambio_de_equipo = ingreso_tiene_cambio_de_equipo(ingreso_id)
            if es_cambio_de_equipo and _has_table_column("ingresos", "serial_cambio"):
                swap = get_ingreso_device_swap(ingreso_id) or {}
                db_serial = (swap.get("equipo_entregado") or "").strip()
                if db_serial:
                    if not serial_confirm:
                        return Response(
                            {"detail": "Verificá la serie del equipo que se entrega."},
                            status=400,
                        )
                    # El snapshot puede venir como "NS | MG 0001": alcanza con que la serie
                    # tipeada coincida con alguna de las dos partes.
                    partes = {p.strip().upper() for p in db_serial.split("|") if p.strip()}
                    if serial_confirm.upper() not in partes:
                        return Response(
                            {
                                "detail": (
                                    f"La serie no coincide con el equipo del cambio ({db_serial})."
                                )
                            },
                            status=409,
                        )
        deteccion = resolve_ingreso_entrega_estado(ingreso_id, remito_salida=remito, venta=marca_venta)
        estado_to = deteccion["estado_to"]
        set_alquilado = deteccion["set_alquilado"]
        clear_alquilado = deteccion["clear_alquilado"]
        comentarios_update = None
        if retira_persona:
            try:
                row = q("SELECT comentarios FROM ingresos WHERE id=%s", [ingreso_id], one=True) or {}
                existing = (row.get("comentarios") or "").strip()
            except Exception:
                existing = ""
            entrega_line = f"Entrega: retira {retira_persona}"
            comentarios_update = f"{existing}\n{entrega_line}" if existing else entrega_line

        _set_audit_user(request)
        with transaction.atomic():
            with connection.cursor() as cur:
                dash_id = _dash_location_id()
                sets = [
                    "estado=%s",
                    "remito_salida=%s",
                    "factura_numero=COALESCE(%s, factura_numero)",
                    "fecha_entrega=COALESCE(%s, now())",
                    "ubicacion_id = COALESCE(%s, ubicacion_id)",
                ]
                params = [estado_to, remito, factura, fecha_entrega, dash_id]
                if set_alquilado:
                    sets.append("alquilado=true")
                if clear_alquilado:
                    sets.append("alquilado=false")
                    sets.append("alquiler_a=NULL")
                    sets.append("alquiler_remito=NULL")
                    sets.append("alquiler_fecha=NULL")
                if comentarios_update is not None:
                    sets.append("comentarios=%s")
                    params.append(comentarios_update)
                params.append(ingreso_id)
                cur.execute(
                    "UPDATE ingresos SET " + ", ".join(sets) + " WHERE id=%s",
                    params,
                )
        try:
            uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
            with transaction.atomic():
                exec_void(
                    """
                    INSERT INTO ingreso_events (ticket_id, a_estado, usuario_id, comentario)
                    SELECT %s, %s, %s, %s
                    WHERE NOT EXISTS (
                        SELECT 1 FROM ingreso_events
                         WHERE ingreso_id=%s AND a_estado=%s
                    )
                    """,
                    [
                        ingreso_id,
                        estado_to,
                        uid,
                        "Alquiler registrado" if estado_to == "alquilado" else (
                            "Entrega de venta registrada" if estado_to == "vendido_entregado" else "Entrega registrada"
                        ),
                        ingreso_id,
                        estado_to,
                    ],
                )
        except Exception:
            # No bloquear la entrega si falla la auditoria de eventos
            pass
        # En un cambio de equipo el que ingresó se queda en SEPID: no hay movimiento de
        # stock hacia el cliente. El que se lleva sale por su propia orden de entrega.
        #
        # En una venta tampoco: el RT o el RTN los emitió comercial y ya movieron el stock en
        # Bejerman. Encolar acá una transferencia más dejaría el depósito en negativo. Se decide
        # por el estado al que cierra y no por el tipo del remito: una venta pendiente cerrada
        # con un comprobante que no empieza con "RT" también es una venta.
        if not es_cambio_de_equipo and estado_to != "vendido_entregado":
            try:
                with transaction.atomic():
                    user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
                    if ingreso_targets_rental_stock(ingreso_id):
                        enqueue_stock_transfer_for_ingreso(ingreso_id, actor_user_id=user_id)
                    elif ingreso_is_demo_return(ingreso_id):
                        enqueue_demo_ready_transfer_for_ingreso(ingreso_id, actor_user_id=user_id)
                    else:
                        enqueue_client_ready_transfer_for_ingreso(ingreso_id, actor_user_id=user_id)
            except Exception:
                logger.exception("No se pudo encolar la transferencia Bejerman", extra={"ingreso_id": ingreso_id})
        return Response({"ok": True})


class DarBajaIngresoView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe", "jefe_veedor", "admin", "recepcion"])
        _set_audit_user(request)
        marked_baja = False
        baja_event_id = None
        bejerman_sync_job = None
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        with transaction.atomic():
            with connection.cursor() as cur:
                dash_id = _dash_location_id()
                row = _fetch_ingreso_device_mg_context(ingreso_id)
                if not row:
                    return Response({"detail": "Ingreso no encontrado"}, status=404)
                estado_actual = (row.get("estado") or "").lower()
                if estado_actual == "baja":
                    return Response({"ok": True, "estado": "baja"})
                if not row.get("es_propietario_mg"):
                    return Response({"detail": "Solo se puede dar de baja un equipo propio MG."}, status=400)
                cur.execute(
                    """
                    UPDATE ingresos
                       SET estado='baja',
                           ubicacion_id = COALESCE(%s, ubicacion_id)
                     WHERE id=%s
                    """,
                    [dash_id, ingreso_id],
                )
                marked_baja = True
        if marked_baja:
            try:
                with transaction.atomic():
                    exec_void(
                        """
                        UPDATE ingreso_baja_requests
                           SET accepted_at = COALESCE(accepted_at, NOW())
                         WHERE ingreso_id = %s
                           AND accepted_at IS NULL
                           AND canceled_at IS NULL
                        """,
                        [ingreso_id],
                    )
            except Exception:
                # Tabla opcional / no bloquear la baja real
                pass
        try:
            event_row = q(
                """
                INSERT INTO ingreso_events (ticket_id, a_estado, comentario)
                VALUES (%s, 'baja', 'Marcado como baja desde la hoja de servicio')
                RETURNING id
                """,
                [ingreso_id],
                one=True,
            )
            baja_event_id = event_row and event_row.get("id")
        except Exception:
            # No bloquear si la tabla o inserción falla
            pass
        if marked_baja:
            try:
                with transaction.atomic():
                    bejerman_sync_job = enqueue_stock_baja_for_ingreso(
                        ingreso_id,
                        ingreso_event_id=baja_event_id,
                        actor_user_id=actor_user_id,
                    )
            except Exception as exc:
                logger.exception("No se pudo encolar la baja Bejerman", extra={"ingreso_id": ingreso_id})
                bejerman_sync_job = {
                    "status": "not_queued",
                    "detail": "No se pudo encolar la sincronización Bejerman de baja.",
                    "error": str(exc),
                }
        patrimonial_email = None
        if marked_baja:
            patrimonial_email = _notify_estado_patrimonial(request, ingreso_id, "baja")
        return Response(
            {
                "ok": True,
                "estado": "baja",
                "bejerman_sync_job": bejerman_sync_job,
                "patrimonial_email": patrimonial_email,
            }
        )


class DarAltaIngresoView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe", "jefe_veedor", "admin", "recepcion"])
        _set_audit_user(request)
        marked_alta = False
        alta_event_id = None
        bejerman_sync_job = None
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        with transaction.atomic():
            with connection.cursor() as cur:
                taller_id = _taller_location_id()
                cur.execute("SELECT estado FROM ingresos WHERE id=%s", [ingreso_id])
                row = cur.fetchone()
                if not row:
                    return Response({"detail": "Ingreso no encontrado"}, status=404)
                estado_actual = (row[0] or "").lower()
                if estado_actual != "baja":
                    return Response({"detail": "Solo se puede dar de alta un ingreso en estado baja"}, status=400)
                cur.execute(
                    """
                    UPDATE ingresos
                       SET estado='ingresado',
                           ubicacion_id = COALESCE(%s, ubicacion_id)
                     WHERE id=%s
                    """,
                    [taller_id, ingreso_id],
                )
                marked_alta = True
        try:
            event_row = q(
                """
                INSERT INTO ingreso_events (ticket_id, a_estado, comentario)
                VALUES (%s, 'ingresado', 'Marcado como alta desde la hoja de servicio')
                RETURNING id
                """,
                [ingreso_id],
                one=True,
            )
            alta_event_id = event_row and event_row.get("id")
        except Exception:
            # No bloquear si la tabla o inserción falla
            pass
        if marked_alta:
            try:
                with transaction.atomic():
                    bejerman_sync_job = enqueue_stock_alta_for_ingreso(
                        ingreso_id,
                        ingreso_event_id=alta_event_id,
                        actor_user_id=actor_user_id,
                    )
            except Exception as exc:
                logger.exception("No se pudo encolar el alta Bejerman", extra={"ingreso_id": ingreso_id})
                bejerman_sync_job = {
                    "status": "not_queued",
                    "detail": "No se pudo encolar la sincronización Bejerman de alta.",
                    "error": str(exc),
                }
        patrimonial_email = None
        if marked_alta:
            patrimonial_email = _notify_estado_patrimonial(request, ingreso_id, "alta")
        return Response(
            {
                "ok": True,
                "estado": "ingresado",
                "bejerman_sync_job": bejerman_sync_job,
                "patrimonial_email": patrimonial_email,
            }
        )


def find_mg_numero_interno_conflict(device_id, numero_interno: str):
    """Devuelve el device que ya ocupa ese número interno MG, o None.

    El número se compara con la clave de `mg_codes`, para que 'MG 12', 'MG12' y 'MG 0012' no
    convivan. La expresión anterior (con `LPAD('\2',4,'0')`) sí los dejaba convivir.
    """
    if connection.vendor == "postgresql":
        return q(
            f"""
            SELECT id
              FROM devices
             WHERE id <> %s
               AND {sql_is_mg_code('numero_interno')}
               AND {sql_mg_key('numero_interno')} = {sql_mg_key('%s')}
             LIMIT 1
            """,
            [device_id, numero_interno],
            one=True,
        )
    return q(
        "SELECT id FROM devices WHERE id <> %s AND numero_interno = %s LIMIT 1",
        [device_id, numero_interno],
        one=True,
    )


def convertir_device_a_propio_mg(
    device_id: int, numero_interno: str, mg_owner_id: int, *, ingreso_id: int | None = None
) -> None:
    """Pasa un equipo al patrimonio de MG BIO.

    Lo comparten la conversión manual desde la hoja de servicio
    (`IngresoConvertirPropioMgView`) y el cambio de equipo, donde el equipo que trajo
    el cliente queda para nosotros como parte de pago.

    **El número interno es opcional**: lo que define el patrimonio es el `customer_id`, no
    el código MG. `resolve_mg_flags` ya trata como propio a todo equipo del cliente MG BIO,
    así que un equipo tomado en parte de pago queda como nuestro y habilita el circuito de
    baja a Desguace aunque nunca se le asigne un MG.

    Con `ingreso_id` se guarda antes quién era el dueño: los papeles de la OS resuelven el
    cliente por `devices.customer_id`, así que sin ese snapshot la orden de salida que firma
    el cliente pasaba a decir 'MG BIO S.R.L.'.

    No valida el duplicado de número interno: eso lo hace el llamador con
    `find_mg_numero_interno_conflict`, porque cada uno arma su propia respuesta de error.
    """
    if ingreso_id and _has_table_column("ingresos", "cliente_original_id"):
        # Sólo si el equipo todavía no era nuestro, y sin pisar un snapshot anterior: una
        # segunda conversión no puede reescribir al cliente que trajo el equipo.
        exec_void(
            """
            UPDATE ingresos t
               SET cliente_original_id = d.customer_id
              FROM devices d
             WHERE t.id = %s
               AND d.id = %s
               AND t.cliente_original_id IS NULL
               AND d.customer_id IS DISTINCT FROM %s
            """,
            [ingreso_id, device_id, mg_owner_id],
        )
    sets = ["customer_id = %s"]
    params = [mg_owner_id]
    if (numero_interno or "").strip():
        sets.append("numero_interno = %s")
        params.append(numero_interno)
    if _has_mg_schema():
        sets.append("mg_estado = 'activo'")
    if _has_table_column("devices", "mg_inactivo_desde"):
        sets.append("mg_inactivo_desde = NULL")
    for column_name in (
        "mg_venta_fecha",
        "mg_venta_factura_numero",
        "mg_venta_remito_numero",
        "mg_venta_observaciones",
        "mg_venta_usuario_id",
        "mg_venta_customer_id",
        "mg_venta_numero_alternativo",
    ):
        if _has_table_column("devices", column_name):
            sets.append(f"{column_name} = NULL")
    if _has_table_column("devices", "alquilado"):
        sets.append("alquilado = FALSE")
    if _has_table_column("devices", "alquiler_a"):
        sets.append("alquiler_a = NULL")
    params.append(device_id)
    exec_void(f"UPDATE devices SET {', '.join(sets)} WHERE id=%s", params)


class IngresoConvertirPropioMgView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    @transaction.atomic
    def post(self, request, ingreso_id: int):
        require_any_permission(request, ["action.devices_preventivos.manage", "action.ingreso.baja_alta"])
        _set_audit_user(request)

        data = request.data or {}
        numero_interno_raw = (data.get("numero_interno") or "").strip()
        if not numero_interno_raw:
            return Response({"detail": "numero_interno requerido"}, status=400)

        numero_interno = normalize_mg(numero_interno_raw)
        if not numero_interno:
            return Response(
                {
                    "detail": "El número interno debe tener formato MG 0001.",
                    "conflict_type": "MG_NUMERO_INTERNO_INVALIDO",
                },
                status=400,
            )

        has_mg_schema = _has_mg_schema()
        mg_estado_sql = (
            "COALESCE(d.mg_estado,'activo') AS mg_estado,"
            if has_mg_schema
            else "'activo' AS mg_estado,"
        )

        ingreso = q(
            f"""
            SELECT
              i.id,
              i.device_id,
              COALESCE(CAST(i.motivo AS TEXT), '') AS motivo,
              d.customer_id,
              COALESCE(d.numero_interno,'') AS numero_interno_actual,
              COALESCE(d.numero_serie,'') AS numero_serie,
              {mg_estado_sql}
              COALESCE(d.alquilado,false) AS alquilado
            FROM ingresos i
            LEFT JOIN devices d ON d.id = i.device_id
            WHERE i.id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not ingreso:
            return Response({"detail": "Ingreso no encontrado"}, status=404)

        device_id = ingreso.get("device_id")
        if not device_id:
            return Response(
                {
                    "detail": "El ingreso no tiene un equipo asociado.",
                    "conflict_type": "INGRESO_SIN_DEVICE",
                },
                status=400,
            )

        mg_owner_id = _mg_owner_customer_id()
        if not mg_owner_id:
            return Response(
                {
                    "detail": "No existe el cliente base 'MG BIO'.",
                    "conflict_type": "MG_OWNER_CUSTOMER_MISSING",
                },
                status=404,
            )

        ingreso_mg_flags = resolve_mg_flags(
            {
                "customer_id": ingreso.get("customer_id"),
                "numero_interno": ingreso.get("numero_interno_actual") or "",
                "numero_serie": ingreso.get("numero_serie") or "",
                "mg_estado": ingreso.get("mg_estado") or "activo",
            },
            mg_owner_id,
        )
        if _is_mg_owner_customer(ingreso.get("customer_id"), mg_owner_id) and not ingreso_mg_flags.get("mg_inactivo_venta"):
            return Response(
                {
                    "detail": "El equipo ya figura como equipo propio de MG BIO.",
                    "conflict_type": "EQUIPO_YA_PROPIO_MG",
                },
                status=400,
            )

        conflict = find_mg_numero_interno_conflict(device_id, numero_interno)
        if conflict:
            return Response(
                {
                    "detail": "El número interno MG ya está asignado a otro equipo.",
                    "conflict_type": "MG_DUPLICATE",
                    "conflict_device_id": conflict.get("id"),
                    "numero_interno_input": numero_interno,
                },
                status=400,
            )

        convertir_device_a_propio_mg(device_id, numero_interno, mg_owner_id, ingreso_id=ingreso_id)

        mg_estado_sql = "COALESCE(d.mg_estado,'activo') AS mg_estado," if _has_mg_schema() else "'activo' AS mg_estado,"
        row = q(
            f"""
            SELECT
              d.id,
              d.customer_id,
              COALESCE(c.razon_social,'') AS customer_nombre,
              COALESCE(d.numero_interno,'') AS numero_interno,
              COALESCE(d.numero_serie,'') AS numero_serie,
              COALESCE(d.alquilado,false) AS alquilado,
              COALESCE(d.alquiler_a,'') AS alquiler_a,
              {mg_estado_sql}
              COALESCE(b.nombre,'') AS marca,
              COALESCE(m.nombre,'') AS modelo
            FROM devices d
            LEFT JOIN customers c ON c.id = d.customer_id
            LEFT JOIN marcas b ON b.id = d.marca_id
            LEFT JOIN models m ON m.id = d.model_id
            WHERE d.id=%s
            """,
            [device_id],
            one=True,
        ) or {}
        row.update(resolve_mg_flags(row, mg_owner_id))
        row["es_cliente_mg_owner"] = bool(row.get("es_propietario_mg"))
        patrimonial_email = _notify_estado_patrimonial(request, ingreso_id, "alta_mg")

        return Response(
            {
                "ok": True,
                "ingreso_id": ingreso_id,
                "device": row,
                "patrimonial_email": patrimonial_email,
            }
        )


# ---- Cambio de equipo (permuta / parte de pago) ----
#
# El cliente trae su equipo, se le entrega otro y el suyo queda para SEPID. Son dos
# movimientos patrimoniales, no un campo de texto: `ingresos.device_id` sigue siendo el
# equipo que INGRESÓ y `ingresos.device_cambio_id` es el que SALIÓ.
#
# Bajo esta OS no sale nada para el cliente: el equipo nuevo se factura por su propia
# orden de entrega (remito RT) y el viejo se queda. Por eso el cambio corta el movimiento
# de stock STR->STC y la emisión del remito RSS; ver `RemitoSalidaPdfView`.

# Qué se hace con el equipo que el cliente deja. 'pendiente' es el único que NO define
# nada: el resto son las razones por las que el equipo **pasa a propiedad de SEPID**, y se
# diferencian en dónde termina la unidad en Bejerman.
#
#   pendiente -> sigue siendo del cliente hasta que alguien decida; no mueve nada.
#   mg        -> propio de SEPID, queda en el parque (alquiler/uso interno). Sin movimiento.
#   venta     -> propio de SEPID y se pone a la venta: se transfiere al depósito VAL.
#   desguace  -> propio de SEPID y se desarma para repuestos: se transfiere a DES.
DEVICE_SWAP_DESTINOS = ("pendiente", "mg", "venta", "desguace")

# Los destinos que mueven patrimonio: pasan el equipo a nuestro nombre y, en 'venta' y
# 'desguace', además mueven stock real en Bejerman. Piden más permiso que registrar el
# cambio; recepción registra con 'pendiente' y jefatura define el destino.
DEVICE_SWAP_DESTINOS_PATRIMONIALES = ("mg", "venta", "desguace")
DEVICE_SWAP_DESTINO_PERMISOS = ("action.ingreso.baja_alta", "action.devices_preventivos.manage")

# Destinos que además de pasar el equipo a nuestro patrimonio lo mueven de depósito en
# Bejerman. El job se encola después de escribir el cambio, nunca antes.
#
# Se guarda el NOMBRE y no la función: el módulo la resuelve en el momento de usarla, así
# `patch("...ingresos_views.enqueue_stock_baja_for_ingreso")` sigue interceptando la
# llamada. Con la referencia directa el diccionario se quedaba con la función original y
# los tests salían a Bejerman de verdad.
DEVICE_SWAP_DESTINOS_CON_MOVIMIENTO = {
    "venta": ("enqueue_stock_venta_for_ingreso", "Pase a Ventas (VAL)"),
    "desguace": ("enqueue_stock_baja_for_ingreso", "Pase a Desguace"),
}

# Estados en los que la OS ya se cerró: ni se registra un cambio ni se deshace uno. El
# `delete` no lo miraba y le borraba la resolución a una orden ya entregada.
#
# `vendido_pendiente_entrega` entra acá aunque la OS todavía no esté entregada: el corte por
# cambio de `resolve_ingreso_entrega_estado` se evalúa **antes** que el de la venta, así que
# registrar un cambio sobre una venta pendiente la cerraba como 'entregado' y le borraba el
# ticket de venta.
DEVICE_SWAP_ESTADOS_CERRADOS = (
    "entregado",
    "baja",
    "vendido_entregado",
    "vendido_pendiente_entrega",
    "alquilado",
)

# Estados en los que el equipo está comprometido y no se puede entregar: o está dentro de
# otra reparación abierta, o ya está en poder de un cliente, o fue dado de baja.
DEVICE_SWAP_BLOCKING_STATES = (
    "ingresado",
    "diagnosticado",
    "presupuestado",
    "reparar",
    "alquilado",
    "vendido_pendiente_entrega",
    "baja",
)


def _has_device_swap_schema() -> bool:
    # Las dos columnas llegan juntas en `apply_device_swap_schema`: si falta alguna, el
    # esquema está a medio aplicar y conviene decirlo con el 500 explicativo antes de
    # que una consulta reviente por una columna inexistente.
    return _has_table_column("ingresos", "device_cambio_id") and _has_table_column(
        "ingresos", "cambio_sin_entrega"
    )


def _device_swap_schema_missing_response():
    return Response(
        {
            "detail": (
                "Schema de cambio de equipo no aplicado. "
                "Ejecuta apply_device_swap_schema."
            )
        },
        status=500,
    )


def _device_label(row) -> str:
    """Etiqueta corta de un equipo: N/S y número interno, lo que haya."""
    row = row or {}
    ns = (row.get("numero_serie") or "").strip()
    ni = (row.get("numero_interno") or "").strip()
    if ns and ni:
        return f"{ns} | {ni}"
    return ns or ni or ""


def _swap_equipment_article_prefix() -> str:
    """Prefijo de los códigos de artículo que son equipos completos.

    En Bejerman los equipos arrancan con '11'; el resto del maestro son repuestos,
    accesorios e insumos, que no se entregan en un cambio.
    """
    return str(getattr(settings, "BEJERMAN_EQUIPMENT_ARTICLE_PREFIX", "") or "11").strip()


def _swap_source_deposits() -> list[str]:
    """Depósitos donde se busca el equipo que se entrega.

    Por defecto VAL (el depósito de ventas, que es de donde sale un equipo vendido).
    Configurable porque puede haber unidades entregables en otros depósitos.
    """
    raw = getattr(settings, "BEJERMAN_SWAP_SOURCE_DEPOSITS", "") or "VAL"
    if isinstance(raw, (list, tuple)):
        codes = list(raw)
    else:
        codes = str(raw).split(",")
    out = [str(c).strip().upper() for c in codes if str(c).strip()]
    return out or ["VAL"]


def _swap_article_model_map(article_codes) -> dict:
    """Artículo de Bejerman -> modelo de NEXORA, para poder fichar el equipo.

    Si el artículo no está mapeado se devuelve vacío y el equipo se carga igual con lo
    que sabe Bejerman: trabar el mostrador por un mapeo faltante sería peor.
    """
    codes = sorted({str(c).strip().upper() for c in (article_codes or []) if str(c).strip()})
    if not codes:
        return {}
    rows = q(
        """
        SELECT
          UPPER(TRIM(am.article_code)) AS article_code,
          am.model_id,
          COALESCE(am.variante,'')     AS variante,
          COALESCE(m.nombre,'')        AS modelo,
          COALESCE(m.tipo_equipo,'')   AS tipo_equipo,
          m.marca_id,
          COALESCE(b.nombre,'')        AS marca
        FROM bejerman_article_mappings am
        JOIN models m ON m.id = am.model_id
        LEFT JOIN marcas b ON b.id = m.marca_id
        WHERE UPPER(TRIM(am.article_code)) = ANY(%s)
        """,
        [codes],
    ) or []
    # Un artículo puede tener varias variantes mapeadas; alcanza con la primera para
    # identificar marca/modelo/tipo, que es lo único que se usa para fichar el equipo.
    out = {}
    for row in rows:
        out.setdefault(row["article_code"], row)
    return out


def _swap_known_devices(serials) -> dict:
    """Serie normalizada -> device ya fichado en NEXORA.

    Sirve para reusar la ficha existente en vez de crear un duplicado, con la misma
    normalización que el índice único `uq_devices_ns_tipo_norm`.
    """
    keys = sorted({_norm_serial_key(s) for s in (serials or []) if _norm_serial_key(s)})
    if not keys:
        return {}
    rows = q(
        """
        SELECT
          UPPER(REPLACE(REPLACE(COALESCE(d.numero_serie,''), ' ', ''), '-', '')) AS clave,
          d.id,
          COALESCE(d.numero_serie,'')   AS numero_serie,
          COALESCE(d.numero_interno,'') AS numero_interno,
          COALESCE(d.tipo_equipo,'')    AS tipo_equipo,
          COALESCE(d.alquilado,false)   AS alquilado,
          COALESCE(ult.estado::text,'')       AS ultimo_estado,
          COALESCE(ult.derivacion_activa, FALSE) AS derivacion_activa
        FROM devices d
        LEFT JOIN LATERAL (
          SELECT i.estado::text AS estado,
                 EXISTS (
                   SELECT 1
                     FROM equipos_derivados active_ed
                    WHERE active_ed.ingreso_id = i.id
                      AND active_ed.fecha_entrega IS NULL
                      AND active_ed.estado IN ('derivado','en_servicio')
                 ) AS derivacion_activa
            FROM ingresos i
           WHERE i.device_id = d.id
           ORDER BY COALESCE(i.fecha_ingreso, i.fecha_creacion) DESC, i.id DESC
           LIMIT 1
        ) ult ON TRUE
        WHERE UPPER(REPLACE(REPLACE(COALESCE(d.numero_serie,''), ' ', ''), '-', '')) = ANY(%s)
        """,
        [keys],
    ) or []
    out = {}
    for row in rows:
        out.setdefault(row["clave"], row)
    return out


def _norm_serial_key(value) -> str:
    return (str(value or "")).upper().replace(" ", "").replace("-", "").strip()


def _swap_stock_candidates(company_key: str, actor_user_id, depositos, term: str, limit: int = 40):
    """Unidades entregables según el stock real de Bejerman.

    Antes esto listaba la tabla `devices` de NEXORA, que es el registro local de equipos:
    mostraba unidades de otros clientes y equipos en el taller, sin mirar disponibilidad.
    Un cambio entrega un equipo **vendido**, así que la fuente correcta es el stock del
    depósito de ventas, filtrado a artículos de equipo.
    """
    prefix = _swap_equipment_article_prefix()
    needle = (term or "").strip().lower()

    client = BejermanSDKClient(company_key=company_key, actor_user_id=actor_user_id)
    filas = []
    for deposito in depositos:
        for row in cached_stock_rows(client, company_key, deposito) or []:
            article_code = as_string(row.get("articleCode")).strip()
            partida = as_string(row.get("partida")).strip()
            if not partida or not article_code.upper().startswith(prefix.upper()):
                continue
            disponible = row.get("availableQuantity")
            if disponible is None:
                disponible = row.get("realQuantity")
            if disponible is None or disponible <= 0:
                continue
            filas.append(
                {
                    "article_code": article_code,
                    "partida": partida,
                    "deposito": as_string(row.get("depositCode")).strip().upper() or deposito,
                    "cantidad": disponible,
                }
            )

    descripciones = {}
    try:
        descripciones = article_descriptions(
            client, [f["article_code"] for f in filas], company_key
        ) or {}
    except Exception:
        # `article_descriptions` no levanta por diseño, pero el catálogo es opcional:
        # sin él la fila se muestra igual con el código.
        logger.exception("No se pudieron leer las descripciones de artículos")

    mapeos = _swap_article_model_map([f["article_code"] for f in filas])
    fichados = _swap_known_devices([f["partida"] for f in filas])

    salida = []
    for fila in filas:
        code_key = fila["article_code"].upper()
        mapeo = mapeos.get(code_key) or {}
        ficha = fichados.get(_norm_serial_key(fila["partida"])) or {}
        descripcion = (descripciones.get(fila["article_code"]) or "").strip()
        item = {
            **fila,
            "descripcion": descripcion,
            "marca": mapeo.get("marca") or "",
            "modelo": mapeo.get("modelo") or "",
            "tipo_equipo": mapeo.get("tipo_equipo") or ficha.get("tipo_equipo") or "",
            "model_id": mapeo.get("model_id"),
            "mapeado": bool(mapeo),
            # Si la unidad ya tiene ficha en NEXORA se reusa en vez de duplicarla.
            "device_id": ficha.get("id"),
            "numero_interno": ficha.get("numero_interno") or "",
        }
        # Bejerman puede tener stock de una unidad que en NEXORA está dentro de otra OS
        # abierta o figura alquilada. Se avisa acá y no al confirmar, que es donde el
        # backend la rechaza.
        estado_local = (ficha.get("ultimo_estado") or "").strip().lower()
        if ficha.get("alquilado"):
            item["motivo_no_disponible"] = "En NEXORA figura alquilado a un cliente"
        elif ficha.get("derivacion_activa"):
            item["motivo_no_disponible"] = "En NEXORA figura como derivado"
        elif estado_local in DEVICE_SWAP_BLOCKING_STATES:
            item["motivo_no_disponible"] = f"En NEXORA tiene un ingreso en estado {estado_local}"
        else:
            item["motivo_no_disponible"] = ""
        item["disponible"] = not item["motivo_no_disponible"]
        if needle:
            blob = " ".join(
                [
                    item["partida"],
                    item["article_code"],
                    descripcion,
                    item["marca"],
                    item["modelo"],
                    item["numero_interno"],
                ]
            ).lower()
            if needle not in blob:
                continue
        salida.append(item)

    salida.sort(key=lambda i: (not i["mapeado"], i["article_code"], i["partida"]))
    return salida[:limit]


def _fichar_device_desde_stock(*, partida: str, article_code: str, customer_id) -> int:
    """Devuelve el `devices.id` de una unidad de stock, fichándola si hace falta.

    En Bejerman la unidad es una partida (la serie); en NEXORA hace falta una fila en
    `devices` porque el cambio la referencia por FK. Si el artículo no está mapeado a un
    modelo, el equipo se ficha igual con la serie: trabar el mostrador por un mapeo
    faltante sería peor que una ficha incompleta, que después se completa desde Equipos.
    """
    serie = (partida or "").strip()
    if not serie:
        raise ValidationError("Falta la partida del equipo que se entrega.")

    mapeo = (_swap_article_model_map([article_code]) or {}).get((article_code or "").strip().upper()) or {}
    tipo_equipo = (mapeo.get("tipo_equipo") or "").strip()

    # La identidad del equipo es serie + tipo (ver uq_devices_ns_tipo_norm).
    existente = q(
        """
        SELECT id FROM devices
         WHERE UPPER(REPLACE(REPLACE(COALESCE(numero_serie,''), ' ', ''), '-', '')) = %s
           AND LOWER(COALESCE(TRIM(tipo_equipo),'')) = LOWER(%s)
         LIMIT 1
        """,
        [_norm_serial_key(serie), tipo_equipo],
        one=True,
    )
    if existente:
        return int(existente["id"])

    if not customer_id:
        raise ValidationError("El ingreso no tiene cliente: no se puede fichar el equipo entregado.")

    creado = q(
        """
        INSERT INTO devices (customer_id, marca_id, model_id, numero_serie, tipo_equipo)
        VALUES (%s, %s, %s, %s, NULLIF(%s,''))
        RETURNING id
        """,
        [customer_id, mapeo.get("marca_id"), mapeo.get("model_id"), serie, tipo_equipo],
        one=True,
    )
    return int(creado["id"])


def _load_swap_target_device(device_id: int):
    return q(
        """
        SELECT
          d.id,
          COALESCE(d.numero_serie,'')   AS numero_serie,
          COALESCE(d.numero_interno,'') AS numero_interno,
          COALESCE(d.alquilado,false)   AS alquilado,
          COALESCE(ult.estado::text,'')       AS ultimo_estado,
          COALESCE(ult.derivacion_activa, FALSE) AS derivacion_activa
        FROM devices d
        LEFT JOIN LATERAL (
          SELECT i.estado::text AS estado,
                 EXISTS (
                   SELECT 1
                     FROM equipos_derivados active_ed
                    WHERE active_ed.ingreso_id = i.id
                      AND active_ed.fecha_entrega IS NULL
                      AND active_ed.estado IN ('derivado','en_servicio')
                 ) AS derivacion_activa
            FROM ingresos i
           WHERE i.device_id = d.id
           ORDER BY COALESCE(i.fecha_ingreso, i.fecha_creacion) DESC, i.id DESC
           LIMIT 1
        ) ult ON TRUE
        WHERE d.id = %s
        """,
        [device_id],
        one=True,
    )


def get_ingreso_device_swap(ingreso_id: int):
    """Datos del cambio de equipo de una OS, o None si no tiene.

    Lo usan la vista de detalle, el PDF de la orden de salida y los cortes de stock.
    """
    if not _has_device_swap_schema():
        return None
    row = q(
        """
        SELECT
          i.device_cambio_id,
          COALESCE(i.serial_cambio,'')          AS serial_cambio,
          COALESCE(i.cambio_destino_viejo,'')   AS destino_viejo,
          COALESCE(i.cambio_sin_entrega, FALSE) AS sin_entrega,
          i.cambio_fecha,
          COALESCE(u.nombre,'')                 AS registrado_por,
          COALESCE(d.numero_serie,'')           AS numero_serie,
          COALESCE(d.numero_interno,'')         AS numero_interno,
          COALESCE(b.nombre,'')                 AS marca,
          COALESCE(m.nombre,'')                 AS modelo
        FROM ingresos i
        LEFT JOIN devices d ON d.id = i.device_cambio_id
        LEFT JOIN marcas b  ON b.id = d.marca_id
        LEFT JOIN models m  ON m.id = d.model_id
        LEFT JOIN users u   ON u.id = i.cambio_registrado_por
        WHERE i.id = %s
          AND (LOWER(COALESCE(i.resolucion,'')) = 'cambio'
               OR COALESCE(i.cambio_sin_entrega, FALSE))
        """,
        [ingreso_id],
        one=True,
    )
    if not row:
        return None
    # Los cambios anteriores a la FK sólo tienen el texto; se devuelven igual para que la
    # pantalla y el PDF los sigan mostrando.
    row["equipo_entregado"] = _device_label(row) or (row.get("serial_cambio") or "")
    row["vinculado"] = bool(row.get("device_cambio_id"))
    # En una parte de pago sin permuta no sale ningún equipo: `equipo_entregado` queda
    # vacío a propósito y esta marca es lo que la distingue de un cambio a medio cargar.
    row["sin_entrega"] = bool(row.get("sin_entrega"))
    return row


def ingreso_tiene_cambio_de_equipo(ingreso_id: int) -> bool:
    """True si la OS se resolvió con un cambio de equipo o una parte de pago.

    Corta el movimiento de stock al cliente y la emisión del remito RSS: bajo una OS así
    no sale nada para el cliente (si hay equipo de reemplazo, va por su propia venta) y el
    equipo que trajo queda en SEPID. Vale igual para la parte de pago sin permuta, donde
    directamente no se entrega nada.

    **La parte de pago no depende de `resolucion`.** Tomar el equipo es un hecho
    patrimonial, no un resultado de la reparación: al equipo tomado se lo repara igual, y
    marcar 'reparado' pisa la resolución. Si el corte dependiera sólo de `resolucion`, esa
    OS volvería a mover stock al cliente y a emitir un RSS por un equipo que es nuestro.
    """
    try:
        sin_entrega_sql = (
            "COALESCE(cambio_sin_entrega, FALSE)" if _has_device_swap_schema() else "FALSE"
        )
        row = q(
            f"""
            SELECT LOWER(COALESCE(resolucion,'')) AS resolucion,
                   {sin_entrega_sql} AS sin_entrega
              FROM ingresos WHERE id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row:
            return False
        return (row.get("resolucion") or "") == "cambio" or bool(row.get("sin_entrega"))
    except Exception:
        logger.exception("No se pudo leer la resolución del ingreso", extra={"ingreso_id": ingreso_id})
        return False


def clear_device_swap_sql_sets(preservar_parte_de_pago: bool = True) -> list[str]:
    """Fragmentos SET para limpiar el cambio cuando se pisa o se anula la resolución.

    Sin esto queda una OS 'reparada' arrastrando la serie de un cambio que ya no existe.

    `preservar_parte_de_pago` es el modo por defecto para las vistas que **pisan** la
    resolución (marcar reparado / no se repara / controlado sin defecto): una parte de pago
    no es un resultado de la reparación, así que reparar el equipo tomado no puede borrar
    el registro de que lo tomamos —y menos dejar sueltos el pase a MG BIO y el movimiento a
    Desguace, que no se revierten—. El deshacer explícito pasa `False` y limpia todo.
    """
    tiene_swap = _has_device_swap_schema()
    columnas = ["serial_cambio"] if _has_table_column("ingresos", "serial_cambio") else []
    if not tiene_swap:
        return [f"{col} = NULL" for col in columnas]
    columnas += ["device_cambio_id", "cambio_destino_viejo", "cambio_fecha", "cambio_registrado_por"]
    if not preservar_parte_de_pago:
        return [f"{col} = NULL" for col in columnas] + ["cambio_sin_entrega = FALSE"]
    guarda = "COALESCE(cambio_sin_entrega, FALSE)"
    return [f"{col} = CASE WHEN {guarda} THEN {col} ELSE NULL END" for col in columnas]


def sets_con_limpieza_de_cambio(*sets: str) -> str:
    """Arma la cláusula SET sumándole la limpieza del cambio de equipo.

    Lo usan las vistas que pisan o anulan `resolucion`. Los fragmentos que recibe no
    llevan parámetros propios más allá de los que ya pasa el llamador: la limpieza son
    todos `= NULL` o `CASE` sin placeholders, así que no altera su orden.

    La parte de pago se conserva: ver `clear_device_swap_sql_sets`.
    """
    return ", ".join([*sets, *clear_device_swap_sql_sets()])


def resolucion_al_anular_estado_final(presupuesto_estado) -> str | None:
    """Resolución que queda al deshacer un estado final de reparación.

    Quitar el "reparado" deshace el resultado del trabajo, no la respuesta del cliente:
    con el presupuesto rechazado la OS sigue resuelta por ese lado y tiene que seguir
    saliendo como "Presupuesto rechazado". Volverla a "sin definir" obligaba a elegir la
    resolución de nuevo a mano en cada salida, con el riesgo de que saliera otra cosa.

    Devuelve None cuando no hay nada que preservar (el caso normal: se limpia).
    """
    if (presupuesto_estado or "").strip().lower() == "rechazado":
        return "presupuesto_rechazado"
    return None


class IngresoCambioEquipoCandidatosView(APIView):
    """Unidades entregables, tomadas del stock real de Bejerman.

    No alcanza con `DevicesListView`: pide `page.devices_preventivos` (que recepción no
    tiene) y además lista el parque local de equipos, sin mirar stock ni disponibilidad.
    """

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ingreso_id: int):
        require_permission(request, "action.ingreso.device_swap")
        term = (request.GET.get("q") or "").strip()

        row = q(
            "SELECT empresa_bejerman, empresa_facturar FROM ingresos WHERE id=%s",
            [ingreso_id],
            one=True,
        )
        if not row:
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        try:
            company_key = require_company(
                row.get("empresa_bejerman") or row.get("empresa_facturar") or "SEPID"
            ).key
        except ValueError:
            return Response({"detail": "Empresa Bejerman no válida"}, status=400)

        disponibles = _swap_source_deposits()
        pedido = (request.GET.get("deposito") or "").strip().upper()
        depositos = [pedido] if pedido else disponibles

        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        try:
            items = _swap_stock_candidates(company_key, actor_user_id, depositos, term)
        except (BejermanSdkConfigError, BejermanSdkUnavailable, BejermanSdkResponseError) as exc:
            # Sin Bejerman no hay lista posible: se dice por qué, en vez de mostrar vacío.
            return Response({"detail": str(exc), "conflict_type": "BEJERMAN_UNAVAILABLE"}, status=502)

        return Response(
            {
                "items": items,
                "depositos": disponibles,
                "deposito": depositos[0] if len(depositos) == 1 else "",
                "prefijo_articulo": _swap_equipment_article_prefix(),
            }
        )


class IngresoCambioEquipoView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ingreso_id: int):
        require_permission(request, "action.ingreso.device_swap")
        return Response(get_ingreso_device_swap(ingreso_id) or {})

    @transaction.atomic
    def post(self, request, ingreso_id: int):
        require_permission(request, "action.ingreso.device_swap")
        if not _has_device_swap_schema():
            return _device_swap_schema_missing_response()
        _set_audit_user(request)

        data = request.data or {}
        destino = (data.get("destino_viejo") or "pendiente").strip().lower()
        if destino not in DEVICE_SWAP_DESTINOS:
            return Response(
                {"detail": f"destino_viejo inválido: {', '.join(DEVICE_SWAP_DESTINOS)}"},
                status=400,
            )
        # Registrar el cambio es del mostrador, pero decidir el destino firme no: 'mg' pasa el
        # equipo al patrimonio de MG BIO y 'desguace' además mueve stock real en Bejerman.
        # Esa misma conversión, por su endpoint propio (`IngresoConvertirPropioMgView`), exige
        # estos permisos; sin este chequeo el cambio de equipo era la puerta de atrás.
        # Recepción registra con 'pendiente' y el destino lo define después quien puede.
        if destino in DEVICE_SWAP_DESTINOS_PATRIMONIALES:
            require_any_permission(request, list(DEVICE_SWAP_DESTINO_PERMISOS))

        ingreso = q(
            """
            SELECT i.id,
                   i.device_id,
                   i.estado::text AS estado,
                   d.customer_id,
                   COALESCE(d.numero_serie,'')   AS numero_serie,
                   COALESCE(d.numero_interno,'') AS numero_interno
              FROM ingresos i
              LEFT JOIN devices d ON d.id = i.device_id
             WHERE i.id = %s
            """,
            [ingreso_id],
            one=True,
        )
        if not ingreso:
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        estado_actual = (ingreso.get("estado") or "").strip().lower()
        if estado_actual in DEVICE_SWAP_ESTADOS_CERRADOS:
            return Response(
                {
                    "detail": (
                        f"La orden ya está en estado {estado_actual}: no se puede registrar "
                        "un cambio de equipo."
                    )
                },
                status=400,
            )

        # Parte de pago sin permuta: tomamos el equipo del cliente y no entregamos ninguno.
        # Es explícito y no se deduce de que falte el equipo, porque "no vino el device_id"
        # también es lo que pasa cuando la pantalla se olvidó de mandarlo.
        sin_entrega = bool(data.get("sin_entrega"))

        device_id = data.get("device_id")
        try:
            device_id = int(device_id) if device_id not in (None, "") else 0
        except (TypeError, ValueError):
            return Response({"detail": "device_id inválido"}, status=400)

        target = None
        if sin_entrega:
            if device_id or (data.get("partida") or "").strip():
                return Response(
                    {
                        "detail": (
                            "Una parte de pago sin entrega no lleva equipo de salida: "
                            "sacá el equipo elegido o destildá la opción."
                        )
                    },
                    status=400,
                )
            device_id = 0
        else:
            # El equipo se elige del stock de Bejerman, donde la unidad es una partida y puede
            # no tener ficha en NEXORA todavía. Si no vino un device_id, se ficha acá.
            partida = (data.get("partida") or "").strip()
            if not device_id and partida:
                try:
                    device_id = _fichar_device_desde_stock(
                        partida=partida,
                        article_code=(data.get("article_code") or "").strip(),
                        customer_id=ingreso.get("customer_id"),
                    )
                except ValidationError as exc:
                    return Response({"detail": str(exc.detail[0] if isinstance(exc.detail, list) else exc.detail)}, status=400)

            if not device_id:
                return Response({"detail": "Elegí el equipo que se entrega."}, status=400)
            if device_id == ingreso.get("device_id"):
                return Response(
                    {"detail": "El equipo que se entrega tiene que ser distinto al que ingresó."},
                    status=400,
                )

            target = _load_swap_target_device(device_id)
            if not target:
                return Response({"detail": "El equipo elegido no existe."}, status=404)
            target_estado = (target.get("ultimo_estado") or "").strip().lower()
            if target.get("alquilado"):
                return Response(
                    {"detail": "El equipo elegido figura alquilado a un cliente."}, status=400
                )
            if target.get("derivacion_activa"):
                return Response(
                    {"detail": "El equipo elegido continúa derivado."}, status=409
                )
            if target_estado in DEVICE_SWAP_BLOCKING_STATES:
                return Response(
                    {"detail": f"El equipo elegido tiene un ingreso en estado {target_estado}."},
                    status=400,
                )

        # Los destinos que pasan el equipo a nuestro patrimonio se validan antes de escribir
        # nada, para no dejar el cambio a medias.
        numero_interno_mg = ""
        mg_owner_id = None
        if destino in DEVICE_SWAP_DESTINOS_PATRIMONIALES:
            mg_owner_id = _mg_owner_customer_id()
            if not mg_owner_id:
                return Response(
                    {
                        "detail": "No existe el cliente base 'MG BIO'.",
                        "conflict_type": "MG_OWNER_CUSTOMER_MISSING",
                    },
                    status=404,
                )

        if destino == "mg":
            # El número interno es OPCIONAL: un equipo tomado en parte de pago es de MG BIO
            # aunque nunca se le asigne un MG. Si viene, tiene que ser válido y único.
            numero_interno_raw = (data.get("numero_interno_mg") or "").strip()
            if numero_interno_raw:
                numero_interno_mg = normalize_mg(numero_interno_raw)
                if not numero_interno_mg:
                    return Response(
                        {
                            "detail": "El número interno debe tener formato MG 0001.",
                            "conflict_type": "MG_NUMERO_INTERNO_INVALIDO",
                        },
                        status=400,
                    )
                conflict = find_mg_numero_interno_conflict(
                    ingreso.get("device_id"), numero_interno_mg
                )
                if conflict:
                    return Response(
                        {
                            "detail": "El número interno MG ya está asignado a otro equipo.",
                            "conflict_type": "MG_DUPLICATE",
                            "conflict_device_id": conflict.get("id"),
                            "numero_interno_input": numero_interno_mg,
                        },
                        status=400,
                    )

        actor_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        serial_snapshot = _device_label(target) if target else ""

        exec_void(
            """
            UPDATE ingresos
               SET resolucion = 'cambio',
                   device_cambio_id = %s,
                   serial_cambio = %s,
                   cambio_destino_viejo = %s,
                   cambio_fecha = now(),
                   cambio_registrado_por = %s,
                   cambio_sin_entrega = %s
             WHERE id = %s
            """,
            [
                device_id or None,
                serial_snapshot or None,
                destino,
                actor_id,
                sin_entrega,
                ingreso_id,
            ],
        )

        # Los dos destinos "firmes" pasan el equipo al patrimonio de MG BIO. Lo que define
        # la propiedad es el cliente, no el código MG: `resolve_mg_flags` ya trata como
        # propio a todo equipo de MG BIO, así que a partir de acá el equipo habilita el
        # circuito de baja/alta aunque no tenga número interno.
        patrimonial_email = None
        bejerman_sync_job = None
        if destino in DEVICE_SWAP_DESTINOS_PATRIMONIALES:
            convertir_device_a_propio_mg(
                ingreso.get("device_id"), numero_interno_mg, mg_owner_id, ingreso_id=ingreso_id
            )
            patrimonial_email = _notify_estado_patrimonial(request, ingreso_id, "alta_mg")

        if destino == "venta":
            # El equipo tomado como parte de pago ya es propio aunque no reciba un número
            # MG. VAL es su depósito físico en Bejerman y "Sepid" su proyección local para
            # que quede visible como disponible para venta, no mezclado con alquiler.
            ensure_default_locations()
            sepid = q(
                "SELECT id FROM locations WHERE LOWER(nombre)=LOWER(%s) LIMIT 1",
                ["Sepid"],
                one=True,
            )
            if not sepid:
                raise ValidationError("No se pudo preparar la ubicación Sepid (VAL).")
            exec_void("UPDATE ingresos SET ubicacion_id=%s WHERE id=%s", [sepid["id"], ingreso_id])

        encolar = DEVICE_SWAP_DESTINOS_CON_MOVIMIENTO.get(destino)
        if encolar:
            # El equipo sale de su depósito actual y va al que corresponde al destino: VAL
            # si se pone a la venta, Desguace si se desarma. Es el mismo mecanismo que la
            # baja de un MG. No se toca el estado de la OS: tiene que poder liberarse y
            # entregarse igual, porque el cliente se lleva el equipo del cambio y necesita
            # su orden de salida.
            enqueue_name, etiqueta = encolar
            try:
                with transaction.atomic():
                    bejerman_sync_job = globals()[enqueue_name](
                        ingreso_id, actor_user_id=actor_id
                    )
            except Exception as exc:
                logger.exception(
                    "No se pudo encolar el movimiento de stock del cambio",
                    extra={"ingreso_id": ingreso_id, "destino": destino},
                )
                bejerman_sync_job = {
                    "status": "not_queued",
                    "detail": f"No se pudo encolar el movimiento en Bejerman ({etiqueta}).",
                    "error": str(exc),
                }

        recibido = _device_label(ingreso) or "equipo del cliente"
        try:
            # Savepoint propio: si la traza falla, Postgres deja la transacción abortada y
            # se perdería el cambio que sí se registró bien.
            #
            # `a_estado` es NOT NULL y el cambio no mueve el estado de la OS, así que se
            # repite el actual en vez de inventar una transición. Sin cast explícito: la
            # columna puede ser el enum `ticket_state` o TEXT según la base.
            with transaction.atomic():
                exec_void(
                    """
                    INSERT INTO ingreso_events (ticket_id, de_estado, a_estado, usuario_id, comentario)
                    VALUES (%s, %s, %s, %s, %s)
                    """,
                    [
                        ingreso_id,
                        estado_actual,
                        estado_actual,
                        actor_id,
                        (
                            f"Parte de pago: queda en SEPID {recibido}, sin entrega de otro "
                            f"equipo (destino: {destino})"
                            if sin_entrega
                            else f"Cambio de equipo: se entrega {serial_snapshot}, queda en "
                            f"SEPID {recibido} (destino: {destino})"
                        ),
                    ],
                )
        except Exception:
            # El cambio ya quedó registrado; no bloquearlo por la traza de eventos.
            logger.exception("No se pudo registrar el evento del cambio", extra={"ingreso_id": ingreso_id})

        return Response(
            {
                "ok": True,
                "cambio": get_ingreso_device_swap(ingreso_id),
                "patrimonial_email": patrimonial_email,
                "bejerman_sync_job": bejerman_sync_job,
                # Cómo llamar al movimiento en la pantalla. Viaja desde acá porque el
                # destino decide a qué depósito va: el front lo tenía escrito como
                # "Pase a Desguace" y con 'venta' habría mentido.
                "bejerman_sync_label": encolar[1] if encolar else "",
            }
        )

    @transaction.atomic
    def delete(self, request, ingreso_id: int):
        """Deshace el cambio.

        Sólo suelta el vínculo y la resolución. **No revierte** el pase del equipo viejo al
        patrimonio de MG BIO ni el movimiento a Desguace: eso se deshace desde Equipos (dar
        de alta / reasignar cliente), con su propio circuito y su propia auditoría.
        """
        require_permission(request, "action.ingreso.device_swap")
        if not _has_device_swap_schema():
            return _device_swap_schema_missing_response()
        _set_audit_user(request)

        # Acotado a las OS resueltas como cambio: si no, deshacer sobre una OS resuelta de
        # otra forma le borraría la resolución sin que nadie lo haya pedido.
        if not ingreso_tiene_cambio_de_equipo(ingreso_id):
            return Response(
                {"detail": "La orden no tiene un cambio de equipo registrado."}, status=400
            )

        # Misma guarda que el `post`: una OS ya cerrada no se toca. Sin esto, deshacer sobre
        # una orden entregada le borraba la resolución con la que se entregó y la dejaba sin
        # el dato que explica por qué no salió un RSS.
        estado_row = q("SELECT estado::text AS estado FROM ingresos WHERE id = %s", [ingreso_id], one=True)
        estado_actual = ((estado_row or {}).get("estado") or "").strip().lower()
        if estado_actual in DEVICE_SWAP_ESTADOS_CERRADOS:
            return Response(
                {
                    "detail": (
                        f"La orden ya está en estado {estado_actual}: no se puede deshacer "
                        "el cambio de equipo."
                    )
                },
                status=400,
            )

        # La resolución sólo se suelta si era el cambio: una parte de pago puede convivir
        # con 'reparado', y borrar esa resolución dejaría la OS sin resultado de taller.
        sets = [
            "resolucion = CASE WHEN LOWER(COALESCE(resolucion,'')) = 'cambio' THEN NULL ELSE resolucion END"
        ] + clear_device_swap_sql_sets(preservar_parte_de_pago=False)
        exec_void(f"UPDATE ingresos SET {', '.join(sets)} WHERE id = %s", [ingreso_id])
        return Response({"ok": True})


class IngresoCorreccionesHistoricasView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    @transaction.atomic
    def post(self, request, ingreso_id: int):
        # La resolución de un conflicto de ingreso abierto entra por acá con la autorización
        # ya chequeada por su propia vista (que además exige la prueba en Bejerman), igual que
        # `_skip_ris_preflight` en el alta por lote. Sin ese permiso propio, sigue pidiendo
        # el de correcciones históricas.
        if not getattr(request, "_historical_correction_authorized", False):
            require_permission(request, "action.ingreso.force_historical")
        _set_audit_user(request)

        if not _has_historical_corrections_schema():
            return Response(
                {
                    "detail": (
                        "Schema de correcciones históricas no aplicado. "
                        "Ejecuta apply_historical_corrections_schema."
                    )
                },
                status=500,
            )

        data = request.data or {}
        accion = str(data.get("accion") or "").strip().lower()
        if accion not in ("entrega", "alta_alquiler", "baja_alquiler", "baja_ingreso", "alta_ingreso"):
            return Response({"detail": "acción inválida"}, status=400)

        motivo = (data.get("motivo") or "").strip()
        if not motivo:
            return Response({"detail": "motivo requerido"}, status=400)

        fecha_efectiva = _parse_historical_effective_datetime(data.get("fecha_efectiva"))
        if not fecha_efectiva:
            return Response({"detail": "fecha_efectiva inválida"}, status=400)
        if fecha_efectiva > timezone.now():
            return Response({"detail": "fecha_efectiva no puede ser futura"}, status=400)

        notificar = _parse_bool_or_default(data.get("notificar"), True)
        uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)

        row = q(
            """
            SELECT id,
                   device_id,
                   estado,
                   COALESCE(alquilado,false) AS alquilado,
                   COALESCE(alquiler_a,'') AS alquiler_a,
                   COALESCE(alquiler_remito,'') AS alquiler_remito,
                   alquiler_fecha,
                   COALESCE(remito_salida,'') AS remito_salida,
                   fecha_entrega
              FROM ingresos
             WHERE id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row:
            return Response({"detail": "Ingreso no encontrado"}, status=404)

        de_estado = (row.get("estado") or "").lower()
        nuevo_estado = de_estado
        comentario = f"Corrección histórica forzada: {accion}"
        sets = []
        params = []
        dash_id = _dash_location_id()
        taller_id = _taller_location_id()
        rental_customer = None
        document_reference = ""

        if accion == "entrega":
            remito_salida = (data.get("remito_salida") or "").strip()
            factura_numero = (data.get("factura_numero") or "").strip()
            fecha_entrega = _parse_historical_effective_datetime(data.get("fecha_entrega"))
            if not fecha_entrega:
                fecha_entrega = fecha_efectiva
            sets.extend(
                [
                    "estado=%s",
                    "remito_salida=NULLIF(%s,'')",
                    "factura_numero=NULLIF(%s,'')",
                    "fecha_entrega=%s",
                ]
            )
            params.extend(["entregado", remito_salida, factura_numero, fecha_entrega])
            document_reference = remito_salida
            if dash_id:
                sets.append("ubicacion_id = COALESCE(%s, ubicacion_id)")
                params.append(dash_id)
            nuevo_estado = "entregado"

        elif accion == "alta_alquiler":
            rental_customer = _historical_rental_customer(data.get("alquiler_customer_id"))
            if not rental_customer:
                return Response(
                    {"detail": "Debe seleccionar un cliente válido de la lista para el alta de alquiler."},
                    status=400,
                )
            alquiler_a = (rental_customer.get("razon_social") or "").strip()
            alquiler_remito = (data.get("alquiler_remito") or "").strip()
            document_reference = alquiler_remito
            alquiler_fecha_raw = data.get("alquiler_fecha")
            if alquiler_fecha_raw in (None, ""):
                alquiler_fecha = timezone.localtime(fecha_efectiva).date()
            else:
                if isinstance(alquiler_fecha_raw, date):
                    alquiler_fecha = alquiler_fecha_raw
                else:
                    parsed = parse_date(str(alquiler_fecha_raw).strip())
                    if not parsed:
                        parsed_dt = _parse_datetime_or_date(alquiler_fecha_raw)
                        parsed = parsed_dt.date() if parsed_dt else None
                    if not parsed:
                        return Response({"detail": "alquiler_fecha inválida"}, status=400)
                    alquiler_fecha = parsed
            sets.extend(
                [
                    "alquilado=TRUE",
                    "estado=%s",
                    "alquiler_a=NULLIF(%s,'')",
                    "alquiler_remito=NULLIF(%s,'')",
                    "alquiler_fecha=%s",
                ]
            )
            params.extend(["alquilado", alquiler_a, alquiler_remito, alquiler_fecha])
            if dash_id:
                sets.append("ubicacion_id = COALESCE(%s, ubicacion_id)")
                params.append(dash_id)
            nuevo_estado = "alquilado"

        elif accion == "baja_alquiler":
            document_reference = (data.get("remito_baja_alquiler") or "").strip()
            sets.extend(
                [
                    "alquilado=FALSE",
                    "alquiler_a=NULL",
                    "alquiler_remito=NULL",
                    "alquiler_fecha=NULL",
                ]
            )
            if de_estado == "alquilado":
                sets.append("estado=%s")
                params.append("ingresado")
                if taller_id:
                    sets.append("ubicacion_id = COALESCE(%s, ubicacion_id)")
                    params.append(taller_id)
                nuevo_estado = "ingresado"
            else:
                nuevo_estado = de_estado or "ingresado"

        elif accion == "baja_ingreso":
            sets.append("estado=%s")
            params.append("baja")
            if dash_id:
                sets.append("ubicacion_id = COALESCE(%s, ubicacion_id)")
                params.append(dash_id)
            nuevo_estado = "baja"

        elif accion == "alta_ingreso":
            sets.append("estado=%s")
            params.append("ingresado")
            if taller_id:
                sets.append("ubicacion_id = COALESCE(%s, ubicacion_id)")
                params.append(taller_id)
            nuevo_estado = "ingresado"

        if not sets:
            return Response({"detail": "No se pudo determinar la corrección a aplicar"}, status=400)

        params.append(ingreso_id)
        exec_void(f"UPDATE ingresos SET {', '.join(sets)} WHERE id=%s", params)

        device_snapshot_updated = _sync_historical_device_snapshot(
            ingreso_id=ingreso_id,
            device_id=int(row.get("device_id")),
            accion=accion,
            alquiler_a=(rental_customer or {}).get("razon_social") or "",
            dash_id=dash_id,
            taller_id=(
                taller_id
                if accion == "alta_ingreso" or (accion == "baja_alquiler" and de_estado == "alquilado")
                else None
            ),
        )

        event_id = None
        try:
            with transaction.atomic():
                event_row = q(
                    """
                    INSERT INTO ingreso_events (ticket_id, de_estado, a_estado, usuario_id, ts, comentario)
                    VALUES (%s, NULLIF(%s,'')::ticket_state, %s::ticket_state, %s, %s, %s)
                    RETURNING id
                    """,
                    [ingreso_id, de_estado, nuevo_estado, uid, fecha_efectiva, comentario],
                    one=True,
                ) or {}
                event_id = event_row.get("id")
        except Exception:
            # Fallback para esquemas sin casts/columnas extendidas.
            try:
                with transaction.atomic():
                    event_row = q(
                        """
                        INSERT INTO ingreso_events (ticket_id, a_estado, usuario_id, ts, comentario)
                        VALUES (%s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        [ingreso_id, nuevo_estado, uid, fecha_efectiva, comentario],
                        one=True,
                    ) or {}
                    event_id = event_row.get("id")
            except Exception:
                pass

        if accion in ("baja_ingreso", "alta_ingreso"):
            try:
                if not ingreso_is_internal_equipment(ingreso_id):
                    bejerman_effect = {
                        "status": "not_applicable",
                        "detail": "La corrección de stock Bejerman no aplica porque el equipo no es propio de MG.",
                    }
                else:
                    enqueue = (
                        enqueue_stock_baja_for_ingreso
                        if accion == "baja_ingreso"
                        else enqueue_stock_alta_for_ingreso
                    )
                    job = enqueue(
                        ingreso_id,
                        ingreso_event_id=event_id,
                        actor_user_id=uid,
                    ) or {}
                    job_status = job.get("status") or "pending"
                    bejerman_effect = {
                        "status": job_status,
                        "job_id": job.get("id"),
                        "detail": (
                            "La corrección de stock quedó encolada para Bejerman."
                            if job_status in ("pending", "running")
                            else "La corrección de stock Bejerman ya estaba aplicada."
                            if job_status == "succeeded"
                            else job.get("last_error") or "La corrección de stock Bejerman quedó bloqueada."
                        ),
                    }
            except Exception as exc:
                logger.exception(
                    "No se pudo encolar la corrección histórica en Bejerman",
                    extra={"ingreso_id": ingreso_id, "accion": accion},
                )
                bejerman_effect = {
                    "status": "not_queued",
                    "detail": "NEXORA se actualizó, pero no se pudo encolar la corrección de stock en Bejerman.",
                    "error": str(exc),
                }
        else:
            bejerman_effect = _historical_bejerman_reference_effect(accion, document_reference)

        audit_payload = {
            **data,
            "accion": accion,
            "estado_anterior": de_estado,
            "estado_resultante": nuevo_estado,
            "alquiler_anterior": {
                "activo": bool(row.get("alquilado")),
                "cliente": row.get("alquiler_a") or "",
                "remito": row.get("alquiler_remito") or "",
                "fecha": row.get("alquiler_fecha"),
            },
            "device_snapshot_updated": device_snapshot_updated,
            "bejerman": bejerman_effect,
        }
        if rental_customer:
            audit_payload["alquiler_customer"] = {
                "id": rental_customer.get("id"),
                "razon_social": rental_customer.get("razon_social") or "",
                "alias_interno": rental_customer.get("alias_interno") or "",
                "cod_empresa": rental_customer.get("cod_empresa") or "",
            }
            audit_payload["alquiler_a"] = rental_customer.get("razon_social") or ""

        payload_json = json.dumps(audit_payload, ensure_ascii=False, default=str)
        if connection.vendor == "postgresql":
            exec_void(
                """
                INSERT INTO ingreso_historical_corrections
                  (ingreso_id, accion, fecha_efectiva, motivo, payload, notificar, usuario_id)
                VALUES (%s, %s, %s, %s, %s::jsonb, %s, %s)
                """,
                [ingreso_id, accion, fecha_efectiva, motivo, payload_json, notificar, uid],
            )
        else:
            exec_void(
                """
                INSERT INTO ingreso_historical_corrections
                  (ingreso_id, accion, fecha_efectiva, motivo, payload, notificar, usuario_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                [ingreso_id, accion, fecha_efectiva, motivo, payload_json, notificar, uid],
            )

        patrimonial_email = None
        if accion in ("baja_ingreso", "alta_ingreso") and notificar:
            patrimonial_email = _notify_estado_patrimonial(
                request,
                ingreso_id,
                "baja" if accion == "baja_ingreso" else "alta",
                fecha_evento=fecha_efectiva,
            )

        return Response(
            {
                "ok": True,
                "accion": accion,
                "estado_anterior": de_estado,
                "estado": nuevo_estado,
                "fecha_efectiva": fecha_efectiva,
                "notificar": bool(notificar),
                "patrimonial_email": patrimonial_email,
                "nexora": {
                    "status": "applied",
                    "detail": "La corrección se aplicó en NEXORA y quedó registrada en el historial.",
                    "device_snapshot_updated": device_snapshot_updated,
                    "alquiler_customer": audit_payload.get("alquiler_customer"),
                },
                "bejerman": bejerman_effect,
            }
        )


class ListosParaRetiroView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        require_permission(request, "page.liberados")
        mg_list_sql = _mg_list_select_sql()
        swap_list_sql = _swap_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        with connection.cursor() as cur:
            _set_audit_user(request)
            # Fallback sin vista: equipos en 'taller' con estado 'liberado'.
            # Devuelve columnas esperadas por el front: id, estado, resolucion,
            # cliente, numero_serie, numero_interno (MG), marca, modelo, tipo_equipo,
            # fecha_ingreso, fecha_listo (desde ingreso_events), fecha_entrega.
            cur.execute(
                f"""
                SELECT
                  t.id,
                  t.estado,
                  t.presupuesto_estado,
                  t.resolucion,
                  {swap_list_sql}
                  c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                  ev.ts AS fecha_listo
                FROM ingresos t
                JOIN devices   d ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                LEFT JOIN (
                  SELECT ingreso_id, MAX(ts) AS ts
                    FROM ingreso_events
                   WHERE a_estado = 'liberado'
                   GROUP BY ingreso_id
                ) ev ON ev.ingreso_id = t.id
                WHERE LOWER(loc.nombre) = LOWER(%s)
                  AND t.estado IN ('liberado','vendido_pendiente_entrega')
                ORDER BY
                  LOWER(COALESCE(c.razon_social, '')) ASC,
                  t.id DESC
                """,
                ["taller"],
            )
            rows = _fetchall_dicts(cur)
        # Sin esto la fila no trae `derivacion_activa` y la marca de derivado del
        # listado no aparece nunca, sin ningún error que lo delate.
        attach_active_derivations(rows)
        return Response(rows)


class GeneralPorClienteView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request, customer_id):
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT
                  t.id, t.estado, t.presupuesto_estado, t.fecha_ingreso, t.ubicacion_id,
                  q.fecha_emitido AS presupuesto_fecha_emision,
                  COALESCE(loc.nombre,'') AS ubicacion_nombre,
                  c.id AS customer_id, c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante
                FROM ingresos t
                JOIN devices   d ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN quotes q ON q.id = (
                  SELECT q2.id FROM quotes q2
                  WHERE q2.ingreso_id = t.id
                  ORDER BY COALESCE(q2.version_num, 1) DESC, q2.id DESC
                  LIMIT 1
                )
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                WHERE c.id = %s
                  AND LOWER(loc.nombre) = LOWER(%s)
                  AND t.estado NOT IN ('entregado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                ORDER BY t.fecha_ingreso DESC;
                """,
                [customer_id, "taller"],
            )
            rows = _fetchall_dicts(cur)
        attach_active_derivations(rows)
        return Response(rows)


class GeneralPorClienteExportView(APIView):
    """
    Exporta a Excel (.xlsx) el "general por cliente" (no entregados / no alquilados).

    GET /api/clientes/<customer_id>/general/export/
      Parametros opcionales:
        - ids: lista separada por comas para limitar la exportación a esos ingresos
        - ids: lista separada por comas para limitar la exportacion a esos ingresos
    Columnas del Excel:
      OS, Cliente, Equipo, N/S, Estado, Fecha ingreso
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, customer_id):
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        ids_raw = (request.GET.get("ids") or "").strip()
        ids = None
        if ids_raw:
            try:
                ids = [int(x) for x in ids_raw.split(",") if x.strip()]
            except Exception:
                return Response({"detail": "Parámetro 'ids' inválido"}, status=400)
            if len(ids) > 1000:
                return Response({"detail": "Demasiados IDs (maximo 1000)"}, status=400)
        with connection.cursor() as cur:
            _set_audit_user(request)
            if ids:
                cur.execute(
                    f"""
                    SELECT
                      t.id,
                      c.razon_social AS cliente,
                      {owner_list_sql}
                      COALESCE(b.nombre,'') AS marca,
                      COALESCE(m.nombre,'') AS modelo,
                      COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                      COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                      d.numero_serie,
                      COALESCE(d.numero_interno,'') AS numero_interno,
                      {mg_list_sql}
                      t.estado,
                      t.presupuesto_estado,
                      t.fecha_ingreso
                    FROM ingresos t
                    JOIN devices   d ON d.id = t.device_id
                    JOIN customers c ON c.id = d.customer_id
                    LEFT JOIN marcas b ON b.id = d.marca_id
                    LEFT JOIN models m ON m.id = d.model_id
                    LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                   WHERE c.id = %s
                     AND t.id = ANY(%s)
                     AND LOWER(loc.nombre) = LOWER(%s)
                     AND t.estado NOT IN ('entregado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                   ORDER BY t.fecha_ingreso DESC, t.id DESC
                    """,
                    [customer_id, ids, "taller"],
                )
            else:
                cur.execute(
                    f"""
                    SELECT
                      t.id,
                      c.razon_social AS cliente,
                      {owner_list_sql}
                      COALESCE(b.nombre,'') AS marca,
                      COALESCE(m.nombre,'') AS modelo,
                      COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                      COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                      d.numero_serie,
                      COALESCE(d.numero_interno,'') AS numero_interno,
                      {mg_list_sql}
                      t.estado,
                      t.presupuesto_estado,
                      t.fecha_ingreso
                    FROM ingresos t
                    JOIN devices   d ON d.id = t.device_id
                    JOIN customers c ON c.id = d.customer_id
                    LEFT JOIN marcas b ON b.id = d.marca_id
                    LEFT JOIN models m ON m.id = d.model_id
                    LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                   WHERE c.id = %s
                     AND LOWER(loc.nombre) = LOWER(%s)
                     AND t.estado NOT IN ('entregado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                   ORDER BY t.fecha_ingreso DESC, t.id DESC
                    """,
                    [customer_id, "taller"],
                )
            rows = _fetchall_dicts(cur)

        wb = Workbook()
        ws = wb.active
        ws.title = "General por cliente"
        headers = ["OS", "Cliente", "Equipo", "N/S", "Estado", "Presupuesto", "Fecha ingreso"]
        ws.append(headers)

        # formato de equipo y N/S via helpers _equipolabel_row y _ns_label

        def _presu_label(v):
            if not v:
                return "-"
            try:
                s = str(v).strip()
            except Exception:
                return str(v)
            if not s:
                return "-"
            if s == "presupuestado":
                return "Presupuestado"
            if s == "no_aplica":
                return "No aplica"
            return s[:1].upper() + s[1:]

        def _fecha_ingreso_label(v):
            if not v:
                return "-"
            try:
                if isinstance(v, datetime):
                    if timezone.is_aware(v):
                        v = timezone.localtime(v)
                    return v.strftime("%d-%m-%y")
                if isinstance(v, date):
                    return v.strftime("%d-%m-%y")
                if isinstance(v, str):
                    s = v.strip()
                    if not s:
                        return "-"
                    parsed = parse_datetime(s) or parse_date(s)
                    if isinstance(parsed, datetime):
                        if timezone.is_aware(parsed):
                            parsed = timezone.localtime(parsed)
                        return parsed.strftime("%d-%m-%y")
                    if isinstance(parsed, date):
                        return parsed.strftime("%d-%m-%y")
                    return s
                return v.strftime("%d-%m-%y")
            except Exception:
                return str(v)

        for r in rows:
            os_txt = os_label(r.get("id"))
            equipo = _equipolabel_row(r)
            ns_val = _ns_label(r)
            estado = (r.get("estado") or "-")
            presu = _presu_label(r.get("presupuesto_estado"))
            fecha_txt = _fecha_ingreso_label(r.get("fecha_ingreso"))

            ws.append([
                os_txt,
                _ingreso_customer_label(r) or "-",
                equipo,
                ns_val,
                estado,
                presu,
                fecha_txt,
            ])

        try:
            widths = [10, 40, 40, 20, 16, 16, 20]
            for idx, w in enumerate(widths, start=1):
                col = ws.column_dimensions[chr(64 + idx)]
                col.width = w
        except Exception:
            pass

        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)

        fname = f"general_cliente_{customer_id}_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
        resp = HttpResponse(
            bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f"attachment; filename=\"{fname}\""
        return resp


class AprobadosParaRepararView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        filtro_cotizacion = ""
        if has_permite_reparacion:
            filtro_cotizacion = """
                  AND NOT (
                        LOWER(CAST(t.motivo AS TEXT)) IN ('cotización de equipo', 'cotizacion de equipo')
                    AND COALESCE(t.permite_reparacion, TRUE) = FALSE
                  )
            """
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT
                  t.id,
                  t.estado,
                  t.presupuesto_estado,
                  COALESCE(t.resolucion, '') AS resolucion,
                  c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                  t.fecha_ingreso,
                  -- El serializer declara `fecha_aprobado`: con el alias viejo
                  -- (`fecha_aprobacion`) la columna se descartaba en silencio.
                  q.fecha_aprobado AS fecha_aprobado
                FROM ingresos t
                JOIN devices d ON d.id=t.device_id
                JOIN customers c ON c.id=d.customer_id
                LEFT JOIN marcas b ON b.id=d.marca_id
                LEFT JOIN models m ON m.id=d.model_id
                LEFT JOIN quotes q ON q.id = (
                  SELECT q2.id FROM quotes q2
                  WHERE q2.ingreso_id = t.id
                  ORDER BY COALESCE(q2.version_num, 1) DESC, q2.id DESC
                  LIMIT 1
                )
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                WHERE LOWER(loc.nombre) = LOWER(%s)
                  AND (
                        (t.presupuesto_estado = 'aprobado'
                        AND t.estado NOT IN ('reparado','controlado_sin_defecto','no_se_repara','entregado','liberado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado'))
                        OR t.estado = 'reparar'
                      )
                  AND t.estado NOT IN ('reparado','controlado_sin_defecto','no_se_repara','entregado','liberado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                  {filtro_cotizacion}
                ORDER BY COALESCE(q.fecha_aprobado, t.fecha_ingreso) ASC;
            """,
                ["taller"],
            )
            data = _fetchall_dicts(cur)
        attach_active_derivations(data)
        return Response(IngresoListItemSerializer(data, many=True).data)


class AprobadosYReparadosView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        mg_list_sql = _mg_list_select_sql()
        swap_list_sql = _swap_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        gate_presupuesto_sql = sql_presupuesto_bloquea_salida("t", "d")
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
              SELECT t.id, t.estado, t.presupuesto_estado,
                     COALESCE(t.resolucion, '') AS resolucion,
                     {swap_list_sql}
                     c.razon_social,
                     {owner_list_sql}
                     d.numero_serie,
                     COALESCE(d.numero_interno,'') AS numero_interno,
                     {mg_list_sql}
                     COALESCE(b.nombre,'') AS marca,
                     COALESCE(m.nombre,'') AS modelo,
                     COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                     COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                     t.motivo,
                     -- El front repite el mismo corte de presupuesto antes de habilitar el
                     -- botón (release-flow.js), y sin las garantías no puede distinguir la
                     -- reparación que se cobra de la que no.
                     COALESCE(t.garantia_fabrica, FALSE) AS garantia_fabrica,
                     COALESCE(t.garantia_reparacion, FALSE) AS garantia_reparacion,
                     t.fecha_ingreso,
                     ev.fecha_reparado
              FROM ingresos t
              JOIN devices d ON d.id=t.device_id
              JOIN customers c ON c.id=d.customer_id
              LEFT JOIN marcas b ON b.id=d.marca_id
              LEFT JOIN models m ON m.id=d.model_id
              LEFT JOIN locations loc ON loc.id = t.ubicacion_id
              LEFT JOIN (
                SELECT ingreso_id, MAX(ts) AS fecha_reparado
                FROM ingreso_events
                WHERE a_estado IN ('reparado','controlado_sin_defecto','no_se_repara')
                GROUP BY ingreso_id
              ) ev ON ev.ingreso_id = t.id
              WHERE LOWER(loc.nombre) = LOWER(%s)
                AND t.estado IN ('reparado','controlado_sin_defecto','no_se_repara')
                -- Nada del cliente sale sin la plata definida: ni el presupuesto emitido y
                -- sin respuesta, ni -desde 09/09/2026- la reparación cobrable que nunca se
                -- presupuestó. Primero se debe resolver el presupuesto: emitirlo y
                -- aprobarlo, rechazarlo o registrar explícitamente que no aplica.
                -- La regla completa, con sus excepciones, está en `service/release_gate.py`.
                AND NOT {gate_presupuesto_sql}
              ORDER BY (ev.fecha_reparado IS NULL) ASC, ev.fecha_reparado ASC, t.fecha_ingreso ASC;
            """,
                ["taller"],
            )
            data = _fetchall_dicts(cur)
        attach_active_derivations(data)
        return Response(IngresoListItemSerializer(data, many=True).data)


class AprobadosView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        mg_list_sql = _mg_list_select_sql()
        swap_list_sql = _swap_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT
                  t.id, t.estado, t.presupuesto_estado,
                  COALESCE(t.resolucion, '') AS resolucion,
                  {swap_list_sql}
                  c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                  t.fecha_ingreso,
                  qa.fecha_aprobado
                FROM ingresos t
                JOIN devices d ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                LEFT JOIN (
                  SELECT DISTINCT ON (ingreso_id) ingreso_id, fecha_aprobado
                  FROM quotes
                  WHERE estado = 'aprobado'
                  ORDER BY ingreso_id, fecha_aprobado DESC
                ) qa ON qa.ingreso_id = t.id
                WHERE t.presupuesto_estado = 'aprobado'
                  AND t.estado NOT IN ('liberado','entregado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                  AND LOWER(loc.nombre) = LOWER(%s)
                ORDER BY
                    COALESCE(qa.fecha_aprobado, t.fecha_ingreso) NULLS LAST,
                    t.fecha_ingreso,
                    t.id;
                """,
                ["taller"],
            )
            data = _fetchall_dicts(cur)
        attach_active_derivations(data)
        return Response(IngresoListItemSerializer(data, many=True).data)

class LiberadosView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        require_permission(request, "page.liberados")
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        with connection.cursor() as cur:
            _set_audit_user(request)
            cur.execute(
                f"""
                SELECT
                  t.id,
                  t.estado,
                  t.presupuesto_estado,
                  c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                  t.fecha_ingreso,
                  t.fecha_entrega
                FROM ingresos t
                JOIN devices d ON d.id=t.device_id
                JOIN customers c ON c.id=d.customer_id
                LEFT JOIN marcas b ON b.id=d.marca_id
                LEFT JOIN models m ON m.id=d.model_id
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                WHERE LOWER(loc.nombre) = LOWER(%s)
                  AND t.estado IN ('liberado','vendido_pendiente_entrega')
                ORDER BY t.fecha_entrega DESC, t.id DESC;
                """,
                ["taller"],
            )
            rows = _fetchall_dicts(cur)
        attach_active_derivations(rows)
        return Response(rows)


class GeneralEquiposView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        require_any_permission(request, ["page.ingresos_history", "page.logistics"])
        ubic = (request.GET.get("ubicacion") or "").strip()
        ubic_id_raw = (request.GET.get("ubicacion_id") or "").strip()
        estado_raw = (request.GET.get("estado") or "").strip()
        excluir_estados_raw = (request.GET.get("excluir_estados") or "").strip()
        solo_taller_raw = (request.GET.get("solo_taller") or "").strip().lower()
        q_raw = (request.GET.get("q") or "").strip()
        delivered_raw = (request.GET.get("delivered") or "").strip().lower()
        from_raw = (request.GET.get("from") or "").strip()
        to_raw = (request.GET.get("to") or "").strip()
        fecha_ingreso_from_raw = (request.GET.get("fecha_ingreso_from") or "").strip()
        fecha_ingreso_to_raw = (request.GET.get("fecha_ingreso_to") or "").strip()
        fecha_liberacion_from_raw = (request.GET.get("fecha_liberacion_from") or "").strip()
        fecha_liberacion_to_raw = (request.GET.get("fecha_liberacion_to") or "").strip()
        fecha_entrega_from_raw = (request.GET.get("fecha_entrega_from") or "").strip()
        fecha_entrega_to_raw = (request.GET.get("fecha_entrega_to") or "").strip()
        os_raw = (request.GET.get("os") or "").strip()
        cliente_raw = (request.GET.get("cliente") or "").strip()
        tipo_equipo_raw = (request.GET.get("tipo_equipo") or "").strip()
        marca_raw = (request.GET.get("marca") or "").strip()
        modelo_raw = (request.GET.get("modelo") or "").strip()
        variante_raw = (request.GET.get("variante") or "").strip()
        estado_q_raw = (request.GET.get("estado_q") or "").strip()
        numero_serie_raw = (request.GET.get("numero_serie") or "").strip()
        numero_interno_raw = (request.GET.get("numero_interno") or "").strip()
        sort_raw = (request.GET.get("sort") or "").strip()

        # paginación opcional (se mantiene respuesta original si no se especifica page_size)
        page_raw = (request.GET.get("page") or "").strip()
        page_size_raw = (request.GET.get("page_size") or "").strip()

        estados = [e.strip().lower() for e in estado_raw.split(",") if e.strip()] if estado_raw else []
        excluir_estados = [e.strip().lower() for e in excluir_estados_raw.split(",") if e.strip()] if excluir_estados_raw else []
        solo_taller_param_present = ("solo_taller" in request.GET)
        solo_taller = solo_taller_raw in ("1", "true", "yes", "y", "t")
        delivered = delivered_raw in ("1", "true", "yes", "y", "t")
        page = int(page_raw) if page_raw.isdigit() and int(page_raw) > 0 else 1
        try:
            page_size = int(page_size_raw) if page_size_raw else 0
        except Exception:
            page_size = 0
        if page_size < 0:
            page_size = 0
        if page_size > 0:
            # cota superior razonable
            page_size = min(page_size, 500)
        historico_filter_ctx = {
            "ubic": canonical_location_name(ubic) if ubic else "",
            "ubic_id_raw": ubic_id_raw,
            "estados": estados,
            "excluir_estados": excluir_estados,
            "solo_taller_param_present": solo_taller_param_present,
            "solo_taller": solo_taller,
            "delivered": delivered,
            "q_raw": q_raw,
            "os_raw": os_raw,
            "cliente_raw": cliente_raw,
            "tipo_equipo_raw": tipo_equipo_raw,
            "marca_raw": marca_raw,
            "modelo_raw": modelo_raw,
            "variante_raw": variante_raw,
            "estado_q_raw": estado_q_raw,
            "numero_serie_raw": numero_serie_raw,
            "numero_interno_raw": numero_interno_raw,
            "fecha_ingreso_from_raw": fecha_ingreso_from_raw,
            "fecha_ingreso_to_raw": fecha_ingreso_to_raw,
            "fecha_liberacion_from_raw": fecha_liberacion_from_raw,
            "fecha_liberacion_to_raw": fecha_liberacion_to_raw,
            "fecha_entrega_from_raw": fecha_entrega_from_raw,
            "fecha_entrega_to_raw": fecha_entrega_to_raw,
        }

        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        has_mg_sale_extended = _has_mg_sale_extended_schema()
        with connection.cursor() as cur:
            _set_audit_user(request)
            wh, params = [], []
            if delivered and not (fecha_entrega_from_raw or fecha_entrega_to_raw):
                fecha_entrega_from_raw = from_raw
                fecha_entrega_to_raw = to_raw

            def _apply_date_range(field_sql, from_val, to_val):
                if from_val:
                    try:
                        f = parse_date(from_val)
                        if f:
                            wh.append(f"DATE({field_sql}) >= %s")
                            params.append(f.isoformat())
                    except Exception:
                        pass
                if to_val:
                    try:
                        tdt = parse_date(to_val)
                        if tdt:
                            wh.append(f"DATE({field_sql}) <= %s")
                            params.append(tdt.isoformat())
                    except Exception:
                        pass
            def _apply_ilike(field_sql, value):
                if value:
                    wh.append(f"{field_sql} ILIKE %s")
                    params.append(f"%{value}%")
            if ubic:
                wh.append("LOWER(loc.nombre) = LOWER(%s)")
                params.append(canonical_location_name(ubic))
            else:
                # Compat: permitir filtrar por ubicacion_id (numérico)
                try:
                    if ubic_id_raw and str(int(ubic_id_raw)) == ubic_id_raw:
                        wh.append("t.ubicacion_id = %s")
                        params.append(int(ubic_id_raw))
                except Exception:
                    pass
            # Optional 'solo_taller' filter: only apply when explicitly requested
            if not (ubic or (ubic_id_raw and ubic_id_raw.isdigit())):
                if solo_taller_param_present and solo_taller:
                    wh.append("LOWER(loc.nombre) = LOWER(%s)")
                    params.append("taller")
            if estados:
                placeholders = ",".join(["%s"] * len(estados))
                wh.append(f"LOWER(TRIM(COALESCE(t.estado::text,''))) IN ({placeholders})")
                params.extend(estados)
            if excluir_estados:
                placeholders = ",".join(["%s"] * len(excluir_estados))
                wh.append(f"LOWER(TRIM(COALESCE(t.estado::text,''))) NOT IN ({placeholders})")
                params.extend(excluir_estados)
            # Busqueda flexible por varios campos (OS, cliente, equipo, estado, serie, interno, ubicacion).
            if q_raw:
                needle = q_raw.strip()
                if needle:
                    needle_ns = needle.replace(" ", "").upper()
                    import re as _re
                    es_codigo_mg = normalize_mg_code(needle) is not None
                    m_os = _re.match(r"^(?:OS\s*)?(\d+)$", needle, _re.IGNORECASE)
                    if es_codigo_mg:
                        # Criterio único de número interno (`mg_codes`): "MG1678", "MG 1678" y
                        # "MG 001678" traen el mismo equipo, esté guardado de la forma que esté.
                        clauses = [
                            f"{sql_mg_key('d.numero_interno')} = {sql_mg_key('%s')}",
                            f"{sql_mg_key('d.numero_serie')} = {sql_mg_key('%s')}",
                        ]
                        if has_mg_sale_extended:
                            clauses.append(f"{sql_mg_key('d.mg_venta_numero_alternativo')} = {sql_mg_key('%s')}")
                        wh.append("(" + " OR ".join(clauses) + ")")
                        params.extend([needle] * len(clauses))
                    else:
                        like = f"%{needle}%"
                        like_ns = f"%{needle_ns}%"
                        has_customer_alias = _has_table_column("customers", "alias_interno")
                        clauses = []
                        if m_os:
                            clauses.append("t.id = %s")
                        clauses.extend([
                            "t.id::text ILIKE %s",
                            "COALESCE(c.razon_social,'') ILIKE %s",
                            # El dueño va además de la razón social, no en su lugar: el
                            # contenedor sigue siendo buscable para quien lo busca a propósito.
                            f"{_owner_name_expr()} ILIKE %s",
                            *(["COALESCE(c.alias_interno,'') ILIKE %s"] if has_customer_alias else []),
                            "COALESCE(b.nombre,'') ILIKE %s",
                            "COALESCE(m.nombre,'') ILIKE %s",
                            "COALESCE(m.tipo_equipo,'') ILIKE %s",
                            "COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,''), '') ILIKE %s",
                            "COALESCE(t.estado::text,'') ILIKE %s",
                            "COALESCE(d.numero_serie,'') ILIKE %s",
                            "COALESCE(d.numero_interno,'') ILIKE %s",
                            "COALESCE(loc.nombre,'') ILIKE %s",
                            "REPLACE(UPPER(COALESCE(d.numero_interno,'')),' ','') LIKE %s",
                            "REPLACE(UPPER(COALESCE(d.numero_serie,'')),' ','') LIKE %s",
                        ])
                        if has_mg_sale_extended:
                            clauses.extend([
                                "COALESCE(d.mg_venta_numero_alternativo,'') ILIKE %s",
                                "REPLACE(UPPER(COALESCE(d.mg_venta_numero_alternativo,'')),' ','') LIKE %s",
                            ])
                        wh.append("(" + " OR ".join(clauses) + ")")
                        if m_os:
                            params.append(int(m_os.group(1)))
                        params.extend([like] * (11 + (1 if has_customer_alias else 0)) + [like_ns, like_ns])
                        if has_mg_sale_extended:
                            params.extend([like, like_ns])

            if os_raw:
                try:
                    import re as _re
                    m_os = _re.match(r"^(?:OS\s*)?(\d+)$", os_raw, _re.IGNORECASE)
                except Exception:
                    m_os = None
                if m_os:
                    wh.append("t.id = %s")
                    params.append(int(m_os.group(1)))
                else:
                    wh.append("t.id::text ILIKE %s")
                    params.append(f"%{os_raw}%")
            if cliente_raw:
                # El filtro de la columna Cliente busca por lo que la columna muestra: el
                # dueño real, y también el contenedor y el alias.
                partes = ["COALESCE(c.razon_social,'') ILIKE %s", f"{_owner_name_expr()} ILIKE %s"]
                if _has_table_column("customers", "alias_interno"):
                    partes.append("COALESCE(c.alias_interno,'') ILIKE %s")
                wh.append("(" + " OR ".join(partes) + ")")
                params.extend([f"%{cliente_raw}%"] * len(partes))
            _apply_ilike("COALESCE(m.tipo_equipo,'')", tipo_equipo_raw)
            _apply_ilike("COALESCE(b.nombre,'')", marca_raw)
            _apply_ilike("COALESCE(m.nombre,'')", modelo_raw)
            _apply_ilike(
                "COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,''), '')",
                variante_raw,
            )
            _apply_ilike("COALESCE(t.estado::text,'')", estado_q_raw)
            _apply_ilike("COALESCE(d.numero_serie,'')", numero_serie_raw)
            _apply_ilike("COALESCE(d.numero_interno,'')", numero_interno_raw)

            # filtros por fechas (rangos)
            _apply_date_range("t.fecha_ingreso", fecha_ingreso_from_raw, fecha_ingreso_to_raw)
            _apply_date_range("ev_lib.fecha_liberacion", fecha_liberacion_from_raw, fecha_liberacion_to_raw)
            _apply_date_range("t.fecha_entrega", fecha_entrega_from_raw, fecha_entrega_to_raw)

            if delivered:
                wh.append("t.fecha_entrega IS NOT NULL")

            where_sql = (" WHERE " + " AND ".join(wh)) if wh else ""

            sort_field = sort_raw[1:] if sort_raw.startswith("-") else sort_raw
            sort_direction = "DESC" if sort_raw.startswith("-") else "ASC"
            sort_map = {
                "os": "t.id",
                # Ordena por el nombre que se ve en la columna, no por el del contenedor.
                "cliente": f"NULLIF(LOWER({_customer_search_expr()}), '')",
                "tipo": "NULLIF(LOWER(COALESCE(m.tipo_equipo,'')), '')",
                "marca": "NULLIF(LOWER(COALESCE(b.nombre,'')), '')",
                "modelo": "NULLIF(LOWER(COALESCE(m.nombre,'')), '')",
                "variante": "NULLIF(LOWER(COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,''), '')), '')",
                "estado": "NULLIF(LOWER(COALESCE(t.estado::text,'')), '')",
                "numero_serie": "NULLIF(LOWER(COALESCE(d.numero_serie,'')), '')",
                "numero_interno": "NULLIF(LOWER(COALESCE(d.numero_interno,'')), '')",
                "fecha_ingreso": "t.fecha_ingreso",
                "fecha_liberacion": "ev_lib.fecha_liberacion",
                "fecha_entrega": "t.fecha_entrega",
            }
            sort_expression = sort_map.get(sort_field)
            order_sql = "ORDER BY t.id DESC"
            if sort_expression:
                order_sql = (
                    f"ORDER BY {sort_expression} {sort_direction} NULLS LAST, "
                    f"t.id {sort_direction}"
                )

            limit_sql = ""
            limit_params = []
            overfetch = 0
            if page_size > 0:
                overfetch = 1
                limit_sql = " LIMIT %s OFFSET %s"
                limit_params.extend([page_size + overfetch, max(0, (page - 1) * page_size)])

            sql = f"""
                SELECT
                  t.id, t.estado, t.presupuesto_estado, t.fecha_ingreso, ev_lib.fecha_liberacion, t.fecha_entrega, t.ubicacion_id,
                  q.fecha_emitido AS presupuesto_fecha_emision,
                  COALESCE(loc.nombre,'') AS ubicacion_nombre,
                  c.id AS customer_id, c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante
                FROM ingresos t
                JOIN devices   d ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN quotes q ON q.id = (
                  SELECT q2.id FROM quotes q2
                  WHERE q2.ingreso_id = t.id
                  ORDER BY COALESCE(q2.version_num, 1) DESC, q2.id DESC
                  LIMIT 1
                )
                LEFT JOIN (
                  SELECT ingreso_id, MAX(ts) AS fecha_liberacion
                  FROM ingreso_events
                  WHERE a_estado = 'liberado'
                  GROUP BY ingreso_id
                ) ev_lib ON ev_lib.ingreso_id = t.id
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                {where_sql}
                {order_sql}
                {limit_sql}
            """
            cur.execute(sql, params + limit_params)
            rows = _fetchall_dicts(cur)
            deleted_rows = _historico_deleted_rows(cur, historico_filter_ctx)
            if page <= 1 or q_raw or os_raw or estado_q_raw or numero_serie_raw or numero_interno_raw:
                rows = _merge_historico_rows(rows, deleted_rows, sort_raw)
        mg_owner_id = _mg_owner_customer_id()
        for row in rows:
            row.update(resolve_mg_flags(row, mg_owner_id))
            row["es_cliente_mg_owner"] = bool(row.get("es_propietario_mg"))
        # Antes de paginar: las filas borradas del histórico no tienen ingreso vivo y
        # el helper las deja en `derivacion_activa=False` sin romper.
        attach_active_derivations(rows)
        # compat: si no hay paginación pedida, seguir devolviendo lista
        if page_size == 0:
            return Response(rows)
        has_next = False
        if len(rows) > page_size:
            has_next = True
            rows = rows[:page_size]
        return Response({
            "items": rows,
            "page": page,
            "page_size": page_size,
            "has_next": bool(has_next),
        })


class GeneralEquiposExportView(APIView):
    """
    Exporta el histórico de ingresos en Excel (.xlsx) respetando los mismos
    filtros de GeneralEquiposView.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        ubic = (request.GET.get("ubicacion") or "").strip()
        ubic_id_raw = (request.GET.get("ubicacion_id") or "").strip()
        estado_raw = (request.GET.get("estado") or "").strip()
        excluir_estados_raw = (request.GET.get("excluir_estados") or "").strip()
        solo_taller_raw = (request.GET.get("solo_taller") or "").strip().lower()
        q_raw = (request.GET.get("q") or "").strip()
        delivered_raw = (request.GET.get("delivered") or "").strip().lower()
        from_raw = (request.GET.get("from") or "").strip()
        to_raw = (request.GET.get("to") or "").strip()
        fecha_ingreso_from_raw = (request.GET.get("fecha_ingreso_from") or "").strip()
        fecha_ingreso_to_raw = (request.GET.get("fecha_ingreso_to") or "").strip()
        fecha_liberacion_from_raw = (request.GET.get("fecha_liberacion_from") or "").strip()
        fecha_liberacion_to_raw = (request.GET.get("fecha_liberacion_to") or "").strip()
        fecha_entrega_from_raw = (request.GET.get("fecha_entrega_from") or "").strip()
        fecha_entrega_to_raw = (request.GET.get("fecha_entrega_to") or "").strip()
        os_raw = (request.GET.get("os") or "").strip()
        cliente_raw = (request.GET.get("cliente") or "").strip()
        tipo_equipo_raw = (request.GET.get("tipo_equipo") or "").strip()
        marca_raw = (request.GET.get("marca") or "").strip()
        modelo_raw = (request.GET.get("modelo") or "").strip()
        variante_raw = (request.GET.get("variante") or "").strip()
        estado_q_raw = (request.GET.get("estado_q") or "").strip()
        numero_serie_raw = (request.GET.get("numero_serie") or "").strip()
        numero_interno_raw = (request.GET.get("numero_interno") or "").strip()

        estados = [e.strip().lower() for e in estado_raw.split(",") if e.strip()] if estado_raw else []
        excluir_estados = [e.strip().lower() for e in excluir_estados_raw.split(",") if e.strip()] if excluir_estados_raw else []
        solo_taller_param_present = ("solo_taller" in request.GET)
        solo_taller = solo_taller_raw in ("1", "true", "yes", "y", "t")
        delivered = delivered_raw in ("1", "true", "yes", "y", "t")
        historico_filter_ctx = {
            "ubic": canonical_location_name(ubic) if ubic else "",
            "ubic_id_raw": ubic_id_raw,
            "estados": estados,
            "excluir_estados": excluir_estados,
            "solo_taller_param_present": solo_taller_param_present,
            "solo_taller": solo_taller,
            "delivered": delivered,
            "q_raw": q_raw,
            "os_raw": os_raw,
            "cliente_raw": cliente_raw,
            "tipo_equipo_raw": tipo_equipo_raw,
            "marca_raw": marca_raw,
            "modelo_raw": modelo_raw,
            "variante_raw": variante_raw,
            "estado_q_raw": estado_q_raw,
            "numero_serie_raw": numero_serie_raw,
            "numero_interno_raw": numero_interno_raw,
            "fecha_ingreso_from_raw": fecha_ingreso_from_raw,
            "fecha_ingreso_to_raw": fecha_ingreso_to_raw,
            "fecha_liberacion_from_raw": fecha_liberacion_from_raw,
            "fecha_liberacion_to_raw": fecha_liberacion_to_raw,
            "fecha_entrega_from_raw": fecha_entrega_from_raw,
            "fecha_entrega_to_raw": fecha_entrega_to_raw,
        }

        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()
        has_mg_sale_extended = _has_mg_sale_extended_schema()
        with connection.cursor() as cur:
            _set_audit_user(request)
            wh, params = [], []
            if delivered and not (fecha_entrega_from_raw or fecha_entrega_to_raw):
                fecha_entrega_from_raw = from_raw
                fecha_entrega_to_raw = to_raw

            def _apply_date_range(field_sql, from_val, to_val):
                if from_val:
                    try:
                        f = parse_date(from_val)
                        if f:
                            wh.append(f"DATE({field_sql}) >= %s")
                            params.append(f.isoformat())
                    except Exception:
                        pass
                if to_val:
                    try:
                        tdt = parse_date(to_val)
                        if tdt:
                            wh.append(f"DATE({field_sql}) <= %s")
                            params.append(tdt.isoformat())
                    except Exception:
                        pass

            def _apply_ilike(field_sql, value):
                if value:
                    wh.append(f"{field_sql} ILIKE %s")
                    params.append(f"%{value}%")

            if ubic:
                wh.append("LOWER(loc.nombre) = LOWER(%s)")
                params.append(canonical_location_name(ubic))
            else:
                try:
                    if ubic_id_raw and str(int(ubic_id_raw)) == ubic_id_raw:
                        wh.append("t.ubicacion_id = %s")
                        params.append(int(ubic_id_raw))
                except Exception:
                    pass

            if not (ubic or (ubic_id_raw and ubic_id_raw.isdigit())):
                if solo_taller_param_present and solo_taller:
                    wh.append("LOWER(loc.nombre) = LOWER(%s)")
                    params.append("taller")

            if estados:
                placeholders = ",".join(["%s"] * len(estados))
                wh.append(f"LOWER(TRIM(COALESCE(t.estado::text,''))) IN ({placeholders})")
                params.extend(estados)
            if excluir_estados:
                placeholders = ",".join(["%s"] * len(excluir_estados))
                wh.append(f"LOWER(TRIM(COALESCE(t.estado::text,''))) NOT IN ({placeholders})")
                params.extend(excluir_estados)

            if q_raw:
                needle = q_raw.strip()
                if needle:
                    needle_ns = needle.replace(" ", "").upper()
                    import re as _re
                    es_codigo_mg = normalize_mg_code(needle) is not None
                    m_os = _re.match(r"^(?:OS\s*)?(\d+)$", needle, _re.IGNORECASE)
                    if es_codigo_mg:
                        # Criterio único de número interno (`mg_codes`): "MG1678", "MG 1678" y
                        # "MG 001678" traen el mismo equipo, esté guardado de la forma que esté.
                        clauses = [
                            f"{sql_mg_key('d.numero_interno')} = {sql_mg_key('%s')}",
                            f"{sql_mg_key('d.numero_serie')} = {sql_mg_key('%s')}",
                        ]
                        if has_mg_sale_extended:
                            clauses.append(f"{sql_mg_key('d.mg_venta_numero_alternativo')} = {sql_mg_key('%s')}")
                        wh.append("(" + " OR ".join(clauses) + ")")
                        params.extend([needle] * len(clauses))
                    else:
                        like = f"%{needle}%"
                        like_ns = f"%{needle_ns}%"
                        has_customer_alias = _has_table_column("customers", "alias_interno")
                        clauses = []
                        if m_os:
                            clauses.append("t.id = %s")
                        clauses.extend([
                            "t.id::text ILIKE %s",
                            "COALESCE(c.razon_social,'') ILIKE %s",
                            # El dueño va además de la razón social, no en su lugar: el
                            # contenedor sigue siendo buscable para quien lo busca a propósito.
                            f"{_owner_name_expr()} ILIKE %s",
                            *(["COALESCE(c.alias_interno,'') ILIKE %s"] if has_customer_alias else []),
                            "COALESCE(b.nombre,'') ILIKE %s",
                            "COALESCE(m.nombre,'') ILIKE %s",
                            "COALESCE(m.tipo_equipo,'') ILIKE %s",
                            "COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,''), '') ILIKE %s",
                            "COALESCE(t.estado::text,'') ILIKE %s",
                            "COALESCE(d.numero_serie,'') ILIKE %s",
                            "COALESCE(d.numero_interno,'') ILIKE %s",
                            "COALESCE(loc.nombre,'') ILIKE %s",
                            "REPLACE(UPPER(COALESCE(d.numero_interno,'')),' ','') LIKE %s",
                            "REPLACE(UPPER(COALESCE(d.numero_serie,'')),' ','') LIKE %s",
                        ])
                        if has_mg_sale_extended:
                            clauses.extend([
                                "COALESCE(d.mg_venta_numero_alternativo,'') ILIKE %s",
                                "REPLACE(UPPER(COALESCE(d.mg_venta_numero_alternativo,'')),' ','') LIKE %s",
                            ])
                        wh.append("(" + " OR ".join(clauses) + ")")
                        if m_os:
                            params.append(int(m_os.group(1)))
                        params.extend([like] * (11 + (1 if has_customer_alias else 0)) + [like_ns, like_ns])
                        if has_mg_sale_extended:
                            params.extend([like, like_ns])

            if os_raw:
                try:
                    import re as _re
                    m_os = _re.match(r"^(?:OS\s*)?(\d+)$", os_raw, _re.IGNORECASE)
                except Exception:
                    m_os = None
                if m_os:
                    wh.append("t.id = %s")
                    params.append(int(m_os.group(1)))
                else:
                    wh.append("t.id::text ILIKE %s")
                    params.append(f"%{os_raw}%")

            if cliente_raw:
                # El filtro de la columna Cliente busca por lo que la columna muestra: el
                # dueño real, y también el contenedor y el alias.
                partes = ["COALESCE(c.razon_social,'') ILIKE %s", f"{_owner_name_expr()} ILIKE %s"]
                if _has_table_column("customers", "alias_interno"):
                    partes.append("COALESCE(c.alias_interno,'') ILIKE %s")
                wh.append("(" + " OR ".join(partes) + ")")
                params.extend([f"%{cliente_raw}%"] * len(partes))
            _apply_ilike("COALESCE(m.tipo_equipo,'')", tipo_equipo_raw)
            _apply_ilike("COALESCE(b.nombre,'')", marca_raw)
            _apply_ilike("COALESCE(m.nombre,'')", modelo_raw)
            _apply_ilike(
                "COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,''), '')",
                variante_raw,
            )
            _apply_ilike("COALESCE(t.estado::text,'')", estado_q_raw)
            _apply_ilike("COALESCE(d.numero_serie,'')", numero_serie_raw)
            _apply_ilike("COALESCE(d.numero_interno,'')", numero_interno_raw)

            _apply_date_range("t.fecha_ingreso", fecha_ingreso_from_raw, fecha_ingreso_to_raw)
            _apply_date_range("ev_lib.fecha_liberacion", fecha_liberacion_from_raw, fecha_liberacion_to_raw)
            _apply_date_range("t.fecha_entrega", fecha_entrega_from_raw, fecha_entrega_to_raw)

            if delivered:
                wh.append("t.fecha_entrega IS NOT NULL")

            where_sql = (" WHERE " + " AND ".join(wh)) if wh else ""

            sql = f"""
                SELECT
                  t.id, t.estado, t.presupuesto_estado, t.fecha_ingreso, ev_lib.fecha_liberacion, t.fecha_entrega, t.ubicacion_id,
                  COALESCE(loc.nombre,'') AS ubicacion_nombre,
                  c.id AS customer_id, c.razon_social,
                  {owner_list_sql}
                  d.numero_serie,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  {mg_list_sql}
                  COALESCE(b.nombre,'') AS marca,
                  COALESCE(m.nombre,'') AS modelo,
                  COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                  COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante
                FROM ingresos t
                JOIN devices   d ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN (
                  SELECT ingreso_id, MAX(ts) AS fecha_liberacion
                  FROM ingreso_events
                  WHERE a_estado = 'liberado'
                  GROUP BY ingreso_id
                ) ev_lib ON ev_lib.ingreso_id = t.id
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                {where_sql}
                ORDER BY t.id DESC
            """
            cur.execute(sql, params)
            rows = _fetchall_dicts(cur)
            rows = _merge_historico_rows(rows, _historico_deleted_rows(cur, historico_filter_ctx))

        wb = Workbook()
        ws = wb.active
        ws.title = "Historico ingresos"
        ws.append([
            "OS",
            "Cliente",
            "Tipo de equipo",
            "Marca",
            "Modelo",
            "Variante",
            "Estado",
            "N/S (serie)",
            "MG",
            "Fecha ingreso",
            "Fecha liberacion",
            "Fecha entrega",
            "Observación",
        ])

        def _dt_txt(value):
            if value is None:
                return "-"
            try:
                return value.strftime("%Y-%m-%d %H:%M")
            except Exception:
                return str(value)

        for r in rows:
            ws.append([
                os_label(r.get("id")),
                _ingreso_customer_label(r) or "-",
                r.get("tipo_equipo") or "-",
                r.get("marca") or "-",
                r.get("modelo") or "-",
                r.get("equipo_variante") or "-",
                r.get("estado") or "-",
                r.get("numero_serie") or "-",
                _mg_label(r),
                _dt_txt(r.get("fecha_ingreso")),
                _dt_txt(r.get("fecha_liberacion")),
                _dt_txt(r.get("fecha_entrega")),
                r.get("eliminacion_motivo") or "",
            ])

        try:
            widths = [10, 40, 20, 20, 24, 24, 18, 22, 14, 20, 20, 20, 80]
            for idx, width in enumerate(widths, start=1):
                ws.column_dimensions[chr(64 + idx)].width = width
        except Exception:
            pass

        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        fname = f"historico_ingresos_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
        resp = HttpResponse(
            bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f"attachment; filename=\"{fname}\""
        return resp


class IngresoAsignarTecnicoView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def patch(self, request, ingreso_id):
        require_roles(request, ["jefe", "admin"]) 
        tecnico_id = request.data.get("tecnico_id")
        if tecnico_id is None:
            return Response({"detail": "tecnico_id requerido"}, status=400)
        ok = q("SELECT id FROM users WHERE id=%s AND activo=true AND rol IN ('tecnico','jefe','jefe_suplente')",
                [tecnico_id], one=True)
        if not ok:
            return Response({"detail": "técnico inválido"}, status=400)
        try:
            prev = q("SELECT asignado_a FROM ingresos WHERE id=%s", [ingreso_id], one=True)
            logger.info(f"[AsignarTecnico] ingreso={ingreso_id} prev={prev and prev.get('asignado_a')} new={tecnico_id}")
        except Exception:
            logger.warning(f"[AsignarTecnico] ingreso={ingreso_id} could not read previous asignado_a")
        _set_audit_user(request)
        exec_void("UPDATE ingresos SET asignado_a=%s WHERE id=%s", [tecnico_id, ingreso_id])
        # Log post-commit para verificar persistencia real en DB
        try:
            def _post_commit_log():
                try:
                    row_pc = q("SELECT asignado_a FROM ingresos WHERE id=%s", [ingreso_id], one=True)
                    logger.info(f"[AsignarTecnico][post_commit] ingreso={ingreso_id} persisted={row_pc and row_pc.get('asignado_a')}")
                except Exception:
                    logger.warning(f"[AsignarTecnico][post_commit] ingreso={ingreso_id} read failed after commit")
            transaction.on_commit(_post_commit_log)
        except Exception:
            pass
        try:
            post = q("SELECT asignado_a FROM ingresos WHERE id=%s", [ingreso_id], one=True)
            logger.info(f"[AsignarTecnico] ingreso={ingreso_id} after_update={post and post.get('asignado_a')}")
        except Exception:
            logger.warning(f"[AsignarTecnico] ingreso={ingreso_id} could not read asignado_a after update")
        # Cerrar los pedidos de asignación del ingreso (aislado en savepoint): se
        # acepta el del técnico elegido y se cancelan los de los demás, que antes
        # quedaban pendientes para siempre y le escondían el botón al resto.
        if _table_exists("ingreso_assignment_requests"):
            try:
                with transaction.atomic():
                    exec_void(
                        """
                        UPDATE ingreso_assignment_requests
                           SET accepted_at = now(), status = 'aceptado'
                         WHERE ingreso_id = %s
                           AND usuario_id = %s
                           AND accepted_at IS NULL
                           AND canceled_at IS NULL
                        """,
                        [ingreso_id, tecnico_id],
                    )
                    exec_void(
                        """
                        UPDATE ingreso_assignment_requests
                           SET canceled_at = now(), status = 'cancelado'
                         WHERE ingreso_id = %s
                           AND usuario_id <> %s
                           AND accepted_at IS NULL
                           AND canceled_at IS NULL
                        """,
                        [ingreso_id, tecnico_id],
                    )
            except Exception:
                # si falla, no romper la transacción principal
                pass
        try:
            resolve_solicitud_asignacion(ingreso_id)
        except Exception:
            logger.exception(
                "No se pudo cerrar el aviso de solicitud de asignación",
                extra={"ingreso_id": ingreso_id},
            )
        try:
            who = q("SELECT COALESCE(nombre,'') AS nombre, COALESCE(email,'') AS email FROM users WHERE id=%s", [tecnico_id], one=True)
            nombre = (who and who.get("nombre")) or ""
            email_to = (who and who.get("email")) or ""
        except Exception:
            nombre = ""
            email_to = ""
        email_sent = False
        email_debug = {}
        try:
            notify_on_assign = getattr(settings, "ASSIGNMENT_NOTIFY_ON_ASSIGN", "1")
            notify_on_assign = str(notify_on_assign).lower() in ("1", "true", "yes")
        except Exception:
            notify_on_assign = True
        try:
            prev_id = prev and prev.get("asignado_a")
        except Exception:
            prev_id = None
        if notify_on_assign and (str(prev_id or "") != str(tecnico_id)):
            try:
                notify_ingreso_asignado(ingreso_id, tecnico_id)
            except Exception:
                pass
        if notify_on_assign and email_to and (str(prev_id or "") != str(tecnico_id)):
            try:
                info = q(
                    """
                    SELECT c.razon_social AS cliente,
                           COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                           COALESCE(b.nombre,'') AS marca,
                           COALESCE(m.nombre,'') AS modelo,
                           COALESCE(d.numero_serie,'') AS numero_serie,
                           COALESCE(d.numero_interno,'') AS numero_interno
                      FROM ingresos t
                      JOIN devices d   ON d.id = t.device_id
                      JOIN customers c ON c.id = d.customer_id
                      LEFT JOIN marcas b ON b.id = d.marca_id
                      LEFT JOIN models m ON m.id = d.model_id
                     WHERE t.id=%s
                    """,
                    [ingreso_id],
                    one=True,
                ) or {}
                os_txt = os_label(ingreso_id)
                cliente = info.get("cliente") or ""
                equipo = " | ".join([p for p in [info.get("tipo_equipo") or "", info.get("marca") or "", info.get("modelo") or ""] if p])
                ns = info.get("numero_serie") or ""
                subject = f"NS: {ns or os_txt} fue reasignado a vos"
                lines = [
                    f"Te asignaron/reasignaron un ingreso.",
                    f"OS: {os_txt}",
                    f"Cliente: {cliente}",
                    f"Equipo: {equipo or '-'}",
                    f"N/S: {ns or '-'}",
                ]
                try:
                    url = _frontend_url(request, f"/ingresos/{ingreso_id}") + "?tab=principal"
                    lines.append("")
                    lines.append(f"Abrir hoja: {url}")
                except Exception:
                    pass
                body = "\n".join(lines)
                try:
                    body = _email_append_footer_text(body)
                except Exception:
                    pass
                cc_list = getattr(settings, "ASSIGNMENT_NOTIFY_CC", []) or []
                if not isinstance(cc_list, (list, tuple)):
                    cc_list = [str(cc_list)] if cc_list else []
                recips = [email_to] + [x for x in cc_list if x]
                try:
                    email_debug.update({
                        "backend": getattr(settings, "EMAIL_BACKEND", None),
                        "host": getattr(settings, "EMAIL_HOST", None),
                        "port": getattr(settings, "EMAIL_PORT", None),
                        "use_tls": getattr(settings, "EMAIL_USE_TLS", None),
                        "use_ssl": getattr(settings, "EMAIL_USE_SSL", None),
                        "from": getattr(settings, "DEFAULT_FROM_EMAIL", None),
                        "recipients": list(recips),
                    })
                except Exception:
                    pass
                if recips:
                    try:
                        _sent, _dbg = _send_mail_with_fallback(subject, body, recips)
                        email_sent = bool(_sent)
                        try:
                            email_debug.update(_dbg or {})
                        except Exception:
                            pass
                        logger.info("assign_notify sent=%s ingreso_id=%s to=%s", email_sent, ingreso_id, email_to)
                    except Exception as e:
                        try:
                            email_debug["error"] = str(e)
                            email_debug["exception"] = e.__class__.__name__
                        except Exception:
                            pass
                        logger.exception("assign_notify failed", extra={"ingreso_id": ingreso_id, "to": email_to, "recipients": recips})
                        email_sent = False
                else:
                    logger.warning("assign_notify no recipient email configured", extra={"ingreso_id": ingreso_id})
            except Exception as e:
                try:
                    email_debug["error"] = str(e)
                    email_debug["exception"] = e.__class__.__name__
                except Exception:
                    pass
                email_sent = False
        try:
            tecnico_id_int = int(tecnico_id)
        except Exception:
            tecnico_id_int = tecnico_id
        resp = {"ok": True, "asignado_a": tecnico_id_int, "asignado_a_nombre": nombre, "email_sent": bool(email_sent)}
        try:
            if getattr(settings, "DEBUG", False) or _rol(request) in ("jefe", "admin", "supervisor"):
                resp["email_debug"] = email_debug
        except Exception:
            pass
        return Response(resp)


class IngresoSolicitarAsignacionView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        # Solo técnicos pueden solicitar
        require_roles(request, ["tecnico"])
        uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        if not uid:
            return Response({"detail": "Usuario inválido"}, status=400)

        # Si ya está asignado a este técnico, responder OK y no hacer nada
        try:
            cur = q("SELECT asignado_a FROM ingresos WHERE id=%s", [ingreso_id], one=True)
            if cur and int(cur.get("asignado_a") or 0) == int(uid):
                return Response({"ok": True, "already_assigned": True, "email_sent": False})
        except Exception:
            pass

        solicitud_row = None
        if _table_exists("ingreso_assignment_requests"):
            try:
                solicitud_row = q(
                    """
                    SELECT id, created_at
                      FROM ingreso_assignment_requests
                     WHERE ingreso_id = %s
                       AND usuario_id = %s
                       AND accepted_at IS NULL
                       AND canceled_at IS NULL
                     ORDER BY created_at DESC, id DESC
                     LIMIT 1
                    """,
                    [ingreso_id, uid],
                    one=True,
                )
            except Exception:
                solicitud_row = None

            if solicitud_row:
                return Response(
                    {
                        "ok": True,
                        "already_pending": True,
                        "solicitud_id": solicitud_row.get("id"),
                        "solicitud_fecha": solicitud_row.get("created_at"),
                        "email_sent": False,
                    }
                )

            try:
                with transaction.atomic():
                    solicitud_row = q(
                        """
                        INSERT INTO ingreso_assignment_requests(ingreso_id, usuario_id, status, created_at)
                        VALUES (%s, %s, 'solicitado', NOW())
                        RETURNING id, created_at
                        """,
                        [ingreso_id, uid],
                        one=True,
                    )
            except Exception:
                try:
                    solicitud_row = q(
                        """
                        SELECT id, created_at
                          FROM ingreso_assignment_requests
                         WHERE ingreso_id = %s
                           AND usuario_id = %s
                           AND accepted_at IS NULL
                           AND canceled_at IS NULL
                         ORDER BY created_at DESC, id DESC
                         LIMIT 1
                        """,
                        [ingreso_id, uid],
                        one=True,
                    )
                except Exception:
                    solicitud_row = None
                if solicitud_row:
                    return Response(
                        {
                            "ok": True,
                            "already_pending": True,
                            "solicitud_id": solicitud_row.get("id"),
                            "solicitud_fecha": solicitud_row.get("created_at"),
                            "email_sent": False,
                        }
                    )
                logger.exception(
                    "No se pudo registrar la solicitud de asignación",
                    extra={"ingreso_id": ingreso_id, "user_id": uid},
                )

        # Enviar notificación por email (reportar éxito/fracaso)
        email_sent = False
        notified = 0
        email_debug = {}
        try:
            # Datos del ingreso para el cuerpo
            info = q(
                """
                SELECT c.razon_social AS cliente,
                       COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                       COALESCE(b.nombre,'') AS marca,
                       COALESCE(m.nombre,'') AS modelo,
                       COALESCE(d.numero_serie,'') AS numero_serie,
                       COALESCE(d.numero_interno,'') AS numero_interno
                  FROM ingresos t
                  JOIN devices d   ON d.id = t.device_id
                  JOIN customers c ON c.id = d.customer_id
                  LEFT JOIN marcas b ON b.id = d.marca_id
                  LEFT JOIN models m ON m.id = d.model_id
                 WHERE t.id=%s
                """,
                [ingreso_id],
                one=True,
            ) or {}
            os_txt = os_label(ingreso_id)
            tech_name = getattr(request.user, "nombre", "")
            cliente = info.get("cliente") or ""
            equipo = " | ".join([p for p in [info.get("tipo_equipo") or "", info.get("marca") or "", info.get("modelo") or ""] if p])
            ns = info.get("numero_serie") or ""

            subject = f"Solicitud de asignacion {os_txt} - {cliente}"
            lines = [
                f"El tecnico {tech_name} solicita asignacion.",
                f"OS: {os_txt}",
                f"Cliente: {cliente}",
                f"Equipo: {equipo or '-'}",
                f"N/S: {ns or '-'}",
            ]
            qs_parts = ["tab=principal"]
            try:
                if uid:
                    qs_parts.append(f"tecnico_id={int(uid)}")
            except Exception:
                pass
            qs = ("?" + "&".join(qs_parts)) if qs_parts else ""
            try:
                url = _frontend_url(request, f"/ingresos/{ingreso_id}") + qs
                lines.append("")
                lines.append(f"Abrir hoja: {url}")
            except Exception:
                pass
            body = "\n".join(lines)
            try:
                body = _email_append_footer_text(body)
            except Exception:
                pass
            # La campana va a la jefatura suscripta al aviso y el correo a quien
            # tenga el Mail prendido (más la lista de entorno). Antes la lista de
            # correo ya filtrada era también la audiencia de la campana: apagar
            # el Mail dejaba a jefatura sin ningún aviso.
            recips = solicitud_asignacion_email_recipients(actor_user_id=uid)
            try:
                notified = notify_solicitud_asignacion(
                    ingreso_id,
                    uid,
                    tecnico_nombre=tech_name,
                    user_ids=solicitud_asignacion_user_ids(actor_user_id=uid),
                    solicitud_id=solicitud_row.get("id") if solicitud_row else None,
                )
            except Exception:
                logger.exception(
                    "No se pudo crear la notificación de solicitud de asignación",
                    extra={"ingreso_id": ingreso_id, "tecnico_id": uid},
                )

            # Debug info (sin secretos)
            try:
                email_debug.update({
                    "backend": getattr(settings, "EMAIL_BACKEND", None),
                    "host": getattr(settings, "EMAIL_HOST", None),
                    "port": getattr(settings, "EMAIL_PORT", None),
                    "use_tls": getattr(settings, "EMAIL_USE_TLS", None),
                    "use_ssl": getattr(settings, "EMAIL_USE_SSL", None),
                    "from": getattr(settings, "DEFAULT_FROM_EMAIL", None),
                    "recipients": list(recips),
                })
            except Exception:
                pass

            if recips:
                try:
                    sent, dbg = _send_mail_with_fallback(subject, body, recips)
                    email_sent = bool(sent)
                    try:
                        email_debug.update(dbg or {})
                    except Exception:
                        pass
                    logger.info(
                        "assignment_request_email sent=%s ingreso_id=%s user_id=%s recipients=%s backend=%s",
                        email_sent, ingreso_id, uid, recips, getattr(settings, "EMAIL_BACKEND", ""),
                    )
                except Exception as e:
                    try:
                        email_debug["error"] = str(e)
                        email_debug["exception"] = e.__class__.__name__
                    except Exception:
                        pass
                    logger.exception(
                        "assignment_request_email failed",
                        extra={"ingreso_id": ingreso_id, "user_id": uid, "recipients": recips},
                    )
                    email_sent = False
            else:
                logger.warning(
                    "assignment_request_email no recipients configured",
                    extra={"ingreso_id": ingreso_id, "user_id": uid},
                )
        except Exception as e:
            try:
                email_debug["error"] = str(e)
                email_debug["exception"] = e.__class__.__name__
            except Exception:
                pass
            email_sent = False

        # Solo exponer debug a roles altos o en DEBUG
        resp = {
            "ok": True,
            "email_sent": bool(email_sent),
            # Avisos de campana/push creados: el front los cuenta como aviso aunque
            # jefatura tenga el Mail apagado.
            "notified": int(notified or 0),
            "solicitud_id": solicitud_row.get("id") if solicitud_row else None,
            "solicitud_fecha": solicitud_row.get("created_at") if solicitud_row else None,
        }
        try:
            if getattr(settings, "DEBUG", False) or _rol(request) in ("jefe", "admin", "supervisor"):
                resp["email_debug"] = email_debug
        except Exception:
            pass
        return Response(resp)


class IngresoSolicitarBajaView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        _set_audit_user(request)
        if user_has_permission(request, "action.ingreso.baja_alta"):
            raise PermissionDenied("Tu usuario ya tiene permiso para dar baja directa")

        require_any_permission(
            request,
            [
                "page.ingresos_history",
                "page.work_queues",
                "page.budget_queues",
                "page.logistics",
                "action.ingreso.edit_basics",
                "action.ingreso.edit_diagnosis",
                "action.ingreso.edit_location",
                "action.ingreso.edit_delivery",
                "action.ingreso.manage_derivations",
                "action.ingreso.repair_transitions",
                "action.presupuesto.manage",
            ],
            message="No autorizado para solicitar baja",
        )

        uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        if not uid:
            return Response({"detail": "Usuario invalido"}, status=400)

        payload = request.data or {}
        motivo = str(payload.get("motivo") or "").strip()
        if not motivo:
            return Response({"detail": "El motivo es obligatorio"}, status=400)

        row = _fetch_ingreso_device_mg_context(ingreso_id)
        if not row:
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        estado_actual = (row.get("estado") or "").lower()
        if estado_actual == "baja":
            return Response({"ok": True, "already_baja": True, "email_sent": False})
        if not row.get("es_propietario_mg"):
            return Response({"detail": "Solo se puede solicitar BAJA para equipos propios MG."}, status=400)

        pending = _pending_baja_request(ingreso_id)
        if pending:
            return Response(
                {
                    "ok": True,
                    "already_pending": True,
                    "email_sent": False,
                    "solicitud_id": pending.get("id"),
                    "baja_solicitada_fecha": pending.get("created_at"),
                }
            )

        req_row = None
        try:
            req_row = q(
                """
                INSERT INTO ingreso_baja_requests (ingreso_id, usuario_id, motivo, created_at)
                VALUES (%s, %s, %s, NOW())
                RETURNING id, created_at
                """,
                [ingreso_id, uid, motivo],
                one=True,
            )
        except Exception:
            pending = _pending_baja_request(ingreso_id)
            if pending:
                return Response(
                    {
                        "ok": True,
                        "already_pending": True,
                        "email_sent": False,
                        "solicitud_id": pending.get("id"),
                        "baja_solicitada_fecha": pending.get("created_at"),
                    }
                )
            logger.exception("baja_request_insert failed", extra={"ingreso_id": ingreso_id, "user_id": uid})
            return Response({"detail": "No se pudo registrar la solicitud de baja"}, status=500)

        info = q(
            """
            SELECT c.razon_social AS cliente,
                   COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                   COALESCE(b.nombre,'') AS marca,
                   COALESCE(m.nombre,'') AS modelo,
                   COALESCE(d.numero_serie,'') AS numero_serie,
                   COALESCE(d.numero_interno,'') AS numero_interno
              FROM ingresos t
              JOIN devices d ON d.id = t.device_id
              JOIN customers c ON c.id = d.customer_id
              LEFT JOIN marcas b ON b.id = d.marca_id
              LEFT JOIN models m ON m.id = d.model_id
             WHERE t.id=%s
            """,
            [ingreso_id],
            one=True,
        ) or {}

        os_txt = os_label(ingreso_id)
        cliente = info.get("cliente") or ""
        equipo = " | ".join(
            [
                p
                for p in [info.get("tipo_equipo") or "", info.get("marca") or "", info.get("modelo") or ""]
                if p
            ]
        )
        numero_serie = info.get("numero_serie") or ""
        numero_interno = info.get("numero_interno") or ""
        actor_nombre = getattr(request.user, "nombre", None) or getattr(request.user, "username", "") or ""
        actor_email = getattr(request.user, "email", "") or ""
        actor_line = actor_nombre or "-"
        if actor_email:
            actor_line = f"{actor_line} ({actor_email})" if actor_line else actor_email

        subject = f"Solicitud de BAJA {os_txt} - {cliente or 'Sin cliente'}"
        lines = [
            "Se solicitó dar de baja un equipo en NEXORA.",
            "",
            "Detalle del ingreso:",
            f"- OS: {os_txt}",
            f"- Cliente: {cliente or '-'}",
            f"- Equipo: {equipo or '-'}",
            f"- Número de serie: {numero_serie or '-'}",
            f"- Número interno: {numero_interno or '-'}",
            f"- Solicitado por: {actor_line or '-'}",
            "",
            "Motivo de la solicitud:",
            motivo,
        ]
        try:
            url = _frontend_url(request, f"/ingresos/{ingreso_id}") + "?tab=principal"
            lines.append("")
            lines.append(f"Hoja de servicio: {url}")
        except Exception:
            pass
        body = "\n".join(lines)
        try:
            body = _email_append_footer_text(body)
        except Exception:
            pass

        email_sent = False
        email_debug = {}
        recips = []
        try:
            rec_rows = q(
                """
                SELECT DISTINCT LOWER(TRIM(CAST(email AS TEXT))) AS email
                  FROM users
                 WHERE COALESCE(activo, TRUE) = TRUE
                   AND LOWER(COALESCE(rol, '')) IN ('jefe', 'jefe_suplente')
                   AND NULLIF(TRIM(CAST(email AS TEXT)), '') IS NOT NULL
                 ORDER BY 1
                """
            ) or []
            campana = [str(r.get("email") or "").strip() for r in rec_rows if str(r.get("email") or "").strip()]
            # La campana usa la lista completa de jefes; el correo, sólo los que
            # dejaron prendido el canal Mail en el centro de notificaciones.
            recips = filter_email_recipients_by_preference("solicitud_baja", campana)
            try:
                notify_solicitud_baja(ingreso_id, solicitante_id=uid, motivo=motivo, emails=campana)
            except Exception:
                logger.exception(
                    "No se pudo crear la notificación de solicitud de baja",
                    extra={"ingreso_id": ingreso_id, "solicitante_id": uid},
                )
            try:
                email_debug.update(
                    {
                        "backend": getattr(settings, "EMAIL_BACKEND", None),
                        "host": getattr(settings, "EMAIL_HOST", None),
                        "port": getattr(settings, "EMAIL_PORT", None),
                        "use_tls": getattr(settings, "EMAIL_USE_TLS", None),
                        "use_ssl": getattr(settings, "EMAIL_USE_SSL", None),
                        "from": getattr(settings, "DEFAULT_FROM_EMAIL", None),
                        "recipients": list(recips),
                    }
                )
            except Exception:
                pass
            if recips:
                try:
                    sent, dbg = _send_mail_with_fallback(subject, body, recips)
                    email_sent = bool(sent)
                    try:
                        email_debug.update(dbg or {})
                    except Exception:
                        pass
                    logger.info(
                        "baja_request_email sent=%s ingreso_id=%s user_id=%s recipients=%s backend=%s",
                        email_sent,
                        ingreso_id,
                        uid,
                        recips,
                        getattr(settings, "EMAIL_BACKEND", ""),
                    )
                except Exception as e:
                    try:
                        email_debug["error"] = str(e)
                        email_debug["exception"] = e.__class__.__name__
                    except Exception:
                        pass
                    logger.exception(
                        "baja_request_email failed",
                        extra={"ingreso_id": ingreso_id, "user_id": uid, "recipients": recips},
                    )
                    email_sent = False
            else:
                logger.warning(
                    "baja_request_email no recipients configured",
                    extra={"ingreso_id": ingreso_id, "user_id": uid},
                )
        except Exception as e:
            try:
                email_debug["error"] = str(e)
                email_debug["exception"] = e.__class__.__name__
            except Exception:
                pass
            email_sent = False

        resp = {
            "ok": True,
            "email_sent": bool(email_sent),
            "solicitud_id": req_row.get("id") if isinstance(req_row, dict) else None,
            "baja_solicitada_fecha": req_row.get("created_at") if isinstance(req_row, dict) else None,
        }
        try:
            if getattr(settings, "DEBUG", False) or _rol(request) in ("jefe", "admin", "supervisor"):
                resp["email_debug"] = email_debug
        except Exception:
            pass
        return Response(resp)


class IngresoSolicitarBajaRechazarView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe", "jefe_veedor", "admin", "recepcion"])
        _set_audit_user(request)

        row = q("SELECT id FROM ingresos WHERE id=%s", [ingreso_id], one=True)
        if not row:
            return Response({"detail": "Ingreso no encontrado"}, status=404)

        try:
            canceled = q(
                """
                UPDATE ingreso_baja_requests
                   SET canceled_at = COALESCE(canceled_at, NOW())
                 WHERE ingreso_id = %s
                   AND accepted_at IS NULL
                   AND canceled_at IS NULL
                RETURNING id
                """,
                [ingreso_id],
                one=True,
            )
        except Exception:
            logger.exception(
                "baja_request_reject failed",
                extra={"ingreso_id": ingreso_id, "user_id": getattr(request, "user_id", None)},
            )
            return Response({"detail": "No se pudo rechazar la solicitud de baja"}, status=500)

        if canceled:
            return Response({"ok": True, "rejected": True})
        return Response({"ok": True, "already_processed": True})


class PendientesGeneralView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        require_roles(request, ["jefe", "admin", "jefe_veedor", "tecnico"])
        tecnico_raw = (request.GET.get("tecnico_id") or "").strip()
        tecnico_filter = tecnico_raw.lower()
        mg_list_sql = _mg_list_select_sql()
        owner_list_sql = _owner_list_select_sql()

        with connection.cursor() as cur:
            _set_audit_user(request)

            sql = """
                SELECT t.id,
                       d.id AS device_id,
                       c.id AS customer_id,
                       t.estado,
                       t.presupuesto_estado,
                       t.motivo,
                       c.razon_social,
                       {owner_list_sql}
                       d.numero_serie,
                       COALESCE(d.numero_interno,'') AS numero_interno,
                       {mg_list_sql}
                       COALESCE(b.nombre,'') AS marca,
                       COALESCE(m.nombre,'') AS modelo,
                       COALESCE(m.tipo_equipo,'') AS tipo_equipo,
                       COALESCE(NULLIF(t.equipo_variante,''), NULLIF(d.variante,''), NULLIF(m.variante,'')) AS equipo_variante,
                       t.asignado_a,
                       COALESCE(u.nombre,'') AS tecnico,
                       COALESCE(u.nombre,'') AS tecnico_nombre,
                       COALESCE(u.nombre,'') AS asignado_a_nombre,
                       t.fecha_ingreso,
                       CASE WHEN ed.estado = 'devuelto' THEN true ELSE false END AS derivado_devuelto
                FROM ingresos t
                JOIN devices d   ON d.id = t.device_id
                JOIN customers c ON c.id = d.customer_id
                LEFT JOIN marcas b ON b.id = d.marca_id
                LEFT JOIN models m ON m.id = d.model_id
                LEFT JOIN users u ON u.id = t.asignado_a
                LEFT JOIN locations loc ON loc.id = t.ubicacion_id
                LEFT JOIN (
                  SELECT e.*, ROW_NUMBER() OVER (
                    PARTITION BY e.ingreso_id ORDER BY e.fecha_deriv DESC, e.id DESC
                  ) AS rn
                  FROM equipos_derivados e
                ) ed ON ed.ingreso_id = t.id AND ed.rn = 1
                WHERE LOWER(loc.nombre) = LOWER(%s)
                  AND t.estado NOT IN ('liberado','entregado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
            """.format(mg_list_sql=mg_list_sql, owner_list_sql=owner_list_sql)
            params = ["taller"]
            if tecnico_raw.isdigit():
                sql += " AND t.asignado_a = %s"
                params.append(int(tecnico_raw))
            elif tecnico_filter in {"sin_asignar", "unassigned", "none", "null"}:
                sql += " AND t.asignado_a IS NULL"

            sql += """
                ORDER BY
                  (CASE WHEN ed.estado = 'devuelto' THEN 1 ELSE 0 END) DESC,
                  (t.motivo = 'urgente control') DESC,
                  t.fecha_ingreso ASC
            """
            cur.execute(sql, params)
            rows = _fetchall_dicts(cur)

        mg_owner_id = _mg_owner_customer_id()
        for row in rows:
            row.update(resolve_mg_flags(row, mg_owner_id))
            row["es_cliente_mg_owner"] = bool(row.get("es_propietario_mg"))

        attach_active_derivations(rows)
        return Response(IngresoListItemSerializer(rows, many=True).data)


def _resolve_location_history_values(rows):
    if not rows:
        return rows

    def _int_history_value(value):
        if value is None:
            return None
        raw = str(value).strip()
        if not raw or raw.lower() in {"null", "none"}:
            return None
        try:
            return int(raw)
        except (TypeError, ValueError):
            return None

    ids = set()
    for row in rows:
        if (row.get("column_name") or "").strip().lower() != "ubicacion_id":
            continue
        for key in ("old_value", "new_value"):
            loc_id = _int_history_value(row.get(key))
            if loc_id is not None:
                ids.add(loc_id)

    if not ids:
        return rows

    try:
        loc_rows = q(
            """
            SELECT id, nombre
              FROM locations
             WHERE id = ANY(%s)
            """,
            [list(ids)],
        ) or []
        names = {int(r["id"]): r.get("nombre") for r in loc_rows if r.get("id") is not None}
    except Exception:
        return rows

    for row in rows:
        if (row.get("column_name") or "").strip().lower() != "ubicacion_id":
            continue
        for key in ("old_value", "new_value"):
            loc_id = _int_history_value(row.get(key))
            if loc_id is not None and names.get(loc_id):
                row[key] = names[loc_id]
    return rows


def _resolve_assigned_history_values(rows):
    if not rows:
        return rows

    def _user_id(value):
        if value is None:
            return None
        raw = str(value).strip()
        if not raw or raw.lower() in {"null", "none"}:
            return None
        try:
            return int(raw)
        except (TypeError, ValueError):
            return None

    ids = set()
    for row in rows:
        if (row.get("column_name") or "").strip().lower() != "asignado_a":
            continue
        for key in ("old_value", "new_value"):
            uid = _user_id(row.get(key))
            if uid is not None:
                ids.add(uid)

    if not ids:
        return rows

    try:
        user_rows = q(
            """
            SELECT id, nombre
              FROM users
             WHERE id = ANY(%s)
            """,
            [list(ids)],
        ) or []
        names = {int(r["id"]): r.get("nombre") for r in user_rows if r.get("id") is not None}
    except Exception:
        return rows

    for row in rows:
        if (row.get("column_name") or "").strip().lower() != "asignado_a":
            continue
        for key in ("old_value", "new_value"):
            uid = _user_id(row.get(key))
            if uid is not None and names.get(uid):
                row[key] = names[uid]
    return rows


def _history_ts(row):
    ts = row.get("ts") if row else None
    if isinstance(ts, datetime):
        return ts
    if isinstance(ts, str):
        try:
            return parse_datetime(ts)
        except Exception:
            return None
    return None


def _history_ts_sort_value(row):
    ts = _history_ts(row)
    if not ts:
        return 0.0
    try:
        return ts.timestamp()
    except Exception:
        return 0.0


def _history_id_sort_value(row):
    try:
        return int(row.get("_id") or row.get("id") or 0)
    except Exception:
        return 0


def _history_value_key(value):
    if value is None:
        return ""
    return str(value).strip()


def _compact_history_rows(rows, max_gap=timedelta(minutes=3)):
    """Une guardados consecutivos del mismo campo, típicos de escritura por tecla."""
    if not rows:
        return rows

    max_gap_seconds = max_gap.total_seconds()
    ordered = sorted(
        enumerate(rows),
        key=lambda item: (
            _history_ts_sort_value(item[1]),
            _history_id_sort_value(item[1]),
            item[0],
        ),
    )
    groups = []
    last_by_key = {}

    for _, original_row in ordered:
        row = dict(original_row)
        row_key = (
            (row.get("table_name") or "").strip().lower(),
            row.get("record_id"),
            (row.get("column_name") or "").strip().lower(),
            row.get("user_id"),
            row.get("user_role"),
        )
        row_ts = _history_ts_sort_value(row)
        last = last_by_key.get(row_key)
        can_join = False
        if last and last["key"] == row_key:
            close_enough = row_ts and last["last_ts"] and (row_ts - last["last_ts"]) <= max_gap_seconds
            same_chain = _history_value_key(last["row"].get("new_value")) == _history_value_key(row.get("old_value"))
            can_join = bool(close_enough and same_chain)

        if can_join:
            last["row"]["new_value"] = row.get("new_value")
            last["row"]["ts"] = row.get("ts")
            if row.get("_id") is not None:
                last["row"]["_id"] = row.get("_id")
            last["row"]["_compact_count"] = int(last["row"].get("_compact_count") or 1) + 1
            last["last_ts"] = row_ts
        else:
            row["_compact_count"] = 1
            group = {"key": row_key, "last_ts": row_ts, "row": row}
            groups.append(group)
            last_by_key[row_key] = group

    compacted = [item["row"] for item in groups]
    compacted.sort(
        key=lambda row: (_history_ts_sort_value(row), _history_id_sort_value(row)),
        reverse=True,
    )
    return compacted


def _drop_mirrored_device_history_rows(rows, max_gap=timedelta(seconds=5)):
    if not rows:
        return rows

    mirrored_fields = {
        "alquilado",
        "alquiler_a",
        "ubicacion_id",
        "faja_garantia",
        "propietario_nombre",
        "propietario_contacto",
        "propietario_doc",
    }
    max_gap_seconds = max_gap.total_seconds()
    ingreso_rows = [
        r for r in rows
        if (r.get("table_name") or "").strip().lower() == "ingresos"
    ]
    filtered = []
    for row in rows:
        table_name = (row.get("table_name") or "").strip().lower()
        column_name = (row.get("column_name") or "").strip().lower()
        if table_name != "devices" or column_name not in mirrored_fields:
            filtered.append(row)
            continue

        row_ts = _history_ts_sort_value(row)
        duplicated = any(
            (other.get("column_name") or "").strip().lower() == column_name
            and other.get("user_id") == row.get("user_id")
            and _history_value_key(other.get("old_value")) == _history_value_key(row.get("old_value"))
            and _history_value_key(other.get("new_value")) == _history_value_key(row.get("new_value"))
            and abs(_history_ts_sort_value(other) - row_ts) <= max_gap_seconds
            for other in ingreso_rows
        )
        if not duplicated:
            filtered.append(row)
    return filtered


class IngresoHistorialView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request, ingreso_id: int):
        require_roles(request, ["jefe", "jefe_veedor", "admin"])
        # Flag opcional para incluir la auditoría HTTP (audit_log) en la respuesta
        def _param_truthy(name: str) -> bool:
            v = (request.query_params.get(name) or "").strip().lower()
            return v in {"1", "true", "yes", "on"}
        include_audit = _param_truthy("include_audit")
        if connection.vendor == "postgresql":
            # change_log (si existe) + audit_log (opcional)
            rows = []
            try:
                rows = q(
                    """
                      SELECT cl.id AS _id, cl.ts, cl.user_id, cl.user_role, COALESCE(u.nombre,'') AS user_nombre, cl.table_name, cl.record_id, cl.column_name, cl.old_value, cl.new_value FROM audit.change_log cl LEFT JOIN users u ON u.id = cl.user_id
                      WHERE ingreso_id = %s
                      ORDER BY cl.ts DESC, cl.id DESC
                    """,
                    [ingreso_id]
                ) or []
            except Exception:
                rows = []
            if include_audit:
                # Complementar con auditoría HTTP (audit_log), pero sin expandir por clave
                pat1 = f"/api/ingresos/{ingreso_id}/%"
                pat2 = f"/api/ingresos/{ingreso_id}/"
                pat3 = f"/api/quotes/{ingreso_id}/%"
                pat4 = f"/api/quotes/{ingreso_id}/"
                al_rows = q(
                    """
                    SELECT al.id AS _id, al.ts, al.user_id, al.role AS user_role, COALESCE(u.nombre,'') AS user_nombre, al.method, al.path, al.body
                      FROM audit_log al LEFT JOIN users u ON u.id = al.user_id
                     WHERE path LIKE %s OR path = %s OR path LIKE %s OR path = %s
                     ORDER BY al.ts DESC, al.id DESC
                    """,
                    [pat1, pat2, pat3, pat4]
                ) or []
                out = []
                for r in (al_rows or []):
                    path = (r.get("path") or "").lower()
                    method = (r.get("method") or "").upper()
                    table_name = "ingresos"
                    record_id = ingreso_id
                    if "/accesorios/" in path:
                        table_name = "ingreso_accesorios"
                    elif "/fotos/" in path:
                        table_name = "ingreso_media"
                    elif "/quotes/" in path or "/presupuestos/" in path:
                        table_name = "quotes"
                    body = r.get("body")
                    try:
                        if isinstance(body, (dict, list)):
                            body_str = json.dumps(body, ensure_ascii=False)
                        elif isinstance(body, str):
                            body_str = body
                        else:
                            body_str = ""
                    except Exception:
                        body_str = ""
                    out.append({
                        "ts": r.get("ts"),
                        "user_id": r.get("user_id"),
                        "user_role": r.get("user_role"),
                        "user_nombre": r.get("user_nombre"),
                        "table_name": table_name,
                        "record_id": record_id,
                        "column_name": f"{method} {r.get('path')}",
                        "old_value": None,
                        "new_value": (body_str or None) and body_str[:512],
                    })
                rows = (rows or []) + out
            rows = _compact_history_rows(rows)
            rows = _drop_mirrored_device_history_rows(rows)
            rows = _resolve_location_history_values(rows)
            rows = _resolve_assigned_history_values(rows)
            rows.sort(key=lambda x: (_history_ts_sort_value(x), _history_id_sort_value(x)), reverse=True)
        else:
            # 1) Cambios de estado (ingreso_events)
            ev_rows = q(
                """
                SELECT 
                  e.id AS _id,
                  e.ts AS ts,
                  e.usuario_id AS user_id,
                  u.rol AS user_role, COALESCE(u.nombre,'') AS user_nombre,
                  'ingresos' AS table_name,
                  e.ingreso_id AS record_id,
                  'estado' AS column_name,
                  e.de_estado AS old_value,
                  e.a_estado AS new_value
                FROM ingreso_events e
                LEFT JOIN users u ON u.id = e.usuario_id
                WHERE e.ingreso_id = %s
                ORDER BY e.ts DESC, e.id DESC
                """,
                [ingreso_id]
            ) or []

            # 2) Cambios por HTTP (audit_log): opcional (?include_audit=1)
            pat1 = f"/api/ingresos/{ingreso_id}/%"
            pat2 = f"/api/ingresos/{ingreso_id}/"
            pat3 = f"/api/quotes/{ingreso_id}/%"
            pat4 = f"/api/quotes/{ingreso_id}/"
            out = list(ev_rows)
            if include_audit:
                al_rows = q(
                    """
                    SELECT al.id AS _id, al.ts, al.user_id, al.role AS user_role, COALESCE(u.nombre,'') AS user_nombre, al.method, al.path, al.body FROM audit_log al LEFT JOIN users u ON u.id = al.user_id
                     WHERE path LIKE %s OR path = %s OR path LIKE %s OR path = %s
                     ORDER BY al.ts DESC, al.id DESC
                    """,
                    [pat1, pat2, pat3, pat4]
                ) or []
                for r in (al_rows or []):
                    path = (r.get("path") or "").lower()
                    method = (r.get("method") or "").upper()
                    table_name = "ingresos"
                    record_id = ingreso_id
                    # Heurística por ruta
                    if "/accesorios/" in path:
                        table_name = "ingreso_accesorios"
                    elif "/fotos/" in path:
                        table_name = "ingreso_media"
                    elif "/quotes/" in path or "/presupuestos/" in path:
                        table_name = "quotes"
                    body = r.get("body")
                    try:
                        if isinstance(body, (dict, list)):
                            body_str = json.dumps(body, ensure_ascii=False)
                        elif isinstance(body, str):
                            body_str = body
                        else:
                            body_str = ""
                    except Exception:
                        body_str = ""
                    out.append({
                        "ts": r.get("ts"),
                        "user_id": r.get("user_id"),
                        "user_role": r.get("user_role"),
                        "user_nombre": r.get("user_nombre"),
                        "table_name": table_name,
                        "record_id": record_id,
                        "column_name": f"{method} {r.get('path')}",
                        "old_value": None,
                        "new_value": (body_str or None) and body_str[:512],
                    })

            # Orden final por fecha desc y sin modificar formato esperado
            out.sort(key=lambda x: (_history_ts_sort_value(x), _history_id_sort_value(x)), reverse=True)
            rows = out
        from ..permissions import user_has_permission
        commercial = any(user_has_permission(request.user, code) for code in ('action.presupuesto.manage','action.presupuesto.view_costs','page.budget_queues'))
        if not commercial:
            rows = [row for row in rows if row.get('table_name') not in ('quotes','quote_items','technical_labor_rates','technical_billing_config')]
        elif not user_has_permission(request.user, 'action.presupuesto.view_costs'):
            rows = [row for row in rows if 'costo' not in str(row.get('column_name','')).lower()]
        return Response(rows)


class CerrarReparacionView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe"])
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        permite_reparacion_sql = (
            "COALESCE(permite_reparacion, TRUE) AS permite_reparacion"
            if has_permite_reparacion
            else "TRUE AS permite_reparacion"
        )
        row_repair = q(
            f"""
            SELECT
              asignado_a,
              motivo,
              {permite_reparacion_sql}
            FROM ingresos
            WHERE id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row_repair:
            raise ValidationError("Ingreso no encontrado")
        if _is_cotizacion_bloqueada_para_reparar(row_repair.get("motivo"), row_repair.get("permite_reparacion")):
            raise ValidationError("Ingreso en cotización: primero habilite la reparación.")
        r = (request.data or {}).get("resolucion")
        if r not in ("reparado","no_reparado","no_se_encontro_falla","presupuesto_rechazado","cambio"):
            return Response({"detail": "resolución inválida"}, status=400)

        serial_cambio = (request.data or {}).get("serial_cambio")
        if r == "cambio":
            if not (serial_cambio or "").strip():
                return Response({"detail": "serial_cambio requerido para resolucion 'cambio'"}, status=400)
            serial_cambio = (serial_cambio or "").strip()

        presupuesto_rechazado_cobro_neto = None
        presupuesto_rechazado_quote_id = None
        if r == "presupuesto_rechazado":
            presupuesto_rechazado_cobro_neto = _parse_rejected_budget_charge_or_raise(
                (request.data or {}).get("presupuesto_rechazado_cobro_neto"),
                required=False,
            )
            if ensure_rejected_budget_charge_schema():
                stored_fields = get_stored_rejected_budget_fields(ingreso_id)
                if presupuesto_rechazado_cobro_neto is None:
                    presupuesto_rechazado_cobro_neto = stored_fields.get("presupuesto_rechazado_cobro_neto")
                latest_rejected_quote = get_rejected_quote_summary(ingreso_id)
                presupuesto_rechazado_quote_id = (
                    latest_rejected_quote.get("quote_id") if latest_rejected_quote else None
                )
            if presupuesto_rechazado_cobro_neto is None:
                return Response(
                    {"detail": "Debe informar el cobro neto del presupuesto rechazado."},
                    status=400,
                )

        with connection.cursor() as cur:
            if r == 'no_se_encontro_falla':
                from .. import technical_interventions as technical
                technical.mark_no_fault(ingreso_id, request.user)
            _set_audit_user(request)
            # Detectar si existe columna serial_cambio
            cur.execute(
                """
                SELECT 1 FROM information_schema.columns
                 WHERE table_name = 'ingresos' AND column_name = 'serial_cambio'
                   AND table_schema = ANY(current_schemas(true))
                 LIMIT 1
                """
            )
            has_serial_col = cur.fetchone() is not None
            has_rejected_budget_cols = has_rejected_budget_charge_schema()
            if r == "cambio" and has_serial_col:
                cur.execute(
                    """
                    UPDATE ingresos
                       SET resolucion = %s,
                           serial_cambio = %s
                     WHERE id = %s
                    """,
                    [r, serial_cambio, ingreso_id],
                )
            elif r == "presupuesto_rechazado" and has_rejected_budget_cols:
                cur.execute(
                    """
                    UPDATE ingresos
                       SET resolucion = %s,
                           presupuesto_rechazado_cobro_neto = %s,
                           presupuesto_rechazado_quote_id = %s
                     WHERE id = %s
                    """,
                    [
                        r,
                        presupuesto_rechazado_cobro_neto,
                        presupuesto_rechazado_quote_id,
                        ingreso_id,
                    ],
                )
            else:
                cur.execute(
                    """
                    UPDATE ingresos
                       SET resolucion = %s
                     WHERE id = %s
                    """,
                    [r, ingreso_id],
                )
        return Response({"ok": True})


# El tipo de equipo forma parte de la identidad del equipo, no es sólo una etiqueta: un CPAP
# y el calentador humidificador que lo acompaña comparten el número de serie (el humidificador
# no trae uno propio), y sin el tipo son indistinguibles. Ver
# apply_devices_ns_tipo_unique_schema y su par apply_models_tipo_equipo_unique_schema.
def _tipo_equipo_de_modelo(modelo_id) -> str:
    try:
        model_id = int(modelo_id)
    except (TypeError, ValueError):
        return ""
    row = q("SELECT COALESCE(tipo_equipo,'') AS tipo_equipo FROM models WHERE id=%s", [model_id], one=True)
    return ((row or {}).get("tipo_equipo") or "").strip()


def _ingreso_duplicado_hoy_response(dup, *, numero_serie: str, numero_interno: str, tipo_equipo: str):
    """La OS que choca es de hoy: esto no es un conflicto, es el mismo remito cargado dos veces.

    Con una OS vieja todavía abierta el ingreso se crea igual y el choque queda anotado, porque
    el cliente está esperando el remito. Acá no: no se crea nada y se nombra la OS que ya
    existe, para que recepción la abra en vez de volver a cargar el equipo. El lote tampoco
    puede saltearlo con ``_allow_open_duplicate``, que es justo por donde entraban las OS
    gemelas del mismo día.
    """

    existing_id = dup["id"]
    fecha_val = dup.get("fecha_ingreso") or dup.get("fecha_creacion")
    try:
        fecha_iso = fecha_val.isoformat()
    except AttributeError:
        fecha_iso = fecha_val
    identificador = (numero_serie or "").strip() or (numero_interno or "").strip()
    sujeto = f"El equipo {identificador}" if identificador else "Este equipo"
    return Response(
        {
            "detail": (
                f"{sujeto} ya se ingresó hoy en la OS {os_label(existing_id)}. "
                "No se puede ingresar dos veces el mismo día: abrí esa orden, o dala de baja "
                "si quedó mal cargada, antes de emitir un remito nuevo."
            ),
            "conflict_type": "INGRESO_DUPLICADO_HOY",
            "payload": {
                "ingreso_id": existing_id,
                "os": os_label(existing_id),
                "fecha_ingreso": fecha_iso,
                "numero_serie": (numero_serie or "").strip(),
                "numero_interno": (numero_interno or "").strip(),
                "tipo_equipo": (tipo_equipo or "").strip(),
            },
        },
        status=409,
    )


class NuevoIngresoView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        require_any_permission(request, ["action.ingreso.create", "page.new_ingreso"])
        # Best-effort: ensure PK sequences are aligned to prevent duplicate key
        # errors due to out-of-sync sequences after legacy imports.
        try:
            from .helpers_impl import repair_pk_sequence
            repair_pk_sequence('devices')
            repair_pk_sequence('ingresos')
        except Exception:
            pass
        data = request.data or {}
        cliente = data.get("cliente") or {}
        equipo = data.get("equipo") or {}
        bejerman_sale = data.get("bejerman_sale") or {}

        try:
            selected_company = require_company(data.get("empresa_bejerman") or data.get("empresa_facturar") or "SEPID")
        except ValueError:
            return Response({"detail": "Empresa Bejerman no válida"}, status=400)
        empresa_bejerman = selected_company.key
        empresa_facturar = selected_company.branding_key

        motivo_raw = (data.get("motivo") or "").strip()
        if not motivo_raw:
            return Response({"detail": "motivo requerido"}, status=400)

        motivo_label_raw = _map_motivo_to_db_label(motivo_raw)
        if not motivo_label_raw:
            valid_motivos = _get_motivo_enum_values()
            return Response({"detail": "motivo inválido", "valid_values": valid_motivos}, status=400)

        raw_vals = _get_motivo_enum_values_raw() or []
        if raw_vals:
            target = None
            norm_target = _norm_txt(motivo_label_raw)
            for rv in raw_vals:
                if _norm_txt(rv) == norm_target:
                    target = rv
                    break
            if not target:
                target = next((x for x in raw_vals if _norm_txt(x) == _norm_txt('otros')), raw_vals[0])
            motivo = target
        else:
            motivo = motivo_label_raw
        permite_reparacion_default = not _motivo_is_cotizacion_equipo(motivo)

        # El número interno no distingue mayúsculas (regla del 14/10/2026): "mg 7359" es el
        # MG 7359. Guardarlo tal cual se tipeó dejó la OS 29098 en una ficha nueva, al lado de
        # la que ya tenía ese MG y su N/S.
        numero_interno = normalize_mg_input(equipo.get("numero_interno"))
        has_mg_schema = _has_mg_schema()
        mg_owner_id = _mg_owner_customer_id()

        ubicacion_id = data.get("ubicacion_id")
        if not ubicacion_id:
            t = q("SELECT id FROM locations WHERE LOWER(nombre)=LOWER(%s) LIMIT 1", ["taller"], one=True)
            if not t:
                return Response({"detail": "No se encontró la ubicación 'Taller' en el catálogo. Creala en 'locations'."}, status=400)
            ubicacion_id = t["id"]

        informe_preliminar = (data.get("informe_preliminar") or "").strip()
        accesorios_text = (data.get("accesorios") or "").strip()
        comentarios_text = (data.get("comentarios") or "").strip() or None
        accesorios_items = data.get("accesorios_items") or []

        remito_ingreso = (data.get("remito_ingreso") or "").strip()
        fecha_ingreso_dt = None
        _fi_raw = data.get("fecha_ingreso")
        if _fi_raw is not None and str(_fi_raw).strip() != "":
            _dt = _parse_datetime_or_date(_fi_raw)
            if not _dt:
                return Response({"detail": "fecha_ingreso inválida (use YYYY-MM-DD o DD/MM/AAAA)"}, status=400)
            if timezone.is_naive(_dt):
                _dt = timezone.make_aware(_dt, timezone.get_current_timezone())
            fecha_ingreso_dt = _dt

        if not equipo.get("marca_id") or not equipo.get("modelo_id"):
            return Response({"detail": "equipo.marca_id y equipo.modelo_id son requeridos"}, status=400)

        c = None
        if cliente.get("id"):
            c = q("SELECT id, cod_empresa, razon_social FROM customers WHERE id=%s", [cliente["id"]], one=True)
        elif cliente.get("cod_empresa"):
            c = q("SELECT id, cod_empresa, razon_social FROM customers WHERE cod_empresa=%s", [cliente["cod_empresa"]], one=True)
        elif cliente.get("razon_social"):
            if _has_table_column("customers", "alias_interno"):
                c = q(
                    """
                    SELECT id, cod_empresa, razon_social
                      FROM customers
                     WHERE LOWER(razon_social)=LOWER(%s)
                        OR LOWER(COALESCE(alias_interno,''))=LOWER(%s)
                     ORDER BY CASE WHEN LOWER(razon_social)=LOWER(%s) THEN 0 ELSE 1 END, id
                     LIMIT 1
                    """,
                    [cliente["razon_social"], cliente["razon_social"], cliente["razon_social"]],
                    one=True,
                )
            else:
                c = q("SELECT id, cod_empresa, razon_social FROM customers WHERE LOWER(razon_social)=LOWER(%s)", [cliente["razon_social"]], one=True)
        else:
            return Response({"detail": "Debe seleccionar un cliente"}, status=400)

        if not c:
            return Response({"detail": "Cliente inexistente"}, status=400)

        if cliente.get("cod_empresa") and c["cod_empresa"] != cliente["cod_empresa"]:
            return Response({"detail": "El código no corresponde a la razón social seleccionada."}, status=400)
        cliente_alias_ok = False
        if cliente.get("razon_social") and _has_table_column("customers", "alias_interno"):
            alias_row = q(
                "SELECT 1 FROM customers WHERE id=%s AND LOWER(COALESCE(alias_interno,''))=LOWER(%s)",
                [c["id"], cliente["razon_social"]],
                one=True,
            )
            cliente_alias_ok = bool(alias_row)
        if cliente.get("razon_social") and c["razon_social"].lower() != (cliente["razon_social"] or "").lower() and not cliente_alias_ok:
            return Response({"detail": "La razón social no corresponde al código seleccionado."}, status=400)

        customer_id = c["id"]

        marca = q("SELECT id FROM marcas WHERE id=%s", [equipo["marca_id"]], one=True)
        model = q("SELECT id FROM models WHERE id=%s AND marca_id=%s", [equipo["modelo_id"], equipo["marca_id"]], one=True)
        if not marca or not model:
            return Response({"detail": "Marca o modelo inexistente"}, status=400)

        prop = data.get("propietario") or {}
        prop_nombre = (prop.get("nombre") or "").strip()
        prop_contacto = (prop.get("contacto") or "").strip()
        prop_doc = (prop.get("doc") or "").strip()

        # Validacion: si el cliente es 'Particular', propietario.nombre y propietario.doc son obligatorios
        try:
            rs_lower = (c.get("razon_social") or "").strip().lower()
        except Exception:
            rs_lower = ""
        if rs_lower == "particular":
            if not prop_nombre or not prop_doc:
                return Response({"detail": "Para cliente 'Particular' es obligatorio completar Nombre y CUIT del propietario"}, status=400)

        if not getattr(request, "_skip_ris_preflight", False):
            user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
            preflight = preflight_ris_for_request_payload(data, user_id=user_id)
            if not preflight.get("can_emit"):
                return Response(preflight, status=409)

        numero_serie = (equipo.get("numero_serie") or "").strip()
        ns_key = (numero_serie or "").replace(" ", "").replace("-", "").upper()
        # El tipo sale del modelo elegido, que es la fuente de verdad del catálogo. El que
        # manda la pantalla queda de respaldo para los modelos que todavía no lo tienen
        # cargado: sin tipo, dos equipos que comparten serie vuelven a ser indistinguibles.
        tipo_equipo = _tipo_equipo_de_modelo(equipo.get("modelo_id")) or (equipo.get("tipo_equipo") or "").strip()
        tipo_key = tipo_equipo.strip().lower()
        # El operador declaró en la pantalla que este equipo comparte el número con otro que
        # ya está en la base (el humidificador que entra con la serie del CPAP).
        otro_equipo_mismo_ns = bool(equipo.get("otro_equipo_mismo_ns"))
        # El lote crea los ingresos uno por uno dentro de una misma transacción, así que los
        # chequeos de duplicado ven lo que acaban de crear los equipos anteriores. Con el tipo
        # en la clave, el humidificador ya no choca contra el CPAP con el que vino.
        allow_open_duplicate = bool(getattr(request, "_allow_open_duplicate", False))
        open_conflict = None
        # Cálculo preliminar de garantía de fábrica desde ventas Bejerman.
        try:
            wcalc = _compute_factory_warranty_from_bejerman_sale(
                numero_serie,
                brand_id=equipo.get("marca_id"),
                model_id=equipo.get("modelo_id"),
                company_key=empresa_bejerman,
                bejerman_sale_payload=bejerman_sale,
            )
        except Exception:
            wcalc = {"garantia": None, "vence_el": None, "fecha_venta": None}

        # Chequeo adicional de duplicado por número interno (MG/NM/NV/CE ####):
        # si existe un ingreso abierto para el mismo MG, se reutiliza ese ingreso.
        if numero_interno:
            has_mg_sale_extended = _has_mg_sale_extended_schema()
            if has_mg_schema and has_mg_sale_extended:
                mg_state_sql = "COALESCE(d.mg_estado,'activo') AS mg_estado,"
                mg_sale_sql = """
                  d.mg_venta_fecha AS mg_venta_fecha,
                  COALESCE(d.mg_venta_factura_numero,'') AS mg_venta_factura_numero,
                  COALESCE(d.mg_venta_remito_numero,'') AS mg_venta_remito_numero,
                  d.mg_venta_customer_id AS mg_venta_customer_id,
                  COALESCE(d.mg_venta_numero_alternativo,'') AS mg_venta_numero_alternativo
                """
            else:
                mg_state_sql = "'activo' AS mg_estado,"
                mg_sale_sql = """
                  NULL AS mg_venta_fecha,
                  '' AS mg_venta_factura_numero,
                  '' AS mg_venta_remito_numero,
                  NULL AS mg_venta_customer_id,
                  '' AS mg_venta_numero_alternativo
                """

            mg_device = q(
                f"""
                SELECT
                  d.id,
                  d.customer_id,
                  COALESCE(d.numero_interno,'') AS numero_interno,
                  COALESCE(d.numero_serie,'') AS numero_serie,
                  COALESCE(d.alquilado,false) AS alquilado,
                  {mg_state_sql}
                  {mg_sale_sql}
                FROM devices d
                WHERE {sql_is_mg_code('d.numero_interno')}
                  AND {sql_mg_key('d.numero_interno')} = {sql_mg_key('%s')}
                LIMIT 1
                """,
                [numero_interno],
                one=True,
            )
            if mg_device:
                mg_flags = resolve_mg_flags(mg_device, mg_owner_id)
                if mg_flags.get("mg_inactivo_venta"):
                    return Response(
                        {
                            "detail": "MG histórico inactivo por venta; no operativo para nuevos ingresos.",
                            "conflict_type": "MG_INACTIVO_VENTA",
                            "payload": {
                                "device_id": mg_device.get("id"),
                                "numero_interno": mg_device.get("numero_interno") or numero_interno,
                                "mg_estado": "inactivo_venta",
                                "mg_venta_fecha": mg_device.get("mg_venta_fecha"),
                                "mg_venta_factura_numero": mg_device.get("mg_venta_factura_numero") or "",
                                "mg_venta_remito_numero": mg_device.get("mg_venta_remito_numero") or "",
                                "mg_venta_customer_id": mg_device.get("mg_venta_customer_id"),
                                "mg_venta_numero_alternativo": mg_device.get("mg_venta_numero_alternativo") or "",
                            },
                        },
                        status=400,
                    )

            dup_mg = q(
                f"""
                SELECT t.id,
                       COALESCE(t.fecha_ingreso, t.fecha_creacion) AS fecha_ingreso,
                       t.fecha_creacion
                  FROM ingresos t
                  JOIN devices d ON d.id = t.device_id
                  LEFT JOIN models m ON m.id = d.model_id
                 WHERE {sql_is_mg_code('d.numero_interno')}
                   AND {sql_mg_key('d.numero_interno')} = {sql_mg_key('%s')}
                   AND LOWER(TRIM(COALESCE(NULLIF(d.tipo_equipo,''), m.tipo_equipo, ''))) = %s
                   AND t.estado NOT IN ('entregado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                 ORDER BY t.id DESC
                 LIMIT 1
                """,
                [numero_interno, tipo_key],
                one=True,
            )
            if dup_mg:
                if ingreso_es_de_hoy(dup_mg):
                    return _ingreso_duplicado_hoy_response(
                        dup_mg,
                        numero_serie=numero_serie,
                        numero_interno=numero_interno,
                        tipo_equipo=tipo_equipo,
                    )
                if not allow_open_duplicate:
                    existing_id = dup_mg["id"]
                    fecha_ingreso_val = dup_mg.get("fecha_ingreso") or dup_mg.get("fecha_creacion")
                    try:
                        fecha_ingreso_iso = fecha_ingreso_val.isoformat()
                    except Exception:
                        fecha_ingreso_iso = fecha_ingreso_val
                    return Response({
                        "ok": True,
                        "ingreso_id": existing_id,
                        "os": os_label(existing_id),
                        "existing": True,
                        "fecha_ingreso": fecha_ingreso_iso,
                    })
                open_conflict = dup_mg

        if ns_key and not open_conflict:
            dup = q(
                """
                SELECT t.id,
                       COALESCE(t.fecha_ingreso, t.fecha_creacion) AS fecha_ingreso,
                       t.fecha_creacion
                  FROM ingresos t
                  JOIN devices d ON d.id = t.device_id
                  LEFT JOIN models m ON m.id = d.model_id
                 WHERE REPLACE(REPLACE(UPPER(d.numero_serie),' ',''),'-','') = %s
                   AND LOWER(TRIM(COALESCE(NULLIF(d.tipo_equipo,''), m.tipo_equipo, ''))) = %s
                   AND t.estado NOT IN ('entregado','alquilado','baja','vendido_pendiente_entrega','vendido_entregado')
                 ORDER BY t.id DESC
                 LIMIT 1
                """,
                [ns_key, tipo_key],
                one=True,
            )
            if dup:
                if ingreso_es_de_hoy(dup):
                    return _ingreso_duplicado_hoy_response(
                        dup,
                        numero_serie=numero_serie,
                        numero_interno=numero_interno,
                        tipo_equipo=tipo_equipo,
                    )
                if not allow_open_duplicate:
                    existing_id = dup["id"]
                    fecha_ingreso_val = dup.get("fecha_ingreso") or dup.get("fecha_creacion")
                    try:
                        fecha_ingreso_iso = fecha_ingreso_val.isoformat()
                    except Exception:
                        fecha_ingreso_iso = fecha_ingreso_val
                    return Response({
                        "ok": True,
                        "ingreso_id": existing_id,
                        "os": os_label(existing_id),
                        "existing": True,
                        "fecha_ingreso": fecha_ingreso_iso,
                    })
                open_conflict = dup

        # Auto-chequeo de garantía de fábrica por N/S si el usuario no la marcó
        # no auto-calculo de garantia de fabrica por ahora (se toma del payload)

        # Resolución consistente de device por NS / MG con validaciones de conflicto
        dev_ns = None
        dev_mg = None
        if ns_key:
            dev_ns = q(
                """
                SELECT id, COALESCE(numero_serie,'') AS numero_serie, COALESCE(numero_interno,'') AS numero_interno
                  FROM devices
                 WHERE REPLACE(REPLACE(UPPER(numero_serie),' ',''),'-','') = %s
                   AND LOWER(TRIM(COALESCE(tipo_equipo,''))) = %s
                 LIMIT 1
                """,
                [ns_key, tipo_key],
                one=True,
            )
            if not dev_ns and not otro_equipo_mismo_ns:
                # Adopción de equipos anteriores al índice por tipo: si hay exactamente uno con
                # esta serie y sin tipo cargado, es el mismo equipo y se lo adopta estampándole
                # el tipo. Crear uno nuevo en su lugar partiría el historial en dos.
                #
                # Salvo que el operador haya dicho lo contrario: cuando declara que está
                # cargando el otro equipo que comparte el número, adoptar sería convertir el
                # CPAP de la base en su humidificador y llevarse su historial puesto.
                legacy = q(
                    """
                    SELECT id, COALESCE(numero_serie,'') AS numero_serie, COALESCE(numero_interno,'') AS numero_interno
                      FROM devices
                     WHERE REPLACE(REPLACE(UPPER(numero_serie),' ',''),'-','') = %s
                       AND COALESCE(TRIM(tipo_equipo),'') = ''
                     LIMIT 2
                    """,
                    [ns_key],
                )
                if len(legacy or []) == 1:
                    dev_ns = legacy[0]
        if numero_interno:
            # Intento exacto primero
            dev_mg = q(
                """
                SELECT id, COALESCE(numero_serie,'') AS numero_serie, COALESCE(numero_interno,'') AS numero_interno
                  FROM devices
                 WHERE numero_interno = %s
                 LIMIT 1
                """,
                [numero_interno],
                one=True,
            )
            if not dev_mg:
                # Búsqueda por número interno normalizado (MG|NM|NV|CE ####)
                dev_mg = q(
                    f"""
                    SELECT id, COALESCE(numero_serie,'') AS numero_serie, COALESCE(numero_interno,'') AS numero_interno
                      FROM devices
                     WHERE {sql_is_mg_code('numero_interno')}
                       AND {sql_mg_key('numero_interno')} = {sql_mg_key('%s')}
                     LIMIT 1
                    """,
                    [numero_interno],
                    one=True,
                )

        conflict_payload = {
            "numero_serie_input": numero_serie or None,
            "numero_interno_input": numero_interno or None,
            "device_ns": dev_ns,
            "device_mg": dev_mg,
        }

        # El número interno identifica una sola máquina. Si el operador declaró que éste es
        # otro equipo con el mismo número de serie y aun así vino un MG que ya está tomado,
        # ese MG es del equipo anterior: sin este corte el ingreso se colgaba del device
        # equivocado y le reescribía marca, modelo y tipo.
        if otro_equipo_mismo_ns and dev_mg and (not dev_ns or dev_ns["id"] != dev_mg["id"]):
            logger.warning("MG tomado por otro equipo con la misma serie: %s", conflict_payload)
            return Response(
                {
                    "detail": (
                        "El número interno ingresado ya está asignado a otro equipo. "
                        "El equipo que comparte el número de serie debe entrar sin ese número "
                        "interno o con el suyo propio."
                    ),
                    "conflict_type": "MG_LINKED_TO_OTHER_NS",
                    "payload": conflict_payload,
                },
                status=400,
            )

        # Clasificar conflictos NS/MG
        if dev_ns and dev_mg and dev_ns["id"] != dev_mg["id"]:
            msg = (
                "Conflicto entre número de serie y número interno. "
                "El número de serie ingresado está asociado a un equipo distinto "
                "del que tiene asignado el número interno ingresado. "
                "Revise que no haya error de tipeo en N/S o MG."
            )
            logger.warning("NS/MG conflict (ids distintos): %s", conflict_payload)
            return Response(
                {
                    "detail": msg,
                    "conflict_type": "NS_MG_MISMATCH",
                    "payload": conflict_payload,
                },
                status=400,
            )

        if dev_mg and numero_serie and (dev_mg.get("numero_serie") or "").strip():
            ns_dev_mg = (dev_mg.get("numero_serie") or "").strip()
            ns_dev_mg_key = ns_dev_mg.replace(" ", "").replace("-", "").upper()
            if ns_dev_mg_key != ns_key:
                msg = (
                    "El número interno ingresado ya está asociado a otro equipo con "
                    "un número de serie distinto. Revise que no haya error de tipeo."
                )
                logger.warning("MG already linked to different NS: %s", conflict_payload)
                return Response(
                    {
                        "detail": msg,
                        "conflict_type": "MG_LINKED_TO_OTHER_NS",
                        "payload": conflict_payload,
                    },
                    status=400,
                )

        # Determinar device_id en base a las resoluciones anteriores
        device_id = None
        if dev_ns:
            device_id = dev_ns["id"]
        elif dev_mg:
            device_id = dev_mg["id"]
        else:
            # No existe device; crear uno nuevo respetando constraints.
            # tipo_equipo va en el INSERT porque forma parte de la identidad: es la columna
            # que el índice único usa junto al número de serie.
            if connection.vendor == "postgresql":
                device_id = exec_returning(
                    """
                    INSERT INTO devices (customer_id, marca_id, model_id, numero_serie, numero_interno, tipo_equipo)
                    VALUES (%s, %s, %s, %s, NULLIF(%s,''), NULLIF(%s,''))
                    ON CONFLICT DO NOTHING
                    RETURNING id
                    """,
                    [customer_id, equipo["marca_id"], equipo["modelo_id"], numero_serie, numero_interno, tipo_equipo],
                )
                if not device_id and ns_key:
                    row_dev = None
                    for _ in range(3):
                        row_dev = q(
                            """
                            SELECT id FROM devices
                             WHERE REPLACE(REPLACE(UPPER(numero_serie),' ',''),'-','') = %s
                               AND LOWER(TRIM(COALESCE(tipo_equipo,''))) = %s
                             LIMIT 1
                            """,
                            [ns_key, tipo_key],
                            one=True,
                        )
                        if row_dev:
                            break
                        time.sleep(0.05)
                    device_id = row_dev and int(row_dev.get("id"))
            else:
                exec_void(
                    """
                    INSERT INTO devices (customer_id, marca_id, model_id, numero_serie, numero_interno, tipo_equipo)
                    VALUES (%s, %s, %s, %s, NULLIF(%s,''), NULLIF(%s,''))
                    """,
                    [customer_id, equipo["marca_id"], equipo["modelo_id"], numero_serie, numero_interno, tipo_equipo],
                )
                device_id = last_insert_id()
        # Blindaje: siempre alinear marca/modelo del device con la selección actual.
        # Evita órdenes con cruces (marca de un lado y modelo de otra marca).
        # El tipo se estampa junto con el modelo: es lo que le da identidad al device frente a
        # otro con la misma serie, y lo que deja adoptado de verdad a un device legacy sin tipo.
        if device_id:
            exec_void(
                "UPDATE devices SET marca_id=%s, model_id=%s, tipo_equipo=COALESCE(NULLIF(%s,''), tipo_equipo) WHERE id=%s",
                [equipo["marca_id"], equipo["modelo_id"], tipo_equipo, device_id],
            )
        # Actualizar snapshot de cliente del device al del ingreso actual
        if device_id and customer_id:
            try:
                with (transaction.atomic() if connection.in_atomic_block else nullcontext()):
                    exec_void("UPDATE devices SET customer_id=%s WHERE id=%s", [customer_id, device_id])
            except Exception:
                logger.exception(
                    "No se pudo actualizar el cliente del equipo durante el ingreso",
                    extra={"device_id": device_id, "customer_id": customer_id},
                )
        # Actualizar fecha de vencimiento de garantía en el device según venta Bejerman.
        if device_id and _has_table_column("devices", "garantia_vence"):
            try:
                with (transaction.atomic() if connection.in_atomic_block else nullcontext()):
                    vence = wcalc.get("vence_el") if isinstance(wcalc, dict) else None
                    if vence is None:
                        exec_void("UPDATE devices SET garantia_vence = NULL WHERE id=%s", [device_id])
                    else:
                        exec_void("UPDATE devices SET garantia_vence = %s WHERE id=%s", [vence, device_id])
            except Exception:
                logger.exception(
                    "No se pudo actualizar el vencimiento de garantía del equipo",
                    extra={"device_id": device_id},
                )
        if numero_interno and device_id:
            try:
                with (transaction.atomic() if connection.in_atomic_block else nullcontext()):
                    exec_void(
                        "UPDATE devices SET numero_interno = NULLIF(%s,'') WHERE id=%s",
                        [numero_interno, device_id],
                    )
            except Exception as e:
                # Seguridad adicional: capturar violaciones de índice único y devolver error entendible
                # sin confirmar ninguna de las escrituras previas de este alta.
                if connection.in_atomic_block:
                    transaction.set_rollback(True)
                msg = (
                    "No se pudo asignar el número interno al equipo porque ya está "
                    "en uso por otro dispositivo. Revise el MG ingresado."
                )
                logger.warning("Error al actualizar numero_interno (device_id=%s, MG=%s): %s", device_id, numero_interno, e)
                return Response(
                    {
                        "detail": msg,
                        "conflict_type": "MG_UNIQUE_CONSTRAINT",
                        "payload": conflict_payload,
                    },
                    status=400,
                )

        auto_gar_rep = False
        last_out_candidates = []
        salida_expr = _salida_taller_expr("t")
        salida_join = _salida_taller_lateral_sql("t")
        if numero_serie:
            row_last_ns = q(
                f"""
                SELECT MAX({salida_expr}) AS last_out
                  FROM ingresos t
                  JOIN devices d ON d.id = t.device_id
                  {salida_join}
                 WHERE d.numero_serie = %s
                   AND {salida_expr} IS NOT NULL
                """,
                [numero_serie],
                one=True,
            )
            last_ns = row_last_ns and row_last_ns.get("last_out")
            if last_ns:
                last_out_candidates.append(last_ns)
        if numero_interno:
            row_mg = q(
                f"""
                SELECT MAX({salida_expr}) AS last_out
                  FROM ingresos t
                  JOIN devices d ON d.id = t.device_id
                  {salida_join}
                 WHERE d.numero_interno = %s
                   AND {salida_expr} IS NOT NULL
                """,
                [numero_interno],
                one=True,
            )
            last_mg = row_mg and row_mg.get("last_out")
            if last_mg:
                last_out_candidates.append(last_mg)
        if last_out_candidates:
            try:
                last_out = max(last_out_candidates)
                auto_gar_rep = (timezone.now() - last_out).days <= 90
            except Exception:
                auto_gar_rep = False
        garantia_rep_payload = bool(data.get("garantia_reparacion"))
        etiq_ok = bool(data.get("etiq_garantia_ok"))
        garantia_rep_final = garantia_rep_payload or auto_gar_rep

        tecnico_id = data.get("tecnico_id")
        if not tecnico_id and _has_table_column("models", "tecnico_id"):
            tdef = q("SELECT tecnico_id FROM models WHERE id=%s", [equipo["modelo_id"]], one=True)
            tecnico_id = tdef["tecnico_id"] if tdef else None
        if not tecnico_id and _has_table_column("marcas", "tecnico_id"):
            tmarca = q("SELECT tecnico_id FROM marcas WHERE id=%s", [equipo["marca_id"]], one=True)
            tecnico_id = (tmarca or {}).get("tecnico_id")
        if tecnico_id:
            tech = q("SELECT id FROM users WHERE id=%s AND activo=true AND rol IN ('tecnico','jefe','jefe_suplente')",
                     [tecnico_id], one=True)
            if not tech:
                return Response({"detail": "Técnico inválido o inactivo"}, status=400)

        uid = getattr(request.user, "id", None) or getattr(request, "user_id", None)
        if not uid:
            return Response({"detail": "Usuario no autenticado"}, status=401)
        _set_audit_user(request)

        # PostgreSQL-only build: motivo es texto válido según enum del modelo

        equipo_variante = (request.data.get("equipo_variante") or "").strip() or None
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        permite_col_sql = ", permite_reparacion" if has_permite_reparacion else ""
        permite_val_sql = ", %s" if has_permite_reparacion else ""
        insert_params = [
            device_id, motivo, ubicacion_id, uid, tecnico_id,
            informe_preliminar, accesorios_text, comentarios_text, equipo_variante,
            prop_nombre, prop_contacto, prop_doc,
            garantia_rep_final, wcalc.get("garantia"), etiq_ok,
        ]
        if has_permite_reparacion:
            insert_params.append(bool(permite_reparacion_default))

        # Asegurar que tenemos un device_id válido antes de crear el ingreso
        if not device_id:
            return Response({
                "detail": "No se pudo identificar el equipo. Ingrese número de serie o número interno válido."
            }, status=400)
        ingreso_id = exec_returning(
            f"""
            INSERT INTO ingresos (
              device_id, motivo, ubicacion_id, recibido_por, asignado_a,
              informe_preliminar, accesorios, comentarios, equipo_variante,
              propietario_nombre, propietario_contacto, propietario_doc,
              garantia_reparacion, garantia_fabrica, etiq_garantia_ok
              {permite_col_sql}
            )
            VALUES (%s,%s,%s,%s,%s,
                    %s,%s,%s,%s,
                    NULLIF(%s,''), NULLIF(%s,''), NULLIF(%s,''),
                    %s, %s, %s
                    {permite_val_sql})
            RETURNING id
            """,
             insert_params
        )

        # Setear empresa Bejerman/branding si las columnas existen en la base.
        try:
            with (transaction.atomic() if connection.in_atomic_block else nullcontext()):
                company_sets = []
                company_params = []
                if _has_table_column("ingresos", "empresa_bejerman"):
                    company_sets.append("empresa_bejerman=%s")
                    company_params.append(empresa_bejerman)
                if _has_table_column("ingresos", "empresa_facturar"):
                    company_sets.append("empresa_facturar=%s")
                    company_params.append(empresa_facturar)
                if company_sets:
                    company_params.append(ingreso_id)
                    exec_void(f"UPDATE ingresos SET {', '.join(company_sets)} WHERE id=%s", company_params)
        except Exception:
            logger.exception(
                "No se pudo guardar la empresa Bejerman del ingreso",
                extra={"ingreso_id": ingreso_id, "empresa_bejerman": empresa_bejerman},
            )

        sets, params = [], []
        if remito_ingreso:
            sets.append("remito_ingreso = NULLIF(%s,'')")
            params.append(remito_ingreso)
        if fecha_ingreso_dt is not None:
            sets.append("fecha_ingreso = %s")
            params.append(fecha_ingreso_dt)
        if fecha_ingreso_dt is None:
            sets.append("fecha_ingreso = COALESCE(fecha_ingreso, NOW())")

        if sets:
            params.append(ingreso_id)
            exec_void(
                f"UPDATE ingresos SET {', '.join(sets)} WHERE id=%s",
                params,
            )

        for it in (accesorios_items or []):
            try:
                acc_id = int(it.get("accesorio_id"))
            except (TypeError, ValueError):
                continue
            ref = (it.get("referencia") or "").strip() or None
            desc = (it.get("descripcion") or "").strip() or None
            exec_void(
              "INSERT INTO ingreso_accesorios(ingreso_id, accesorio_id, referencia, descripcion) VALUES (%s, %s, %s, %s)",
              [ingreso_id, acc_id, ref, desc]
            )

        payload = {"ok": True, "ingreso_id": ingreso_id, "os": os_label(ingreso_id)}
        if open_conflict:
            # El ingreso se creó igual (lo pidió el lote con `_allow_open_duplicate`), pero
            # arrastra un conflicto: el equipo ya tenía una OS abierta. Quien llamó decide qué
            # hacer con el dato; acá sólo se informa.
            fecha_conflicto = open_conflict.get("fecha_ingreso") or open_conflict.get("fecha_creacion")
            try:
                fecha_conflicto = fecha_conflicto.isoformat()
            except Exception:
                pass
            payload["open_conflict"] = {
                "conflicting_ingreso_id": open_conflict["id"],
                "conflicting_os": os_label(open_conflict["id"]),
                "conflicting_fecha_ingreso": fecha_conflicto,
                "numero_serie": numero_serie,
                "tipo_equipo": tipo_equipo,
            }
        return Response(payload, status=201)


class NuevoIngresoLoteError(Exception):
    def __init__(self, payload, status_code=400, item_index=None):
        super().__init__((payload or {}).get("detail") if isinstance(payload, dict) else str(payload))
        self.payload = payload if isinstance(payload, dict) else {"detail": str(payload)}
        self.status_code = status_code
        self.item_index = item_index


def _child_ingreso_request(request, payload, **flags):
    user = getattr(request, "user", None)
    child = SimpleNamespace(
        data=payload,
        user=user,
        user_id=getattr(request, "user_id", None) or getattr(user, "id", None),
        user_obj=getattr(request, "user_obj", None) or user,
        user_role=getattr(request, "user_role", None) or getattr(user, "rol", ""),
        _skip_ris_preflight=True,
        **flags,
    )
    cached_permissions = getattr(request, "_effective_permissions_v2", None)
    if isinstance(cached_permissions, dict):
        setattr(child, "_effective_permissions_v2", cached_permissions)
    return child


def _register_ingreso_conflicts(conflicts, request):
    """Anota los conflictos del lote y avisa al jefe. Nunca corta la emisión.

    Si esto falla, el remito igual tiene que salir: el precio de perder el aviso es mucho
    menor que el de dejar al cliente sin comprobante.
    """

    if not conflicts:
        return []
    user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
    registered = []
    for conflict in conflicts:
        try:
            conflict_id = record_conflict(
                conflict["ingreso_id"],
                conflict["conflicting_ingreso_id"],
                numero_serie=conflict.get("numero_serie") or "",
                tipo_equipo=conflict.get("tipo_equipo") or "",
                user_id=user_id,
            )
        except Exception:
            logger.exception(
                "ingreso_open_conflict_record_failed",
                extra={"ingreso_id": conflict.get("ingreso_id")},
            )
            continue
        registered.append({**conflict, "conflict_id": conflict_id})
        try:
            notify_ingreso_conflicto_abierto(
                conflict["ingreso_id"],
                conflict["conflicting_ingreso_id"],
            )
        except Exception:
            logger.exception(
                "ingreso_open_conflict_notify_failed",
                extra={"ingreso_id": conflict.get("ingreso_id")},
            )
    return registered


def _ris_lote_urls(first_ingreso_id, remitos=None):
    representatives = [
        item.get("ingreso_id")
        for item in (remitos or [])
        if isinstance(item, dict) and item.get("ingreso_id")
    ]
    if len(representatives) > 1:
        raw_ids = ",".join(str(item) for item in representatives)
        return {
            "pdf_url": f"/api/ingresos/ris/lote/pdf/?ids={raw_ids}",
            "print_url": f"/api/ingresos/ris/lote/print/?ids={raw_ids}",
        }
    return {
        "pdf_url": f"/api/ingresos/{first_ingreso_id}/ris/pdf/",
        "print_url": f"/api/ingresos/{first_ingreso_id}/ris/print/",
    }


def _ris_lote_remitos(ingreso_ids, batch_result=None):
    raw_entries = batch_result.get("remitos") if isinstance(batch_result, dict) else None
    entries = []
    for raw in raw_entries if isinstance(raw_entries, list) else []:
        if not isinstance(raw, dict):
            continue
        raw_ids = raw.get("ingreso_ids") or raw.get("ingresoIds") or []
        clean_ids = []
        for value in raw_ids if isinstance(raw_ids, (list, tuple)) else []:
            try:
                ingreso_id = int(value)
            except (TypeError, ValueError):
                continue
            if ingreso_id > 0 and ingreso_id not in clean_ids:
                clean_ids.append(ingreso_id)
        if not clean_ids:
            try:
                ingreso_id = int(raw.get("ingreso_id") or raw.get("ingresoId") or 0)
            except (TypeError, ValueError):
                ingreso_id = 0
            if ingreso_id > 0:
                clean_ids.append(ingreso_id)
        remito_number = str(raw.get("remito_number") or raw.get("remitoNumber") or "").strip()
        if not clean_ids or not remito_number:
            continue
        representative_id = clean_ids[0]
        entries.append(
            {
                "ingreso_id": representative_id,
                "ingreso_ids": clean_ids,
                "remito_number": remito_number,
                "comprobante_tipo": raw.get("comprobante_tipo") or "RIS",
                "line_count": raw.get("line_count") or raw.get("lineCount") or 0,
                "equipment_count": len(clean_ids),
                "pdf_status": raw.get("pdf_status") or "pending",
                "pdf_url": f"/api/ingresos/{representative_id}/ris/pdf/",
                "print_url": f"/api/ingresos/{representative_id}/ris/print/",
            }
        )
    if entries:
        return entries

    grouped = {}
    for ingreso_id in ingreso_ids:
        status = get_ris_status_for_ingreso(ingreso_id)
        remito_number = str(status.get("remito_number") or "").strip()
        if not remito_number:
            continue
        entry = grouped.setdefault(
            remito_number,
            {
                "ingreso_id": ingreso_id,
                "ingreso_ids": [],
                "remito_number": remito_number,
                "comprobante_tipo": status.get("comprobante_tipo") or "RIS",
                "line_count": 0,
                "equipment_count": 0,
                "pdf_status": status.get("pdf_status") or "pending",
                "pdf_url": f"/api/ingresos/{ingreso_id}/ris/pdf/",
                "print_url": f"/api/ingresos/{ingreso_id}/ris/print/",
            },
        )
        entry["ingreso_ids"].append(ingreso_id)
        entry["equipment_count"] = len(entry["ingreso_ids"])
    return list(grouped.values())


def _ris_lote_result_fields(ingreso_ids, first_ingreso_id, batch_result=None):
    remitos = _ris_lote_remitos(ingreso_ids, batch_result)
    remito_numbers = [item["remito_number"] for item in remitos]
    return {
        "multiple": len(remito_numbers) > 1,
        "remitos": remitos,
        "remito_numbers": remito_numbers,
        "remitoNumbers": remito_numbers,
        **_ris_lote_urls(first_ingreso_id, remitos),
    }


@method_decorator(transaction.non_atomic_requests, name="dispatch")
class NuevoIngresoLoteView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        require_any_permission(request, ["action.ingreso.create", "page.new_ingreso"])
        data = request.data or {}
        ris_mode = parse_ris_document_mode(data.get("ris_mode") or data.get("document_mode"))
        manual_remito_number = (
            data.get("manual_remito_number")
            or data.get("manualRemitoNumber")
            or data.get("remito_ingreso")
            or ""
        )
        raw_items = data.get("items") or []
        if not isinstance(raw_items, list) or not raw_items:
            return Response({"detail": "Debe agregar al menos un equipo al lote."}, status=400)
        user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        preflight = preflight_ris_for_request_payload(data, user_id=user_id)
        if not preflight.get("can_emit"):
            return Response(preflight, status=409)

        shared_keys = (
            "cliente",
            "propietario",
            "empresa_bejerman",
            "empresa_facturar",
            "fecha_ingreso",
            "ubicacion_id",
        )
        ingreso_view = NuevoIngresoView()
        ingresos = []
        conflicts = []
        try:
            with transaction.atomic():
                for index, item in enumerate(raw_items):
                    if not isinstance(item, dict):
                        raise NuevoIngresoLoteError({"detail": "Cada equipo del lote debe ser un objeto."}, 400, index)
                    payload = dict(item)
                    for key in shared_keys:
                        if key in data and key not in payload:
                            payload[key] = data.get(key)
                    # Un equipo con la OS anterior todavía abierta ya no frena el lote: el
                    # cliente está esperando el remito y el resto de los equipos no tienen la
                    # culpa. Se crea igual y el conflicto queda anotado para resolverlo después.
                    child_response = ingreso_view.post(
                        _child_ingreso_request(request, payload, _allow_open_duplicate=True)
                    )
                    body = getattr(child_response, "data", None) or {}
                    status_code = getattr(child_response, "status_code", 500)
                    if status_code >= 400:
                        raise NuevoIngresoLoteError(body, status_code, index)
                    ingresos.append(dict(body))
                    conflict = body.get("open_conflict") if isinstance(body, dict) else None
                    if conflict and body.get("ingreso_id"):
                        conflicts.append({**conflict, "ingreso_id": body["ingreso_id"], "item_index": index})
        except NuevoIngresoLoteError as exc:
            payload = {**exc.payload, "item_index": exc.item_index}
            return Response(payload, status=exc.status_code)

        ingreso_ids = [item.get("ingreso_id") for item in ingresos if item.get("ingreso_id")]
        first_ingreso_id = ingreso_ids[0] if ingreso_ids else None
        if not first_ingreso_id:
            return Response({"detail": "No se pudo identificar el primer ingreso creado."}, status=500)

        # Antes de emitir: los conflictos tienen que estar anotados, porque de esa tabla sale
        # la decisión de imprimir el renglón sin número de OS.
        conflicts = _register_ingreso_conflicts(conflicts, request)

        user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        is_register_mode = ris_mode == RIS_DOCUMENT_MODE_REGISTER
        try:
            row = (
                register_ris_batch(ingreso_ids, manual_remito_number, user_id=user_id)
                if is_register_mode
                else emit_or_get_ris_batch(ingreso_ids, user_id=user_id)
            )
            ris = get_ris_status_for_ingreso(first_ingreso_id)
            lote_result = _ris_lote_result_fields(ingreso_ids, first_ingreso_id, row)
            remito_number = (
                (lote_result.get("remito_numbers") or [""])[0]
                or ris.get("remito_number")
                or row.get("remito_number")
                or ""
            )
            if conflicts and remito_number:
                # Deja rastro de en qué remito salió el equipo trabado, para poder ubicarlo
                # cuando alguien vaya a resolver el conflicto.
                try:
                    attach_conflict_remito_number(
                        [item["ingreso_id"] for item in conflicts], remito_number
                    )
                except Exception:
                    logger.exception("ingreso_open_conflict_attach_remito_failed")
            response_payload = {
                "ok": True,
                "ingresos": ingresos,
                "ingreso_ids": ingreso_ids,
                "remito_number": remito_number,
                "ris": ris,
                "document_mode": ris.get("document_mode") or ris_mode,
                "pdf_status": ris.get("pdf_status") or ("not_applicable" if is_register_mode else "pending"),
                "open_conflicts": conflicts,
                **{
                    key: value
                    for key, value in lote_result.items()
                    if not is_register_mode or key not in {"pdf_url", "print_url"}
                },
            }
            if is_register_mode:
                response_payload["detail"] = (
                    f"Ingresos creados y remito {remito_number} registrado en Bejerman."
                    if remito_number
                    else "Ingresos creados y remito registrado en Bejerman."
                )
            else:
                response_payload["retry_after_ms"] = 2500
            return Response(
                response_payload,
                status=201,
            )
        except BejermanRisPreflightError as exc:
            ris = get_ris_status_for_ingreso(first_ingreso_id)
            lote_result = _ris_lote_result_fields(ingreso_ids, first_ingreso_id)
            response_payload = {
                "ok": True,
                "ingresos": ingresos,
                "ingreso_ids": ingreso_ids,
                "remito_number": ris.get("remito_number") or "",
                "ris": ris,
                "document_mode": ris.get("document_mode") or ris_mode,
                "pdf_status": ris.get("pdf_status") or ("not_applicable" if is_register_mode else "pending"),
                "open_conflicts": conflicts,
                **{
                    key: value
                    for key, value in lote_result.items()
                    if not is_register_mode or key not in {"pdf_url", "print_url"}
                },
                **exc.payload,
            }
            if not is_register_mode:
                response_payload["retry_after_ms"] = 2500
            return Response(response_payload, status=409)
        except BejermanRisError as exc:
            ris = get_ris_status_for_ingreso(first_ingreso_id)
            lote_result = _ris_lote_result_fields(ingreso_ids, first_ingreso_id)
            action = "registrar remito" if is_register_mode else "emitir RIS"
            response_payload = {
                "ok": True,
                "ingresos": ingresos,
                "ingreso_ids": ingreso_ids,
                "remito_number": ris.get("remito_number") or "",
                "ris": ris,
                "document_mode": ris.get("document_mode") or ris_mode,
                "pdf_status": ris.get("pdf_status") or ("not_applicable" if is_register_mode else "pending"),
                "open_conflicts": conflicts,
                **{
                    key: value
                    for key, value in lote_result.items()
                    if not is_register_mode or key not in {"pdf_url", "print_url"}
                },
                "detail": f"Ingresos creados. No se pudo {action}: {exc}",
            }
            if not is_register_mode:
                response_payload["retry_after_ms"] = 2500
            return Response(response_payload, status=201)


class GarantiaReparacionCheckView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        ns = (request.GET.get("numero_serie") or "").strip()
        # Normalizado igual que se guarda: con "mg 7359" la comparación exacta de abajo no
        # encontraba la ficha "MG 7359" y la garantía de reparación no saltaba.
        mg = normalize_mg_input(request.GET.get("numero_interno") or request.GET.get("mg"))
        # El tipo de equipo desempata los equipos que comparten número de serie: el
        # humidificador entra con la serie del CPAP y no hereda su garantía de reparación.
        # Es opcional: sin tipo el chequeo se comporta como antes.
        tipo_key = (request.GET.get("tipo_equipo") or "").strip().lower()
        if not ns and not mg:
            return Response({"within_90_days": False, "last_ingreso": None})

        # Los equipos anteriores al índice por tipo pueden no tenerlo cargado: esos siguen
        # contando, porque descartarlos sería inventar que no hay garantía.
        tipo_filter_sql = """
                   AND (
                        %s = ''
                     OR LOWER(TRIM(COALESCE(NULLIF(d.tipo_equipo,''), m.tipo_equipo, ''))) IN ('', %s)
                   )
        """

        last_out_candidates = []
        salida_expr = _salida_taller_expr("t")
        salida_join = _salida_taller_lateral_sql("t")
        if ns:
            row_ns = q(
                f"""
                SELECT MAX({salida_expr}) AS last_out
                  FROM ingresos t
                  JOIN devices d ON d.id = t.device_id
                  LEFT JOIN models m ON m.id = d.model_id
                  {salida_join}
                 WHERE d.numero_serie = %s
                   AND {salida_expr} IS NOT NULL
                   {tipo_filter_sql}
                """,
                [ns, tipo_key, tipo_key],
                one=True,
            )
            last_ns = row_ns and row_ns.get("last_out")
            if last_ns:
                last_out_candidates.append(last_ns)

        if mg:
            row_mg = q(
                f"""
                SELECT MAX({salida_expr}) AS last_out
                  FROM ingresos t
                  JOIN devices d ON d.id = t.device_id
                  LEFT JOIN models m ON m.id = d.model_id
                  {salida_join}
                 WHERE d.numero_interno = %s
                   AND {salida_expr} IS NOT NULL
                   {tipo_filter_sql}
                """,
                [mg, tipo_key, tipo_key],
                one=True,
            )
            last_mg = row_mg and row_mg.get("last_out")
            if last_mg:
                last_out_candidates.append(last_mg)

        if not last_out_candidates:
            return Response({"within_90_days": False, "last_ingreso": None})

        last_out = max(last_out_candidates)
        within = (timezone.now() - last_out).days <= 90
        return Response({"within_90_days": within, "last_ingreso": last_out})

class GarantiaFabricaCheckView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def get(self, request):
        # Cálculo basado en venta Bejerman cacheada.
        ns = (request.GET.get("numero_serie") or "").strip()
        mg = normalize_mg_input(request.GET.get("numero_interno") or request.GET.get("mg"))

        if not ns and not mg:
            return Response({"within_365_days": None, "fecha_venta": None, "found": False})

        # Permitir que el front envíe brand_id/model_id explícitos
        brand_id_explicit = request.GET.get("brand_id")
        model_id_explicit = request.GET.get("model_id")
        try:
            brand_id_explicit = int(brand_id_explicit) if brand_id_explicit not in (None, "", "null") else None
        except Exception:
            brand_id_explicit = None
        try:
            model_id_explicit = int(model_id_explicit) if model_id_explicit not in (None, "", "null") else None
        except Exception:
            model_id_explicit = None

        device = None
        if ns:
            device = q(
                "SELECT id, numero_serie, marca_id, model_id FROM devices WHERE numero_serie=%s LIMIT 1",
                [ns],
                one=True,
            )
        if (not device) and mg:
            device = q(
                "SELECT id, numero_serie, marca_id, model_id FROM devices WHERE numero_interno=%s LIMIT 1",
                [mg],
                one=True,
            )

        numero_serie = (device or {}).get("numero_serie") or ns
        # Prioridad: IDs explícitos enviados por el front; si no, los del device (si existe)
        brand_id = brand_id_explicit if brand_id_explicit is not None else (device or {}).get("marca_id")
        model_id = model_id_explicit if model_id_explicit is not None else (device or {}).get("model_id")
        try:
            calc = _compute_factory_warranty_from_bejerman_sale(
                numero_serie,
                brand_id=brand_id,
                model_id=model_id,
                company_key=request.GET.get("company_key") or request.GET.get("empresa_bejerman") or None,
            )
        except Exception:
            # Sin este log un error de cálculo se ve idéntico a "no hay venta": las dos
            # cosas devuelven garantía nula y el operador no tiene cómo distinguirlas.
            logger.warning("garantia_fabrica_check_failed", exc_info=True, extra={"numeroSerie": numero_serie})
            calc = {"garantia": None, "fecha_venta": None, "vence_el": None, "meta": {}}

        en_garantia = calc.get("garantia")
        fecha_venta = calc.get("fecha_venta")
        vence = calc.get("vence_el")
        meta = calc.get("meta") or {}

        return Response({
            "within_365_days": en_garantia,
            "fecha_venta": (fecha_venta.isoformat() if fecha_venta else None),
            "garantia_vence": (vence.isoformat() if vence else None),
            "days": calc.get("days"),
            "warranty_days": calc.get("days"),
            "found": bool(fecha_venta),
            # La fecha sale de un remito, no de una factura: quien la consuma tiene que
            # mostrarla como estimada y no dar la garantía por cierta.
            "estimated": bool(meta.get("estimated")),
            "source": "bejerman_sale",
            "meta": meta,
        })


def _can_read_ingreso_detail(request, ingreso_id: int) -> bool:
    has_broad_detail = user_has_any_permission(request, INGRESO_DETAIL_READ_PERMISSIONS)
    if has_broad_detail:
        return True
    has_liberados_detail = user_has_permission(request, "page.liberados")
    if not has_liberados_detail:
        raise PermissionDenied("No autorizado")
    estado_row = q("SELECT estado FROM ingresos WHERE id=%s", [ingreso_id], one=True)
    if not estado_row:
        return False
    if (estado_row.get("estado") or "").strip().lower() != "liberado":
        raise PermissionDenied("Solo se pueden ver ingresos liberados")
    return True


def _bejerman_status_deposit_names(client: BejermanSDKClient) -> dict[str, str]:
    try:
        response = cached_deposit_response(client)
    except Exception:
        # Los nombres son decoración del panel: si Bejerman no contesta se
        # muestran los códigos y el resto del estado igual se arma.
        response = last_good_deposit_response(getattr(client, "company_key", "")) or {}
    names: dict[str, str] = {}
    for record in records_from_response(response):
        code = as_string(first_value(record, BEJERMAN_STATUS_DEPOSIT_CODE_FIELDS)).upper()
        name = as_string(first_value(record, BEJERMAN_STATUS_DEPOSIT_NAME_FIELDS))
        if code:
            names.setdefault(code, name or code)
    return names


def _bejerman_status_partida_expirations(client: BejermanSDKClient) -> dict[str, str]:
    try:
        rows = build_partida_rows(client.obtener_partidas())
    except Exception:
        return {}
    expirations: dict[str, str] = {}
    for row in rows:
        partida = as_string(row.get("partida")).upper()
        deposit = as_string(row.get("depositCode")).upper()
        expiration = as_string(row.get("expirationDate"))
        if not partida or not expiration:
            continue
        if deposit:
            expirations.setdefault(f"{deposit}|{partida}", expiration)
        expirations.setdefault(partida, expiration)
    return expirations


def _stock_effective_quantity(row: dict) -> float:
    available = as_number(row.get("availableQuantity"))
    real = as_number(row.get("realQuantity"))
    value = available if available is not None else real
    return float(value or 0)


def _bejerman_equipment_stock_item(
    row: dict,
    *,
    deposit_names: dict[str, str],
    partida_expirations: dict[str, str],
) -> dict:
    raw = row.get("raw") if isinstance(row.get("raw"), dict) else {}
    deposit_code = as_string(row.get("depositCode")).upper()
    partida = as_string(row.get("partida"))
    expiration = as_string(first_value(raw, BEJERMAN_STATUS_EXPIRATION_FIELDS))
    if not expiration:
        expiration = (
            partida_expirations.get(f"{deposit_code}|{partida.upper()}")
            or partida_expirations.get(partida.upper())
            or ""
        )
    return {
        "depositCode": deposit_code,
        "depositName": deposit_names.get(deposit_code, ""),
        "articleCode": as_string(row.get("articleCode")),
        "articleDescription": as_string(first_value(raw, BEJERMAN_STATUS_ARTICLE_DESCRIPTION_FIELDS)),
        "partida": partida,
        "realQuantity": row.get("realQuantity"),
        "committedQuantity": row.get("committedQuantity"),
        "availableQuantity": row.get("availableQuantity"),
        "expirationDate": expiration or None,
    }


def build_ingreso_bejerman_estado(ingreso_id: int, actor_user_id: int | None) -> dict:
    empresa_bejerman_sql = (
        "COALESCE(NULLIF(t.empresa_bejerman, ''), 'SEPID') AS empresa_bejerman"
        if _has_table_column("ingresos", "empresa_bejerman")
        else "'SEPID' AS empresa_bejerman"
    )
    row = q(
        f"""
        SELECT t.id,
               {empresa_bejerman_sql},
               COALESCE(d.numero_serie, '') AS numero_serie,
               COALESCE(d.numero_interno, '') AS numero_interno
          FROM ingresos t
          JOIN devices d ON d.id = t.device_id
         WHERE t.id = %s
        """,
        [ingreso_id],
        one=True,
    )
    if not row:
        return {"detail": "Ingreso no encontrado", "_status_code": 404}

    numero_serie = as_string(row.get("numero_serie"))
    numero_interno = as_string(row.get("numero_interno"))
    identifier = numero_serie or numero_interno
    identifier_source = "numero_serie" if numero_serie else ("numero_interno" if numero_interno else "")
    company_key = as_string(row.get("empresa_bejerman")) or "SEPID"
    checked_at = timezone.now().isoformat()

    base_payload = {
        "status": "missing_identifier",
        "companyKey": company_key,
        "identifier": identifier,
        "identifierSource": identifier_source,
        "checkedAt": checked_at,
        "items": [],
        "warning": "",
        "action": None,
    }
    if not identifier:
        base_payload["warning"] = "El equipo no tiene número de serie ni número interno para consultar Bejerman."
        return base_payload

    client = BejermanSDKClient(company_key=company_key, actor_user_id=actor_user_id)
    stock_response = client.obtener_stock_partida(identifier)
    stock_rows = build_stock_rows(stock_response)
    deposit_names = _bejerman_status_deposit_names(client) if stock_rows else {}
    partida_expirations = _bejerman_status_partida_expirations(client) if stock_rows else {}
    items = [
        _bejerman_equipment_stock_item(
            stock_row,
            deposit_names=deposit_names,
            partida_expirations=partida_expirations,
        )
        for stock_row in stock_rows
    ]
    items.sort(
        key=lambda item: (
            0 if _stock_effective_quantity(item) > 0 else 1,
            as_string(item.get("depositCode")),
            as_string(item.get("articleCode")),
        )
    )

    if not items:
        status = "not_found"
        warning = f"Bejerman no devolvió registros para la partida {identifier}."
    elif any(_stock_effective_quantity(item) > 0 for item in items):
        status = "found"
        warning = ""
    else:
        status = "zero_stock"
        warning = f"Bejerman devolvió la partida {identifier}, pero sin stock positivo."

    action = None
    if status == "zero_stock":
        release_remito = _service_release_remito_salida_for_ingreso(ingreso_id)
        if release_remito:
            warning = (
                f"Bejerman devolvió la partida {identifier} sin stock positivo, "
                f"pero es lo esperado: el equipo ya fue liberado ({release_remito})."
            )
        else:
            ris_status = get_ris_status_for_ingreso(ingreso_id)
            if ris_status.get("status") == "generated" and as_string(ris_status.get("remito_number")):
                action = {
                    "type": "force_stock_entry",
                    "label": "Forzar ingreso de stock en Bejerman",
                }

    return {
        **base_payload,
        "status": status,
        "items": items,
        "warning": warning,
        "action": action,
    }


def _bejerman_estado_error_response(exc: Exception, ingreso_id: int | None = None) -> Response:
    detail = str(exc) or "No fue posible consultar Bejerman."
    lower = detail.lower()
    if isinstance(exc, BejermanSdkConfigError) and "credencial" in lower:
        code = "BEJERMAN_CREDENTIALS_REQUIRED"
    elif isinstance(exc, BejermanSdkConfigError):
        code = "BEJERMAN_SDK_NOT_CONFIGURED"
    elif isinstance(exc, BejermanSdkUnavailable):
        code = "BEJERMAN_SDK_UNAVAILABLE"
    elif isinstance(exc, BejermanSdkResponseError):
        code = "BEJERMAN_SDK_ERROR"
    else:
        code = "BEJERMAN_ERROR"
    # Hasta acá estos 503 sólo dejaban rastro como "Service Unavailable" en el log
    # de django.request, sin el motivo: las excepciones del SDK se atrapan y se
    # traducen sin loguearse. El motivo va en el mensaje y no en `extra` porque el
    # formatter configurado ("%(asctime)s %(levelname)s %(name)s: %(message)s") no
    # imprime los campos extra.
    logger.warning(
        "ingreso_bejerman_estado_sdk_error ingreso_id=%s code=%s exc=%s detail=%s",
        ingreso_id if ingreso_id is not None else "-",
        code,
        type(exc).__name__,
        detail[:300],
    )
    return Response({"detail": detail, "code": code}, status=503)


class IngresoBejermanEstadoView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ingreso_id: int):
        if not _can_read_ingreso_detail(request, ingreso_id):
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        try:
            payload = build_ingreso_bejerman_estado(ingreso_id, actor_user_id)
        except (BejermanSdkConfigError, BejermanSdkUnavailable, BejermanSdkResponseError) as exc:
            return _bejerman_estado_error_response(exc, ingreso_id)
        except Exception as exc:
            logger.warning("ingreso_bejerman_estado_failed", exc_info=True, extra={"ingreso_id": ingreso_id})
            return _bejerman_estado_error_response(exc, ingreso_id)
        status_code = int(payload.pop("_status_code", 200))
        action = payload.get("action") or {}
        if action.get("type") == "force_stock_entry" and not user_has_any_permission(
            request, ["action.bejerman_sync.force_stock_entry"]
        ):
            # El botón se ofrece sólo a quien puede usarlo: antes lo veía cualquiera que
            # abriera el estado y al apretarlo recibía "No autorizado".
            payload["action"] = None
        return Response(payload, status=status_code)


class IngresoBejermanEstadoAccionView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        if not _can_read_ingreso_detail(request, ingreso_id):
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        require_roles(request, ["jefe", "jefe_veedor", "supervisor"])
        _set_audit_user(request)
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        body = request.data if isinstance(request.data, dict) else {}
        action = str(body.get("action") or "").strip()
        if action != "force_stock_entry":
            return Response({"detail": "Acción de Bejerman inválida."}, status=400)
        try:
            result = force_stock_entry_for_ingreso(ingreso_id, user_id=actor_user_id)
            state = build_ingreso_bejerman_estado(ingreso_id, actor_user_id)
        except BejermanRisError as exc:
            return Response({"detail": str(exc)}, status=409)
        except (BejermanSdkConfigError, BejermanSdkUnavailable, BejermanSdkResponseError) as exc:
            return _bejerman_estado_error_response(exc, ingreso_id)
        except Exception as exc:
            logger.warning("ingreso_bejerman_forzar_stock_failed", exc_info=True, extra={"ingreso_id": ingreso_id})
            return _bejerman_estado_error_response(exc, ingreso_id)
        state.pop("_status_code", None)
        return Response({"ok": True, "result": result, "state": state})


class IngresoAlquilerRemitoView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ingreso_id: int):
        if not _can_read_ingreso_detail(request, ingreso_id):
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        remito_number = request.GET.get("remito") or None
        try:
            state = get_rental_remito_state(ingreso_id, remito_number, actor_user_id=actor_user_id)
        except RentalRemitoError as exc:
            return Response({"detail": exc.detail, "code": exc.code}, status=exc.status_code)
        except (BejermanSdkConfigError, BejermanSdkUnavailable, BejermanSdkResponseError) as exc:
            return _bejerman_estado_error_response(exc, ingreso_id)
        except Exception as exc:
            logger.warning("ingreso_alquiler_remito_estado_failed", exc_info=True, extra={"ingreso_id": ingreso_id})
            return _bejerman_estado_error_response(exc, ingreso_id)
        return Response(state)


class IngresoAlquilerRemitoPdfView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ingreso_id: int):
        if not _can_read_ingreso_detail(request, ingreso_id):
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        try:
            pdf_bytes, content_type, filename = fetch_rental_remito_pdf(ingreso_id, actor_user_id=actor_user_id)
        except RentalRemitoError as exc:
            return Response({"detail": exc.detail, "code": exc.code}, status=exc.status_code)
        except Exception as exc:
            logger.warning("ingreso_alquiler_remito_pdf_failed", exc_info=True, extra={"ingreso_id": ingreso_id})
            return _bejerman_estado_error_response(exc, ingreso_id)
        response = HttpResponse(pdf_bytes, content_type=content_type or "application/pdf")
        response["Content-Disposition"] = f'inline; filename="{filename}"'
        return response


class IngresoAlquilerRemitoAccionView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        require_roles(request, ["jefe", "jefe_veedor", "admin", "recepcion"])
        require_any_permission(request, ["action.ingreso.fix_ris_preflight", "action.bejerman_sync.manage"])
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        body = request.data if isinstance(request.data, dict) else {}
        action = str(body.get("action") or "").strip()
        remito_number = body.get("remitoNumber") or body.get("remito_number") or None
        try:
            result = execute_rental_remito_action(
                ingreso_id,
                action,
                remito_number=remito_number,
                actor_user_id=actor_user_id,
            )
        except RentalRemitoError as exc:
            return Response({"detail": exc.detail, "code": exc.code}, status=exc.status_code)
        except (BejermanSdkConfigError, BejermanSdkUnavailable, BejermanSdkResponseError) as exc:
            return _bejerman_estado_error_response(exc, ingreso_id)
        except Exception as exc:
            logger.warning(
                "ingreso_alquiler_remito_accion_failed",
                exc_info=True,
                extra={"ingreso_id": ingreso_id, "action": action},
            )
            return _bejerman_estado_error_response(exc, ingreso_id)
        return Response(result)


class IngresoDetalleView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ingreso_id: int):
        try:
            strong = str(request.GET.get("strong", "")).strip().lower() in ("1","true","yes")
        except Exception:
            strong = False
        if not _can_read_ingreso_detail(request, ingreso_id):
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        permite_reparacion_sql = (
            "COALESCE(t.permite_reparacion, TRUE) AS permite_reparacion,"
            if has_permite_reparacion
            else "TRUE AS permite_reparacion,"
        )
        has_mg_schema = _has_mg_schema()
        has_mg_sale_extended = _has_mg_sale_extended_schema()
        mg_sale_customer_join = (
            "LEFT JOIN customers c_mg_sale ON c_mg_sale.id = d.mg_venta_customer_id"
            if has_mg_sale_extended
            else "LEFT JOIN customers c_mg_sale ON 1=0"
        )
        if has_mg_schema and has_mg_sale_extended:
            mg_select_sql = """
              COALESCE(d.mg_estado,'activo') AS mg_estado,
              d.mg_inactivo_desde AS mg_inactivo_desde,
              d.mg_venta_fecha AS mg_venta_fecha,
              COALESCE(d.mg_venta_factura_numero,'') AS mg_venta_factura_numero,
              COALESCE(d.mg_venta_remito_numero,'') AS mg_venta_remito_numero,
              COALESCE(d.mg_venta_observaciones,'') AS mg_venta_observaciones,
              d.mg_venta_usuario_id AS mg_venta_usuario_id,
              d.mg_venta_customer_id AS mg_venta_customer_id,
              COALESCE(c_mg_sale.razon_social,'') AS mg_venta_customer_nombre,
              COALESCE(d.mg_venta_numero_alternativo,'') AS mg_venta_numero_alternativo,
            """
        else:
            mg_select_sql = """
              'activo' AS mg_estado,
              NULL AS mg_inactivo_desde,
              NULL AS mg_venta_fecha,
              '' AS mg_venta_factura_numero,
              '' AS mg_venta_remito_numero,
              '' AS mg_venta_observaciones,
              NULL AS mg_venta_usuario_id,
              NULL AS mg_venta_customer_id,
              '' AS mg_venta_customer_nombre,
              '' AS mg_venta_numero_alternativo,
            """
        if strong:
            # Pequeño delay best-effort para evitar carrera read-after-write
            try:
                time.sleep(0.12)
            except Exception:
                pass
        empresa_bejerman_sql = (
            "COALESCE(t.empresa_bejerman, 'SEPID') AS empresa_bejerman,"
            if _has_table_column("ingresos", "empresa_bejerman")
            else "'SEPID' AS empresa_bejerman,"
        )
        empresa_facturar_sql = (
            "COALESCE(t.empresa_facturar, 'SEPID') AS empresa_facturar,"
            if _has_table_column("ingresos", "empresa_facturar")
            else "'SEPID' AS empresa_facturar,"
        )
        row = q(
            f"""
            SELECT
              t.id,
              t.motivo,
              {empresa_bejerman_sql}
              {empresa_facturar_sql}
              {permite_reparacion_sql}
              t.estado,
              t.presupuesto_estado,
              t.resolucion,
              t.fecha_ingreso,
              t.fecha_creacion,
              t.fecha_servicio,
              t.garantia_reparacion,
              COALESCE(t.etiq_garantia_ok, true) AS etiq_garantia_ok,
              t.faja_garantia,
              t.remito_ingreso,
              t.remito_salida,
              t.factura_numero,
              t.fecha_entrega,
              t.alquilado,
              t.alquiler_a,
              t.alquiler_remito,
              t.alquiler_fecha,
              t.informe_preliminar,
              t.descripcion_problema,
              t.trabajos_realizados,
              t.comentarios,
              t.accesorios,
              t.equipo_variante,
              t.ubicacion_id,
              COALESCE(l.nombre,'') AS ubicacion_nombre,
              t.asignado_a,
              COALESCE(u.nombre,'') AS asignado_a_nombre,
              t.recibido_por AS ingresado_por_id,
              COALESCE(ur.nombre,'') AS ingresado_por_nombre,
              COALESCE(d.propietario_nombre, d.propietario) AS propietario_nombre,
              COALESCE(d.propietario_contacto, '') AS propietario_contacto,
              COALESCE(d.propietario_doc, '') AS propietario_doc,
              d.id AS device_id,
              COALESCE(d.numero_serie,'') AS numero_serie,
              COALESCE(d.numero_interno,'') AS numero_interno,
              {mg_select_sql}
              COALESCE(t.garantia_fabrica,false) AS garantia,
              d.marca_id,
              COALESCE(b.nombre,'') AS marca,
              d.model_id,
              COALESCE(m.nombre,'') AS modelo,
              COALESCE(m.tipo_equipo,'') AS tipo_equipo,
              c.id AS customer_id,
              c.razon_social,
              c.cod_empresa,
              c.telefono
            FROM ingresos t
            JOIN devices d ON d.id = t.device_id
            JOIN customers c ON c.id = d.customer_id
            LEFT JOIN marcas b ON b.id = d.marca_id
            LEFT JOIN models m ON m.id = d.model_id
            LEFT JOIN locations l ON l.id = t.ubicacion_id
            LEFT JOIN users u ON u.id = t.asignado_a
            LEFT JOIN users ur ON ur.id = t.recibido_por
            {mg_sale_customer_join}
            WHERE t.id = %s
            """,
            [ingreso_id],
            one=True,
        )

        if not row:
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        # Snapshot inmediato en devices (fase 2: fuente de verdad)
        try:
            exec_void(
                "UPDATE devices SET propietario_nombre = NULLIF(%s,''), propietario_contacto = NULLIF(%s,''), propietario_doc = NULLIF(%s,'') WHERE id = %s",
                [
                    row.get("propietario_nombre") or "",
                    row.get("propietario_contacto") or "",
                    row.get("propietario_doc") or "",
                    row.get("device_id"),
                ],
            )
        except Exception:
            pass
        mg_owner_id = _mg_owner_customer_id()
        row.update(resolve_mg_flags(row, mg_owner_id))
        row["es_cliente_mg_owner"] = bool(row.get("es_propietario_mg"))
        row["ris"] = get_ris_status_for_ingreso(ingreso_id)
        remito_salida_bejerman = _service_release_remito_salida_for_ingreso(ingreso_id)
        if remito_salida_bejerman and not (row.get("remito_salida") or "").strip():
            try:
                exec_void(
                    """
                    UPDATE ingresos
                       SET remito_salida = %s
                     WHERE id = %s
                       AND (remito_salida IS NULL OR TRIM(remito_salida) = '')
                    """,
                    [remito_salida_bejerman, ingreso_id],
                )
            except Exception:
                pass
            row["remito_salida"] = remito_salida_bejerman
        row["remito_salida_bejerman"] = remito_salida_bejerman
        # Agregar serial_cambio si existe la columna
        try:
            with connection.cursor() as cur:
                cur.execute(
                    """
                    SELECT 1 FROM information_schema.columns
                     WHERE table_name='ingresos' AND column_name='serial_cambio'
                       AND table_schema = ANY(current_schemas(true))
                     LIMIT 1
                    """
                )
                if cur.fetchone() is not None:
                    cur.execute("SELECT serial_cambio FROM ingresos WHERE id=%s", [row["id"]])
                    sc = cur.fetchone()
                    row["serial_cambio"] = (sc[0] if sc else None)
        except Exception:
            pass
        # Datos completos del cambio de equipo: qué se entregó y qué destino tiene el que
        # quedó acá. `serial_cambio` de arriba queda como snapshot para los cambios viejos.
        try:
            row["cambio_equipo"] = get_ingreso_device_swap(row["id"])
        except Exception:
            logger.exception("No se pudo leer el cambio de equipo", extra={"ingreso_id": row.get("id")})
            row["cambio_equipo"] = None
        try:
            if has_rejected_budget_charge_schema():
                row.update(get_stored_rejected_budget_fields(row["id"]))
            else:
                row["presupuesto_rechazado_cobro_neto"] = None
                row["presupuesto_rechazado_quote_id"] = None
        except Exception:
            row["presupuesto_rechazado_cobro_neto"] = None
            row["presupuesto_rechazado_quote_id"] = None
        # Si está en garantía de reparación, traer trabajos realizados del último servicio entregado
        if row.get("garantia_reparacion"):
            try:
                salida_expr = _salida_taller_expr("t")
                # "El último servicio" es la última REPARACIÓN anterior a ésta. Se saltean los
                # alquileres, ventas, bajas y los "no se reparó": el técnico mira este texto
                # para decidir si la falla nueva la cubre el trabajo anterior, y un ingreso sin
                # reparación no le dice nada.
                # El orden va por cuándo entró el ingreso y no por cuándo se registró la salida:
                # una entrega cargada tarde no puede convertir a un servicio posterior en "el
                # anterior".
                orden_t = "COALESCE(t.fecha_ingreso, t.fecha_creacion)"
                orden_cur = "COALESCE(cur.fecha_ingreso, cur.fecha_creacion)"
                prev = q(
                    f"""
                    SELECT t.id, t.trabajos_realizados
                      FROM ingresos t
                      JOIN ingresos cur ON cur.id = %s
                      {_salida_taller_lateral_sql("t")}
                     WHERE t.device_id = cur.device_id
                       AND t.id <> cur.id
                       AND t.resolucion::text = 'reparado'
                       AND {salida_expr} IS NOT NULL
                       AND (
                            {orden_t} < {orden_cur}
                         OR ({orden_t} = {orden_cur} AND t.id < cur.id)
                       )
                     ORDER BY {orden_t} DESC, t.id DESC
                     LIMIT 1
                    """,
                    [row.get("id")],
                    one=True,
                )
            except Exception:
                prev = None
            if prev:
                row["garantia_reparacion_trabajos"] = prev.get("trabajos_realizados")
            else:
                row["garantia_reparacion_trabajos"] = None
        row["os"] = os_label(row["id"])
        if not row.get("fecha_creacion"):
            row["fecha_creacion"] = row.get("fecha_ingreso")
        if "tipo_equipo_nombre" not in row:
            row["tipo_equipo_nombre"] = row.get("tipo_equipo") or ""
        try:
            from datetime import date as _date, datetime as _dt
            fs = row.get("fecha_servicio")
            if fs and isinstance(fs, _date) and not isinstance(fs, _dt):
                row["fecha_servicio"] = _dt(fs.year, fs.month, fs.day, 0, 0, 0, tzinfo=None)
        except Exception:
            pass
        accs = q(
            """
          SELECT ia.id, ia.accesorio_id, ca.nombre AS accesorio_nombre, ia.referencia, ia.descripcion
          FROM ingreso_accesorios ia
          JOIN catalogo_accesorios ca ON ca.id = ia.accesorio_id AND ca.activo
          WHERE ia.ingreso_id=%s
          ORDER BY ia.id
            """,
            [ingreso_id],
        )
        row["accesorios_items"] = accs
        # Accesorios vinculados a alquiler
        try:
            accs_alq = q(
                """
              SELECT ia.id, ia.accesorio_id, ca.nombre AS accesorio_nombre, ia.referencia, ia.descripcion
              FROM ingreso_alquiler_accesorios ia
              JOIN catalogo_accesorios ca ON ca.id = ia.accesorio_id AND ca.activo
              WHERE ia.ingreso_id=%s
              ORDER BY ia.id
                """,
                [ingreso_id],
            )
        except Exception:
            accs_alq = []
        row["alquiler_accesorios_items"] = accs_alq
        # Resolver tecnico solicitado (ultima solicitud pendiente si existe)
        req = None
        if _table_exists("ingreso_assignment_requests"):
            try:
                req = q(
                    """
                    SELECT r.usuario_id AS uid, COALESCE(u.nombre,'') AS nombre
                      FROM ingreso_assignment_requests r
                      LEFT JOIN users u ON u.id = r.usuario_id
                     WHERE r.ingreso_id = %s
                       AND r.canceled_at IS NULL
                       AND (r.accepted_at IS NULL)
                     ORDER BY r.created_at DESC, r.id DESC
                     LIMIT 1
                    """,
                    [ingreso_id],
                    one=True,
                )
            except Exception:
                req = None
        if not req and _table_exists("audit_log"):
            # Fallback: usar audit_log si esta habilitada/cargada
            try:
                path1 = f"/api/ingresos/{ingreso_id}/solicitar-asignacion/"
                path2 = f"/api/ingresos/{ingreso_id}/solicitar-asignacion"
                req = q(
                    """
                    SELECT al.user_id AS uid, COALESCE(u.nombre,'') AS nombre
                      FROM audit_log al
                      LEFT JOIN users u ON u.id = al.user_id
                     WHERE al.method = 'POST'
                       AND (al.path = %s OR al.path = %s)
                     ORDER BY al.ts DESC, al.id DESC
                     LIMIT 1
                    """,
                    [path1, path2],
                    one=True,
                )
            except Exception:
                req = None
        if req:
            row["tecnico_solicitado_id"] = req.get("uid")
            row["tecnico_solicitado_nombre"] = req.get("nombre") or ""
        else:
            row["tecnico_solicitado_id"] = None
            row["tecnico_solicitado_nombre"] = ""
        baja_req = _pending_baja_request(ingreso_id)
        if baja_req:
            row["baja_solicitada_id"] = baja_req.get("id")
            row["baja_solicitada_nombre"] = baja_req.get("nombre") or ""
            row["baja_solicitada_motivo"] = baja_req.get("motivo") or ""
            row["baja_solicitada_fecha"] = baja_req.get("created_at")
        else:
            row["baja_solicitada_id"] = None
            row["baja_solicitada_nombre"] = ""
            row["baja_solicitada_motivo"] = ""
            row["baja_solicitada_fecha"] = None
        event_hitos = _ingreso_event_date_hitos(ingreso_id)
        quote_hitos = _ingreso_quote_date_hitos(ingreso_id)
        row["fechas_hitos"] = {
            "fecha_creacion": row.get("fecha_creacion"),
            "fecha_ingreso": row.get("fecha_ingreso"),
            # `fecha_servicio` se auto-setea en la transición ingresado -> diagnosticado y es
            # editable a mano en la pestaña Diagnóstico: es el dato que el taller realmente
            # carga, así que manda sobre el evento cuando existe.
            "fecha_diagnosticado": row.get("fecha_servicio") or event_hitos.get("fecha_diagnosticado"),
            "fecha_presupuestado": quote_hitos.get("fecha_presupuestado"),
            "fecha_reparado": event_hitos.get("fecha_reparado"),
            "fecha_liberacion": event_hitos.get("fecha_liberacion"),
            "fecha_servicio": row.get("fecha_servicio"),
            "fecha_entrega": row.get("fecha_entrega"),
            "fecha_alquiler": row.get("alquiler_fecha"),
            "fecha_venta_mg": row.get("mg_venta_fecha"),
            "fecha_baja_solicitada": row.get("baja_solicitada_fecha"),
        }
        attach_active_derivations(row)
        return Response(IngresoDetailWithAccesoriosSerializer(row).data)

    def patch(self, request, ingreso_id: int):
        rol = _rol(request)
        d = request.data or {}
        _set_audit_user(request)
        sync_actor_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        permite_reparacion_sql = (
            "COALESCE(permite_reparacion, TRUE) AS permite_reparacion"
            if has_permite_reparacion
            else "TRUE AS permite_reparacion"
        )
        should_sync_preventivo = False
        basic_fields = {
            "propietario_nombre", "propietario_contacto", "propietario_doc",
            "customer_id", "cliente_id", "razon_social", "cod_empresa", "telefono",
            "numero_serie", "equipo_variante", "remito_ingreso", "informe_preliminar",
            "motivo", "garantia", "comentarios", "marca_id", "modelo_id",
        }
        diagnosis_fields = {
            "descripcion_problema", "trabajos_realizados", "fecha_servicio",
            "garantia_reparacion", "etiq_garantia_ok", "faja_garantia", "numero_interno",
        }
        location_fields = {
            "ubicacion_id", "accesorios", "alquilado", "alquiler_a", "alquiler_remito", "alquiler_fecha",
        }
        delivery_fields = {"remito_salida", "factura_numero", "fecha_entrega"}

        if any(k in d for k in delivery_fields):
            derivacion_activa = get_active_derivation(ingreso_id)
            if derivacion_activa:
                return Response(active_derivation_block_payload(derivacion_activa), status=409)

        if any(k in d for k in basic_fields):
            require_permission(request, "action.ingreso.edit_basics")
        if any(k in d for k in diagnosis_fields):
            require_permission(request, "action.ingreso.edit_diagnosis")
        if any(k in d for k in location_fields):
            require_permission(request, "action.ingreso.edit_location")
        if any(k in d for k in delivery_fields):
            require_permission(request, "action.ingreso.edit_delivery")
        if connection.vendor == "postgresql":
            exec_void("SET LOCAL app.ingreso_id = %s", [ingreso_id])

        row_est = q(
            f"""
            SELECT t.estado,
                   t.asignado_a,
                   t.alquilado,
                   t.motivo,
                   t.fecha_entrega,
                   t.alquiler_remito,
                   t.remito_salida,
                   COALESCE(loc.nombre, '') AS ubicacion_nombre,
                   {permite_reparacion_sql}
              FROM ingresos t
              LEFT JOIN locations loc ON loc.id=t.ubicacion_id
             WHERE t.id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row_est:
            return Response({"detail": "Ingreso no encontrado"}, status=404)
        estado_actual = (row_est["estado"] or "").lower()
        asignado_a = row_est["asignado_a"]
        alquilado_actual = bool(row_est["alquilado"])
        motivo_actual = row_est.get("motivo")
        fecha_entrega_actual = row_est.get("fecha_entrega")
        alquiler_remito_actual = (row_est.get("alquiler_remito") or "").strip()
        alquiler_con_salida_confirmada = bool(alquilado_actual and alquiler_remito_actual)

        if alquiler_con_salida_confirmada and any(
            key in d for key in ("alquilado", "alquiler_a", "alquiler_remito", "alquiler_fecha")
        ):
            cambios = {
                "alquilado": bool(d.get("alquilado")),
                "alquiler_remito": (d.get("alquiler_remito") or "").strip(),
            }
            if (
                ("alquilado" in d and not cambios["alquilado"])
                or ("alquiler_remito" in d and cambios["alquiler_remito"] != alquiler_remito_actual)
                or "alquiler_a" in d
                or "alquiler_fecha" in d
            ):
                return Response(
                    {"detail": "No se puede modificar el alquiler: el equipo ya fue entregado con RTA."},
                    status=409,
                )

        sets_no_estado, params_no_estado = [], []
        needs_warranty_recompute = False
        fecha_entrega_nueva = fecha_entrega_actual

        rental_shelf_requested = False
        rental_shelf_applies = False
        location_sync_plan = None
        location_sync_result = None
        if "ubicacion_id" in d:
            ubicacion_id = d.get("ubicacion_id")
            if not ubicacion_id:
                raise ValidationError("ubicacion_id requerido")
            u = q("SELECT id, nombre FROM locations WHERE id=%s", [ubicacion_id], one=True)
            if not u:
                raise ValidationError("Ubicación inexistente")
            rental_shelf_requested = is_rental_shelf_location_name(u.get("nombre"))
            source_location = (row_est.get("ubicacion_nombre") or "").strip()
            target_location = (u.get("nombre") or "").strip()
            source_deposit = bejerman_deposit_for_location(source_location)
            target_deposit = bejerman_deposit_for_location(target_location)
            if rental_shelf_requested:
                rental_shelf_applies = ingreso_targets_rental_stock(ingreso_id)

            # Sólo las ubicaciones físicas con origen y destino distintos generan
            # movimiento. El guion y las ubicaciones sin mapeo Bejerman siguen siendo
            # referencias locales, por lo que no se les inventa un depósito de origen.
            if source_deposit and target_deposit and source_deposit != target_deposit:
                if rental_shelf_requested:
                    if rental_shelf_applies:
                        location_sync_plan = {
                            "kind": "rental_shelf",
                            "source_deposit": source_deposit,
                            "target_deposit": target_deposit,
                            "source_location": source_location,
                            "target_location": target_location,
                        }
                else:
                    location_sync_plan = {
                        "kind": "location_transfer",
                        "source_deposit": source_deposit,
                        "target_deposit": target_deposit,
                        "source_location": source_location,
                        "target_location": target_location,
                    }

            confirmed = d.get("confirm_bejerman_location_change") is True or str(
                d.get("confirm_bejerman_location_change") or ""
            ).strip().lower() in {"1", "true", "si", "sí"}
            if location_sync_plan and not confirmed:
                return Response(
                    {
                        "detail": (
                            "Confirmá el cambio de ubicación: también solicita un traslado de stock en Bejerman."
                        ),
                        "confirmation_required": True,
                        "source_deposit": location_sync_plan["source_deposit"],
                        "target_deposit": location_sync_plan["target_deposit"],
                    },
                    status=400,
                )
            sets_no_estado.append("ubicacion_id=%s")
            params_no_estado.append(ubicacion_id)

        if any(key in d for key in ('descripcion_problema','trabajos_realizados')):
            from .. import technical_interventions as technical
            if technical.available() and q("SELECT ingreso_id FROM technical_interventions WHERE ingreso_id=%s", [ingreso_id], one=True):
                if 'trabajos_realizados' in d:
                    raise ValidationError("Los trabajos se registran mediante la lista de la intervención técnica.")
                old_diagnosis = q("SELECT descripcion_problema FROM ingresos WHERE id=%s FOR UPDATE", [ingreso_id], one=True)
                if (old_diagnosis['descripcion_problema'] or '') != str(d.get('descripcion_problema') or '').strip():
                    exec_void("UPDATE technical_interventions SET revision=revision+1, scope_revision=scope_revision+1 WHERE ingreso_id=%s", [ingreso_id])
        desc_present = False
        trab_present = False
        fecha_present = False
        if "descripcion_problema" in d:
            desc = (d.get("descripcion_problema") or "").strip()
            desc_present = bool(desc)
            sets_no_estado.append("descripcion_problema=%s")
            params_no_estado.append(desc)
        if "trabajos_realizados" in d:
            trab_present = bool((d.get("trabajos_realizados") or "").strip())
            sets_no_estado.append("trabajos_realizados=%s")
            params_no_estado.append(d.get("trabajos_realizados"))
        if "accesorios" in d:
            sets_no_estado.append("accesorios=%s")
            params_no_estado.append(d.get("accesorios"))
        if "fecha_servicio" in d:
            should_sync_preventivo = True
            val = d.get("fecha_servicio")
            if val is None or (isinstance(val, str) and val.strip() == ""):
                sets_no_estado.append("fecha_servicio=NULL")
            else:
                dt = parse_datetime(val)
                if not dt:
                    raise ValidationError("fecha_servicio inválida")
                if timezone.is_naive(dt):
                    dt = timezone.make_aware(dt, timezone.get_current_timezone())
                sets_no_estado.append("fecha_servicio=%s")
                params_no_estado.append(dt)
                fecha_present = True
        # Si el rol es tecnico/jefe_veedor, solo el asignado puede editar diagnostico/trabajos/fecha
        if rol in ("tecnico", "jefe_veedor") and any(k in d for k in ("descripcion_problema", "trabajos_realizados", "fecha_servicio")):
            uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
            if int(asignado_a or 0) != int(uid or 0):
                raise PermissionDenied("Solo el técnico asignado puede editar diagnóstico y reparación")
        if any(k in d for k in ("remito_salida", "factura_numero", "fecha_entrega")):
            if "remito_salida" in d:
                sets_no_estado.append("remito_salida = NULLIF(%s,'')")
                params_no_estado.append((d.get("remito_salida") or "").strip())
            if "factura_numero" in d:
                sets_no_estado.append("factura_numero = NULLIF(%s,'')")
                params_no_estado.append((d.get("factura_numero") or "").strip())
            if "fecha_entrega" in d:
                val = d.get("fecha_entrega")
                if val is None or (isinstance(val, str) and val.strip() == ""):
                    fecha_entrega_nueva = None
                    sets_no_estado.append("fecha_entrega=NULL")
                else:
                    dt = parse_datetime(val)
                    if not dt:
                        raise ValidationError("fecha_entrega inválida")
                    if timezone.is_naive(dt):
                        dt = timezone.make_aware(dt, timezone.get_current_timezone())
                    fecha_entrega_nueva = dt
                    sets_no_estado.append("fecha_entrega=%s")
                    params_no_estado.append(dt)
        if "garantia_reparacion" in d:
            sets_no_estado.append("garantia_reparacion=%s")
            params_no_estado.append(bool(d.get("garantia_reparacion")))
        if "etiq_garantia_ok" in d:
            sets_no_estado.append("etiq_garantia_ok=%s")
            params_no_estado.append(bool(d.get("etiq_garantia_ok")))
        if "faja_garantia" in d:
            sets_no_estado.append("faja_garantia = NULLIF(%s,'')")
            params_no_estado.append((d.get("faja_garantia") or "").strip())
        if "numero_interno" in d:
            val = normalize_mg_input(d.get("numero_interno"))
            try:
                dev_row = q("SELECT device_id FROM ingresos WHERE id=%s", [ingreso_id], one=True)
                device_id = dev_row and dev_row.get("device_id")
            except Exception:
                device_id = None
            if val and device_id:
                # Verificar conflicto con otro device por MG normalizado
                conflict = q(
                    f"""
                    SELECT id
                      FROM devices
                     WHERE id <> %s
                       AND {sql_is_mg_code('numero_interno')}
                       AND {sql_mg_key('numero_interno')} = {sql_mg_key('%s')}
                     LIMIT 1
                    """,
                    [device_id, val],
                    one=True,
                )
                if conflict:
                    return Response(
                        {
                            "detail": "El número interno ya está asignado a otro equipo.",
                            "conflict_type": "MG_DUPLICATE",
                            "conflict_device_id": conflict["id"],
                            "current_device_id": device_id,
                            "numero_interno_input": val,
                        },
                        status=400,
                    )
            try:
                exec_void(
                    "UPDATE devices SET numero_interno = NULLIF(%s,'') WHERE id = (SELECT device_id FROM ingresos WHERE id=%s)",
                    [val, ingreso_id]
                )
            except Exception as e:
                return Response(
                    {
                        "detail": "No se pudo asignar el número interno (posible duplicado).",
                        "conflict_type": "MG_UNIQUE_CONSTRAINT",
                        "current_device_id": device_id,
                        "numero_interno_input": val,
                        "error": str(e),
                    },
                    status=400,
                )
        if any(k in d for k in ("propietario_nombre", "propietario_contacto", "propietario_doc")):
            if "propietario_nombre" in d:
                sets_no_estado.append("propietario_nombre = NULLIF(%s,'')")
                params_no_estado.append((d.get("propietario_nombre") or "").strip())
            if "propietario_contacto" in d:
                sets_no_estado.append("propietario_contacto = NULLIF(%s,'')")
                params_no_estado.append((d.get("propietario_contacto") or "").strip())
            if "propietario_doc" in d:
                sets_no_estado.append("propietario_doc = NULLIF(%s,'')")
                params_no_estado.append((d.get("propietario_doc") or "").strip())
            # Fase 2: persistir también en devices
            try:
                exec_void(
                    "UPDATE devices SET propietario_nombre = NULLIF(%s,''), propietario_contacto = NULLIF(%s,''), propietario_doc = NULLIF(%s,'') WHERE id = (SELECT device_id FROM ingresos WHERE id=%s)",
                    [
                        (d.get("propietario_nombre") or "").strip(),
                        (d.get("propietario_contacto") or "").strip(),
                        (d.get("propietario_doc") or "").strip(),
                        ingreso_id,
                    ],
                )
            except Exception:
                pass
        if "customer_id" in d or "cliente_id" in d:
            try:
                new_cid = int(d.get("customer_id") or d.get("cliente_id"))
            except Exception:
                return Response({"detail": "customer_id inválido"}, status=400)
            if new_cid <= 0:
                return Response({"detail": "customer_id inválido"}, status=400)
            c_row = q("SELECT id FROM customers WHERE id=%s", [new_cid], one=True)
            if not c_row:
                return Response({"detail": "Cliente inexistente"}, status=404)
            exec_void(
                "UPDATE devices SET customer_id=%s WHERE id = (SELECT device_id FROM ingresos WHERE id=%s)",
                [new_cid, ingreso_id],
            )
        if any(k in d for k in ("razon_social", "cod_empresa", "telefono")):
            rs = d.get("razon_social")
            ce = d.get("cod_empresa")
            tel = d.get("telefono")
            sets, params = [], []
            if rs is not None:
                sets.append("razon_social = NULLIF(%s,'')")
                params.append((rs or "").strip())
            if ce is not None:
                sets.append("cod_empresa = NULLIF(%s,'')")
                params.append((ce or "").strip())
            if tel is not None:
                sets.append("telefono = NULLIF(%s,'')")
                params.append((tel or "").strip())
            if sets:
                params.append(ingreso_id)
                exec_void(
                    f"""
                    UPDATE customers
                       SET {', '.join(sets)}
                     WHERE id = (
                        SELECT d.customer_id FROM devices d
                        JOIN ingresos t ON t.device_id = d.id
                        WHERE t.id=%s
                     )
                    """,
                    params,
                )
        if "numero_serie" in d:
            ns = (d.get("numero_serie") or "").strip()
            needs_warranty_recompute = True
            ns_key = (ns or "").replace(" ", "").replace("-", "").upper()
            # Reasociar el ingreso al device que ya tenga ese N/S (normalizado) **y el mismo
            # tipo de equipo**: la serie sola no identifica al equipo desde que el
            # humidificador entra con la del CPAP. Sin el tipo, corregir la serie del
            # humidificador lo colgaba del CPAP o daba "ya está asignado a otro equipo".
            if ns_key:
                row_target = q(
                    """
                    SELECT d.id FROM devices d
                     LEFT JOIN models m ON m.id = d.model_id
                     WHERE REPLACE(REPLACE(UPPER(d.numero_serie),' ','') ,'-','') = %s
                       AND LOWER(TRIM(COALESCE(NULLIF(d.tipo_equipo,''), m.tipo_equipo, ''))) = (
                             SELECT LOWER(TRIM(COALESCE(NULLIF(dd.tipo_equipo,''), mm.tipo_equipo, '')))
                               FROM ingresos tt
                               JOIN devices dd ON dd.id = tt.device_id
                          LEFT JOIN models mm ON mm.id = dd.model_id
                              WHERE tt.id = %s
                           )
                     LIMIT 1
                    """,
                    [ns_key, ingreso_id],
                    one=True,
                )
                if row_target:
                    target_id = int(row_target["id"])
                    row_cur = q("SELECT device_id FROM ingresos WHERE id=%s", [ingreso_id], one=True)
                    if not row_cur:
                        return Response({"detail": "Ingreso no encontrado"}, status=404)
                    cur_dev_id = int(row_cur["device_id"])
                    if cur_dev_id != target_id:
                        return Response(
                            {
                                "detail": "El número de serie ya está asignado a otro equipo.",
                                "conflict_type": "NS_DUPLICATE",
                                "conflict_device_id": target_id,
                                "current_device_id": cur_dev_id,
                                "numero_serie_input": ns,
                            },
                            status=400,
                        )
                    # No modificar numero_serie en devices si ya existe ese device
                else:
                    # Fallback: si no existe device con ese N/S, actualizar el N/S del device actual
                    exec_void(
                        "UPDATE devices SET numero_serie = NULLIF(%s,'') WHERE id = (SELECT device_id FROM ingresos WHERE id=%s)",
                        [ns, ingreso_id],
                    )
            else:
                # Si viene vacío, limpiar N/S del device actual
                exec_void(
                    "UPDATE devices SET numero_serie = NULLIF(%s,'') WHERE id = (SELECT device_id FROM ingresos WHERE id=%s)",
                    [ns, ingreso_id],
                )
            # Si no vino 'garantia' explícito, auto-chequear por N/S
            # removed: auto-chequeo de garantia por N/S (fuente: ultimo ingreso)
                # (bloque removido)
        # Permitir actualizar variante de equipo (texto libre)
        if "equipo_variante" in d:
            sets_no_estado.append("equipo_variante = NULLIF(%s,'')")
            params_no_estado.append((d.get("equipo_variante") or "").strip())
        if "remito_ingreso" in d:
            sets_no_estado.append("remito_ingreso = NULLIF(%s,'')")
            params_no_estado.append((d.get("remito_ingreso") or "").strip())
        if "informe_preliminar" in d:
            sets_no_estado.append("informe_preliminar = NULLIF(%s,'')")
            params_no_estado.append((d.get("informe_preliminar") or "").strip())
        if "motivo" in d:
            motivo_raw = (d.get("motivo") or "").strip()
            if not motivo_raw:
                raise ValidationError("motivo requerido")
            motivo_label_raw = _map_motivo_to_db_label(motivo_raw)
            if not motivo_label_raw:
                valid_motivos = _get_motivo_enum_values()
                return Response({"detail": "motivo inválido", "valid_values": valid_motivos}, status=400)
            raw_vals = _get_motivo_enum_values_raw() or []
            if raw_vals:
                target = None
                norm_target = _norm_txt(motivo_label_raw)
                for rv in raw_vals:
                    if _norm_txt(rv) == norm_target:
                        target = rv
                        break
                if not target:
                    target = next((x for x in raw_vals if _norm_txt(x) == _norm_txt("otros")), raw_vals[0])
                motivo_db = target
            else:
                motivo_db = motivo_label_raw
            sets_no_estado.append("motivo=%s")
            params_no_estado.append(motivo_db)
            if has_permite_reparacion:
                nuevo_es_cotizacion = _motivo_is_cotizacion_equipo(motivo_db)
                anterior_es_cotizacion = _motivo_is_cotizacion_equipo(motivo_actual)
                if nuevo_es_cotizacion and not anterior_es_cotizacion:
                    sets_no_estado.append("permite_reparacion=FALSE")
                elif (not nuevo_es_cotizacion) and anterior_es_cotizacion:
                    sets_no_estado.append("permite_reparacion=TRUE")
        # (actualizacion de garantia en devices removida)
            # bloque removido





        # Actualizar garantia de fábrica a nivel de ingreso cuando se envía en el PATCH
        if "garantia" in d:
            sets_no_estado.append("garantia_fabrica=%s")
            params_no_estado.append(bool(d.get("garantia")))
        if "comentarios" in d:
            sets_no_estado.append("comentarios = NULLIF(%s,'')")
            params_no_estado.append((d.get("comentarios") or "").strip())
        # Actualizar marca/modelo del dispositivo asociado al ingreso
        if ("marca_id" in d) or ("modelo_id" in d):
            needs_warranty_recompute = True
            marca_id = d.get("marca_id")
            modelo_id = d.get("modelo_id")
            # Traer device actual
            cur_dev = q("SELECT device_id FROM ingresos WHERE id=%s", [ingreso_id], one=True)
            if not cur_dev:
                return Response({"detail": "Ingreso no encontrado"}, status=404)
            device_id = cur_dev["device_id"]
            if modelo_id is not None:
                # Validar modelo y obtener su marca
                md = q("SELECT id, marca_id FROM models WHERE id=%s", [modelo_id], one=True)
                if not md:
                    raise ValidationError("modelo_id inválido")
                modelo_marca = md["marca_id"]
                if marca_id is not None and int(marca_id) != int(modelo_marca):
                    raise ValidationError("modelo_id no pertenece a la marca_id indicada")
                # Ajustar marca al del modelo si no fue provista
                exec_void("UPDATE devices SET model_id=%s, marca_id=%s WHERE id=%s", [modelo_id, modelo_marca, device_id])
            elif marca_id is not None:
                # Validar marca
                mk = q("SELECT id FROM marcas WHERE id=%s", [marca_id], one=True)
                if not mk:
                    raise ValidationError("marca_id inválida")
                # Si el modelo actual no corresponde a la nueva marca, ponerlo en NULL
                row = q("SELECT model_id FROM devices WHERE id=%s", [device_id], one=True)
                cur_model = row and row.get("model_id")
                if cur_model:
                    ok = q("SELECT 1 FROM models WHERE id=%s AND marca_id=%s", [cur_model, marca_id], one=True)
                    if not ok:
                        exec_void("UPDATE devices SET marca_id=%s, model_id=NULL WHERE id=%s", [marca_id, device_id])
                    else:
                        exec_void("UPDATE devices SET marca_id=%s WHERE id=%s", [marca_id, device_id])
        if "alquilado" in d:
            new_alquilado = bool(d.get("alquilado"))
            if alquilado_actual and not new_alquilado and not user_has_permission(
                request, "action.ingreso.uncheck_rental"
            ):
                raise PermissionDenied("No tenés permiso para destildar el alquiler")
            if alquilado_actual and not new_alquilado:
                alquiler_remito_bloqueante = alquiler_remito_actual or (d.get("alquiler_remito") or "").strip()
                entregado_bloqueante = bool(estado_actual == "entregado" or fecha_entrega_nueva)
                if alquiler_remito_bloqueante and entregado_bloqueante:
                    raise ValidationError(
                        "No se puede destildar alquiler: el equipo ya fue entregado con remito de alquiler."
                    )
            sets_no_estado.append("alquilado=%s")
            params_no_estado.append(new_alquilado)
            try:
                if new_alquilado:
                    sets_no_estado.append("estado='alquilado'")
                    dash_id = _dash_location_id()
                    if dash_id and not any(str(s).strip().startswith("ubicacion_id=") for s in sets_no_estado):
                        sets_no_estado.append("ubicacion_id=%s")
                        params_no_estado.append(dash_id)
            except Exception:
                pass
        if "alquiler_a" in d:
            sets_no_estado.append("alquiler_a=NULLIF(%s,'')")
            params_no_estado.append((d.get("alquiler_a") or "").strip())
        if "alquiler_remito" in d:
            alquiler_remito_nuevo = (d.get("alquiler_remito") or "").strip()
            sets_no_estado.append("alquiler_remito=NULLIF(%s,'')")
            params_no_estado.append(alquiler_remito_nuevo)
            # El RTA es un remito de salida: mantener ambas proyecciones evita que la
            # cabecera de la hoja oculte un equipo que ya está con el cliente.
            if alquiler_remito_nuevo.upper().startswith("RTA"):
                sets_no_estado.append("remito_salida=COALESCE(NULLIF(remito_salida,''), %s)")
                params_no_estado.append(alquiler_remito_nuevo)
        if "alquiler_fecha" in d:
            sets_no_estado.append("alquiler_fecha=%s")
            params_no_estado.append(d.get("alquiler_fecha") or None)

        # Recalcular garantía si cambió N/S o marca/modelo
        if needs_warranty_recompute:
            try:
                row_dev = q(
                    """
                    SELECT d.id, d.numero_serie, d.marca_id, d.model_id
                      FROM devices d
                      JOIN ingresos t ON t.device_id = d.id
                     WHERE t.id=%s
                    """,
                    [ingreso_id],
                    one=True,
                )
                if row_dev:
                    calc = _compute_factory_warranty_from_bejerman_sale(
                        (row_dev.get("numero_serie") or "").strip(),
                        brand_id=row_dev.get("marca_id"),
                        model_id=row_dev.get("model_id"),
                    )
                    if not any(str(s).strip().startswith("garantia_fabrica=") for s in sets_no_estado):
                        sets_no_estado.append("garantia_fabrica=%s")
                        params_no_estado.append(calc.get("garantia"))
                    vence = calc.get("vence_el")
                    if vence is None:
                        exec_void("UPDATE devices SET garantia_vence=NULL WHERE id=%s", [row_dev.get("id")])
                    else:
                        exec_void("UPDATE devices SET garantia_vence=%s WHERE id=%s", [vence, row_dev.get("id")])
            except Exception:
                pass

        def _with_location_sync(payload: dict) -> dict:
            """Adjunta el resultado de la sincronización ya encolada con el cambio."""
            if location_sync_result:
                payload["bejerman_location_sync"] = location_sync_result
            if rental_shelf_requested:
                payload["rental_shelf"] = {
                    "applies": bool(rental_shelf_applies),
                    "stock_queued": bool((location_sync_result or {}).get("queued")),
                    "stock_error": (location_sync_result or {}).get("error") or "",
                    "reason": (
                        ""
                        if rental_shelf_applies
                        else "El equipo no es del parque de alquiler: su stock no va a STL."
                    ),
                }
            return payload

        def _enqueue_location_sync() -> None:
            """Encola dentro de la misma transacción que persiste la ubicación."""
            nonlocal location_sync_result
            if not location_sync_plan:
                return
            if location_sync_plan["kind"] == "rental_shelf":
                job = enqueue_stock_transfer_for_ingreso(
                    ingreso_id,
                    actor_user_id=sync_actor_id,
                    source_deposit=location_sync_plan["source_deposit"],
                    requeue_succeeded=True,
                )
            else:
                job = enqueue_location_transfer_for_ingreso(
                    ingreso_id,
                    source_deposit=location_sync_plan["source_deposit"],
                    target_deposit=location_sync_plan["target_deposit"],
                    source_location=location_sync_plan["source_location"],
                    target_location=location_sync_plan["target_location"],
                    actor_user_id=sync_actor_id,
                )
            location_sync_result = {
                "queued": True,
                "job_id": job.get("id"),
                "source_deposit": location_sync_plan["source_deposit"],
                "target_deposit": location_sync_plan["target_deposit"],
                "source_location": location_sync_plan["source_location"],
                "target_location": location_sync_plan["target_location"],
            }

        diag_present = desc_present or trab_present or fecha_present
        promote_from_ingresado = diag_present and estado_actual == "ingresado"
        promote_from_asignado = diag_present and estado_actual == "asignado"
        if promote_from_ingresado:
            if not asignado_a:
                raise ValidationError("Antes de diagnosticar, asigne un técnico al ingreso")
            try:
                with transaction.atomic():
                    # Auto-setear fecha_servicio solo en la transición ingresado -> diagnosticado
                    # si no vino explícita en el payload y aún está NULL en DB.
                    sets_tmp = list(sets_no_estado)
                    params_tmp = list(params_no_estado)
                    if not fecha_present:
                        # Establecer solo si actualmente es NULL (COALESCE mantiene existente)
                        sets_tmp.append("fecha_servicio=COALESCE(fecha_servicio, %s)")
                        params_tmp.append(timezone.now())
                        should_sync_preventivo = True
                    if sets_tmp:
                        params_all = list(params_tmp) + [ingreso_id]
                        q(f"UPDATE ingresos SET {', '.join(sets_tmp)} WHERE id=%s", params_all)
                    q("UPDATE ingresos SET estado='diagnosticado' WHERE id=%s AND estado='ingresado'", [ingreso_id])
                    _enqueue_location_sync()
            except Exception as exc:
                if location_sync_plan:
                    return Response(
                        {"detail": f"No se pudo encolar el traslado de ubicación en Bejerman: {exc}"},
                        status=409,
                    )
                raise
            registrar_ingreso_event(
                ingreso_id,
                "diagnosticado",
                de_estado="ingresado",
                usuario_id=sync_actor_id,
                comentario="Diagnóstico cargado desde la hoja de servicio",
            )
            if should_sync_preventivo:
                sync_result = sync_plan_from_ingreso_fecha_servicio(ingreso_id, sync_actor_id)
                if (sync_result or {}).get("status") == "error":
                    logger.warning(
                        "preventivo sync returned error status after promote_from_ingresado: ingreso_id=%s",
                        ingreso_id,
                    )
            return Response(_with_location_sync({"ok": True}))
        if promote_from_asignado:
            sets_no_estado.append("estado='diagnosticado'")
        if not sets_no_estado:
            return Response({"ok": True})
        params_no_estado.append(ingreso_id)
        try:
            with transaction.atomic():
                q(f"UPDATE ingresos SET {', '.join(sets_no_estado)} WHERE id=%s", params_no_estado)
                _enqueue_location_sync()
        except Exception as exc:
            if location_sync_plan:
                return Response(
                    {"detail": f"No se pudo encolar el traslado de ubicación en Bejerman: {exc}"},
                    status=409,
                )
            raise
        if promote_from_asignado:
            registrar_ingreso_event(
                ingreso_id,
                "diagnosticado",
                de_estado="asignado",
                usuario_id=sync_actor_id,
                comentario="Diagnóstico cargado desde la hoja de servicio",
            )
        if should_sync_preventivo:
            sync_result = sync_plan_from_ingreso_fecha_servicio(ingreso_id, sync_actor_id)
            if (sync_result or {}).get("status") == "error":
                logger.warning(
                    "preventivo sync returned error status after ingreso patch: ingreso_id=%s",
                    ingreso_id,
                )
        return Response(_with_location_sync({"ok": True}))


# Nuevo: marcar controlado sin defecto (equipos propios revisados sin falla)
class MarcarControladoSinDefectoView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    def post(self, request, ingreso_id: int):
        require_roles(request, ["tecnico", "jefe", "jefe_veedor"])
        # Solo el técnico asignado puede marcar (salvo jefes)
        try:
            if _rol(request) in ("tecnico", "jefe_veedor"):
                row = q("SELECT asignado_a FROM ingresos WHERE id=%s", [ingreso_id], one=True)
                uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
                if not row or int(row.get("asignado_a") or 0) != int(uid or 0):
                    raise PermissionDenied("Solo el técnico asignado puede marcar como controlado")
        except PermissionDenied:
            raise
        except Exception:
            if _rol(request) in ("tecnico", "jefe_veedor"):
                raise PermissionDenied("Solo el técnico asignado puede marcar como controlado")

        _prev_estado = _estado_actual_txt(ingreso_id)
        from .. import technical_interventions as technical
        technical.mark_no_fault(ingreso_id, request.user)
        _set_audit_user(request)
        # El pase automático a 'no_aplica' solo aplica a los equipos propios de MG BIO
        # (patrimonio de la empresa / alquiler / demo). Para los equipos de clientes, no
        # se toca el presupuesto: queda en el estado que tenga (ej: 'pendiente') para
        # permitir cotizar la revisión técnica o definir su cobro manualmente.
        mg_ctx = _fetch_ingreso_device_mg_context(ingreso_id)
        es_propio_mg = bool(mg_ctx and mg_ctx.get("es_propietario_mg"))

        presupuesto_set = []
        if es_propio_mg:
            presupuesto_no_aplica = marcar_quote_no_aplica(ingreso_id)
            if presupuesto_no_aplica:
                presupuesto_set = ["presupuesto_estado='no_aplica'"]

        sets = sets_con_limpieza_de_cambio(
            "estado='controlado_sin_defecto'",
            "resolucion='no_se_encontro_falla'",
            *presupuesto_set,
        )
        exec_void(f"UPDATE ingresos SET {sets} WHERE id=%s", [ingreso_id])
        if _prev_estado != "controlado_sin_defecto":
            registrar_ingreso_event(
                ingreso_id,
                "controlado_sin_defecto",
                de_estado=_prev_estado,
                usuario_id=getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None),
                comentario="Marcado como controlado sin defecto",
            )

        # Revisado y sin falla es fin de taller igual que reparado: el equipo propio
        # vuelve al parque de alquiler.
        actor_user_id = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
        shelf = move_ingreso_to_rental_shelf(ingreso_id, actor_user_id=actor_user_id)
        if not shelf.get("applies"):
            try:
                if ingreso_is_demo_return(ingreso_id):
                    enqueue_demo_ready_transfer_for_ingreso(ingreso_id, actor_user_id=actor_user_id)
            except Exception:
                logger.exception("No se pudo encolar transferencia Bejerman STR a VAL", extra={"ingreso_id": ingreso_id})

        resp = {
            "ok": True,
            "estado": "controlado_sin_defecto",
            "presupuesto_estado": _presupuesto_estado_actual_txt(ingreso_id),
        }
        if shelf.get("moved"):
            location_name = shelf.get("location_name") or RENTAL_SHELF_LOCATION_NAME
            resp["auto_moved"] = True
            resp["auto_moved_to"] = location_name
            resp["ubicacion_id"] = shelf.get("location_id")
            resp["ubicacion_nombre"] = location_name
        return Response(resp)


def _latest_quote_is_approved(ingreso_id: int) -> bool:
    if not (
        _has_table("quotes")
        and _has_table_column("quotes", "id")
        and _has_table_column("quotes", "ingreso_id")
        and _has_table_column("quotes", "estado")
    ):
        return False
    order_sql = "COALESCE(version_num, 1) DESC, id DESC" if _has_table_column("quotes", "version_num") else "id DESC"
    try:
        row = q(
            f"""
            SELECT estado::text AS estado
              FROM quotes
             WHERE ingreso_id=%s
             ORDER BY {order_sql}
             LIMIT 1
            """,
            [ingreso_id],
            one=True,
        ) or {}
        return (row.get("estado") or "").strip().lower() == "aprobado"
    except Exception:
        return False


def _estado_previo_a_estado(ingreso_id: int, estado_cur: str, fallback: str = "diagnosticado") -> tuple[str, dict]:
    if connection.vendor != "postgresql":
        return fallback, {}
    try:
        row = q("SELECT to_regclass('audit.change_log') AS table_name", one=True) or {}
        if not row.get("table_name"):
            return fallback, {}
        params = [ingreso_id, estado_cur, list(RESTOREABLE_REPAIR_STATES)]
        estado_row = q(
            """
            SELECT id, ts, old_value, new_value
              FROM audit.change_log
             WHERE table_name='ingresos'
               AND record_id=%s
               AND column_name='estado'
               AND LOWER(COALESCE(new_value, ''))=%s
               AND LOWER(COALESCE(old_value, '')) = ANY(%s)
             ORDER BY ts DESC, id DESC
             LIMIT 1
            """,
            params,
            one=True,
        )
        if not estado_row:
            estado_row = q(
                """
                SELECT id, ts, old_value, new_value
                  FROM audit.change_log
                 WHERE table_name='ingresos'
                   AND record_id=%s
                   AND column_name='estado'
                   AND LOWER(COALESCE(new_value, '')) = ANY(%s)
                   AND LOWER(COALESCE(old_value, '')) = ANY(%s)
                 ORDER BY ts DESC, id DESC
                 LIMIT 1
                """,
                [ingreso_id, list(TECH_FINAL_REPAIR_STATES), list(RESTOREABLE_REPAIR_STATES)],
                one=True,
            )
        previous = (estado_row or {}).get("old_value")
        previous_lower = (previous or "").strip().lower()
        if previous_lower in RESTOREABLE_REPAIR_STATES:
            return previous_lower, estado_row or {}
    except Exception:
        return fallback, {}
    return fallback, {}


def _presupuesto_previo_en_misma_transicion(ingreso_id: int, estado_transition: dict, presupuesto_actual):
    if connection.vendor != "postgresql" or not estado_transition:
        return presupuesto_actual
    try:
        row = q("SELECT to_regclass('audit.change_log') AS table_name", one=True) or {}
        if not row.get("table_name"):
            return presupuesto_actual
        change_id = estado_transition.get("id")
        change_ts = estado_transition.get("ts")
        if not change_id or not change_ts:
            return presupuesto_actual
        presupuesto_row = q(
            """
            SELECT old_value
              FROM audit.change_log
             WHERE table_name='ingresos'
               AND record_id=%s
               AND column_name='presupuesto_estado'
               AND id BETWEEN %s - 10 AND %s + 10
               AND ts BETWEEN %s::timestamptz - interval '2 seconds'
                          AND %s::timestamptz + interval '2 seconds'
               AND COALESCE(new_value, '') = COALESCE(%s, '')
             ORDER BY ABS(id - %s) ASC
             LIMIT 1
            """,
            [ingreso_id, change_id, change_id, change_ts, change_ts, presupuesto_actual, change_id],
            one=True,
        )
        previous = (presupuesto_row or {}).get("old_value")
        return previous if previous not in (None, "") else presupuesto_actual
    except Exception:
        return presupuesto_actual


class QuitarEstadoFinalReparacionView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        require_roles(request, ["tecnico", "jefe", "jefe_veedor"])
        row = q(
            """
            SELECT
              estado::text AS estado,
              presupuesto_estado::text AS presupuesto_estado,
              asignado_a
            FROM ingresos
            WHERE id=%s
            """,
            [ingreso_id],
            one=True,
        )
        if not row:
            raise ValidationError("Ingreso no encontrado")

        estado_cur = (row.get("estado") or "").strip().lower()
        presupuesto_actual = row.get("presupuesto_estado")
        presupuesto_cur = (presupuesto_actual or "").strip().lower()
        if estado_cur in REPAIR_CLOSE_STATES:
            raise ValidationError("No se puede quitar el estado final porque la OS ya avanzó a salida o cierre")
        if estado_cur not in TECH_FINAL_REPAIR_STATES:
            raise ValidationError("Solo se puede quitar un estado final de reparación")

        if _rol(request) in ("tecnico", "jefe_veedor"):
            uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
            if int(row.get("asignado_a") or 0) != int(uid or 0):
                raise PermissionDenied("Solo el técnico asignado puede quitar el estado final")

        estado_next, _estado_transition = _estado_previo_a_estado(ingreso_id, estado_cur)
        presupuesto_next = "aprobado" if presupuesto_cur == "no_aplica" and _latest_quote_is_approved(ingreso_id) else presupuesto_actual
        _set_audit_user(request)
        # Al anular la resolución hay que soltar también el cambio de equipo: si no, la OS
        # queda arrastrando la serie de un cambio que ya no existe.
        resolucion_next = resolucion_al_anular_estado_final(presupuesto_next)
        exec_void(
            f"""
            UPDATE ingresos
               SET {sets_con_limpieza_de_cambio('estado=%s', 'resolucion=%s', 'presupuesto_estado=%s')}
             WHERE id=%s
            """,
            [estado_next, resolucion_next, presupuesto_next, ingreso_id],
        )
        return Response(
            {
                "ok": True,
                "estado_anterior": estado_cur,
                "estado": estado_next,
                "resolucion": resolucion_next,
                "presupuesto_estado": presupuesto_next,
            }
        )


class MarcarNoSeReparaView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, ingreso_id: int):
        require_roles(request, ["tecnico", "jefe", "jefe_veedor"])
        has_permite_reparacion = _ingresos_has_permite_reparacion_col()
        permite_reparacion_sql = (
            "COALESCE(permite_reparacion, TRUE) AS permite_reparacion"
            if has_permite_reparacion
            else "TRUE AS permite_reparacion"
        )
        try:
            row = q(
                f"""
                SELECT asignado_a, motivo, {permite_reparacion_sql}
                  FROM ingresos
                 WHERE id=%s
                """,
                [ingreso_id],
                one=True,
            )
            if not row:
                raise ValidationError("Ingreso no encontrado")
            if _rol(request) in ("tecnico", "jefe_veedor"):
                uid = getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None)
                if int(row.get("asignado_a") or 0) != int(uid or 0):
                    raise PermissionDenied("Solo el técnico asignado puede marcar como no se repara")
            if _is_cotizacion_bloqueada_para_reparar(row.get("motivo"), row.get("permite_reparacion")):
                raise ValidationError("Ingreso en cotización: primero habilite la reparación.")
        except PermissionDenied:
            raise
        except ValidationError:
            raise
        except Exception:
            if _rol(request) in ("tecnico", "jefe_veedor"):
                raise PermissionDenied("Solo el técnico asignado puede marcar como no se repara")
            raise

        _prev_estado = _estado_actual_txt(ingreso_id)
        _set_audit_user(request)
        # Igual que en "controlado sin defecto": el 'no_aplica' vive en `quotes` o no
        # vive. Y si el cliente rechazó un presupuesto, ese 'rechazado' se conserva:
        # taparlo con 'no_aplica' le hacía tomar a métricas el costo en vez del cobro.
        presupuesto_no_aplica = marcar_quote_no_aplica(ingreso_id)
        sets = sets_con_limpieza_de_cambio(
            "estado='no_se_repara'",
            "resolucion='no_reparado'",
            *(["presupuesto_estado='no_aplica'"] if presupuesto_no_aplica else []),
        )
        exec_void(f"UPDATE ingresos SET {sets} WHERE id=%s", [ingreso_id])
        if _prev_estado != "no_se_repara":
            registrar_ingreso_event(
                ingreso_id,
                "no_se_repara",
                de_estado=_prev_estado,
                usuario_id=getattr(getattr(request, "user", None), "id", None) or getattr(request, "user_id", None),
                comentario="Marcado como no se repara",
            )

        # El equipo ya no vuelve al banco: sólo falta la orden de salida.
        try:
            notify_no_se_repara_for_remito(ingreso_id, request=request)
        except Exception:
            logger.exception(
                "No se pudo avisar que el equipo sale sin reparar",
                extra={"ingreso_id": ingreso_id},
            )
        # Si el aviso venía de un rechazo con la reparación hecha, esto es
        # justamente lo que se le había pedido al técnico: se cierra solo.
        try:
            resolve_presupuesto_rechazado_reparado(ingreso_id)
        except Exception:
            logger.exception(
                "No se pudo cerrar el aviso de deshacer la reparación",
                extra={"ingreso_id": ingreso_id},
            )
        return Response({
            "ok": True,
            "estado": "no_se_repara",
            "resolucion": "no_reparado",
            "presupuesto_estado": _presupuesto_estado_actual_txt(ingreso_id),
        })


__all__ = [
    'MisPendientesView',
    'PendientesPresupuestoView',
    'PresupuestadosView',
    'AprobadosParaRepararView',
    'AprobadosYReparadosView',
    'LiberadosView',
    'GeneralEquiposView',
    'GeneralPorClienteView',
    'MarcarControladoSinDefectoView',
    'MarcarNoSeReparaView',
    'QuitarEstadoFinalReparacionView',
    'MarcarParaRepararView',
    'QuitarRepararView',
    'MarcarReparadoView',
    'HabilitarReparacionCotizacionView',
    'EntregarIngresoView',
    'DarBajaIngresoView',
    'DarAltaIngresoView',
    'IngresoCorreccionesHistoricasView',
    'IngresoBejermanEstadoView',
    'IngresoBejermanEstadoAccionView',
    'IngresoAsignarTecnicoView',
    'IngresoSolicitarAsignacionView',
    'IngresoSolicitarBajaView',
    'IngresoSolicitarBajaRechazarView',
    'PendientesGeneralView',
    'IngresoHistorialView',
    'CerrarReparacionView',
]
